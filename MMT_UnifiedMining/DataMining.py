from __future__ import annotations

import argparse
import ast
import csv
import hashlib
import importlib.util
import inspect
import json
import os
import re
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable


PIPELINE_CAN = "can"
PIPELINE_TOPIC = "topic"
METRIC_CATEGORY_MPI = "MPI"
METRIC_CATEGORY_CPI = "CPI"
DEFAULT_CROSS_DONE_METRICS = {"*"}
DEFAULT_H5_META_CACHE_SIZE = 512
DEFAULT_TOPIC_FRAMES_CACHE_SIZE = 2
RUN_STATE_DIR_NAME = ".run_state"
RUN_STATE_SCHEMA_VERSION = 1
RESUME_MODE_ASK = "ask"
RESUME_MODE_AUTO = "auto"
RESUME_MODE_NEVER = "never"
TASK_STATUS_RUNNING = "running"
TASK_STATUS_SUCCESS = "success"
TASK_STATUS_FAILED = "failed"
TASK_STATUS_SKIPPED_NO_H5 = "skipped_no_h5"
TASK_STATUS_SKIPPED_EMPTY_TOPIC_FRAMES = "skipped_empty_topic_frames"
TASK_STATUS_SKIPPED_MISSING_SIGNALS = "skipped_missing_signals"
COMPLETED_TASK_STATUSES = {
    TASK_STATUS_SUCCESS,
    TASK_STATUS_SKIPPED_NO_H5,
    TASK_STATUS_SKIPPED_EMPTY_TOPIC_FRAMES,
    TASK_STATUS_SKIPPED_MISSING_SIGNALS,
}
INCOMPLETE_TASK_STATUSES = {
    TASK_STATUS_RUNNING,
    TASK_STATUS_FAILED,
}

ODOMETER_DATASET_NAME = "Vehicle_Odometer"
ODOMETER_FIELD_NAME = "IVehOdo"
ODOMETER_VALID_FIELD_NAME = "IVehOdoV"
APA_CPI_TYPE_SIGNAL = "APA_Brake_Propulsion_Torqu_1F0_M.IAPABT_APAType"
APA_CPI_STATUS_SIGNAL = "Advanced_Park_Assist_Status_190_M.IAPASP_APAStsAuth"
DONE_FOLDER_TIME_PATTERN = re.compile(r"^(\d{8})-(\d{6})")
EVENT_TIME_DETAIL_KEYS = {
    "event_start",
    "event_end",
    "event_start_timestamp",
    "event_end_timestamp",
}
EVENT_TIME_TEXT_PATTERN = re.compile(
    r"(?P<label>\b(?:event_start|event_end|event_start_timestamp|event_end_timestamp)\b\s*[:=]\s*)"
    r"(?P<value>[-+]?\d+(?:\.\d+)?)",
    re.IGNORECASE,
)
EVENT_BLOCK_PATTERN = re.compile(r"\[Event[^\]]*\]", re.IGNORECASE)
EVENT_BLOCK_START_END_PATTERN = re.compile(
    r"(?P<label>\b(?:start|end)\s*=\s*)(?P<value>[-+]?\d+(?:\.\d+)?)",
    re.IGNORECASE,
)
_H5_CONTAINS_SIGNALS_CACHE: dict[tuple[tuple[str, int, int], tuple[str, ...]], bool] = {}
_H5_TIME_RANGE_CACHE: dict[tuple[str, int, int], tuple[float, float] | None] = {}
_H5_MILEAGE_CACHE: dict[tuple[str, int, int], float] = {}
_TOPIC_FRAMES_CACHE: dict[str, list[dict[str, Any]]] = {}
_TOPIC_SCENE_FINGERPRINT_CACHE: dict[int, tuple[object, str | None]] = {}
TOPIC_SCENE_FINGERPRINT_FUNCTIONS = (
    "_find_risk_segments",
    "_get_frame_risk_params",
    "_get_unp_set_speed",
    "_is_non_straight_driving_scene",
    "_is_primary_sequence_frame",
    "_is_set_speed_stable",
)


def _configure_console_encoding() -> None:
    for stream_name in ("stdout", "stderr"):
        stream = getattr(sys, stream_name, None)
        if hasattr(stream, "reconfigure"):
            try:
                stream.reconfigure(encoding="utf-8", errors="replace")
            except Exception:
                pass


_configure_console_encoding()


@dataclass
class PipelineConfig:
    name: str
    type: str
    enabled: bool
    func_folder_path: Path
    output_dir: Path
    file_pattern: str = "*.h5"
    preprocess_signals: list[str] = field(default_factory=list)
    default_metric_category: str = METRIC_CATEGORY_CPI
    enable_cross_done_continuation: bool = False
    cross_done_max_gap_seconds: float = 10.0
    cross_done_context_seconds: float = 10.0
    cross_done_metric_names: set[str] = field(default_factory=set)
    skip_empty_topic_frames: bool = True


@dataclass
class UnifiedConfig:
    source_path: str
    query_date: str
    query_software: str
    max_depth: int
    suffix: str
    output_dir: Path
    report_prefix: str
    pipelines: list[PipelineConfig]


@dataclass
class MetricModule:
    pipeline_name: str
    pipeline_type: str
    name: str
    file_name: str
    path: Path
    module: object
    runner: Callable[[Any], Any]
    category: str

    @property
    def output_prefix(self) -> str:
        return f"{self.pipeline_name}__{self.name}"


@dataclass
class BadCaseExtractRule:
    enabled: bool = True
    min_severity: float | None = 0.0
    top_n: int | None = None
    max_cases: int | None = 20
    sort_by: str = "severity_desc"


@dataclass
class BadCaseExtractConfig:
    enabled: bool
    output_file: str
    global_rule: BadCaseExtractRule
    pipeline_rules: dict[str, BadCaseExtractRule]
    metric_rules: dict[str, BadCaseExtractRule]


@dataclass
class RunStateSummary:
    expected_count: int
    completed_count: int
    success_count: int
    skipped_count: int
    failed_count: int
    running_count: int
    missing_count: int

    @property
    def pending_count(self) -> int:
        return self.failed_count + self.running_count + self.missing_count


@dataclass
class RunState:
    path: Path
    data: dict[str, Any]
    dirty: bool = field(default=False, init=False, repr=False)
    last_saved_monotonic: float = field(default_factory=time.monotonic, init=False, repr=False)

    @property
    def run_timestamp(self) -> str:
        return str(self.data["run_timestamp"])

    @property
    def tasks(self) -> dict[str, dict[str, Any]]:
        tasks = self.data.setdefault("tasks", {})
        if not isinstance(tasks, dict):
            tasks = {}
            self.data["tasks"] = tasks
        return tasks

    def task_key(self, pipeline: PipelineConfig, folder_path: str, metric: MetricModule) -> str:
        return make_task_key(pipeline.name, folder_path, metric.name)

    def should_skip(self, pipeline: PipelineConfig, folder_path: str, metric: MetricModule) -> bool:
        task = self.tasks.get(self.task_key(pipeline, folder_path, metric), {})
        return str(task.get("status", "")) in COMPLETED_TASK_STATUSES

    def start_task(self, pipeline: PipelineConfig, folder_path: str, metric: MetricModule) -> None:
        key = self.task_key(pipeline, folder_path, metric)
        now_text = now_iso()
        self.tasks[key] = {
            **self.tasks.get(key, {}),
            "status": TASK_STATUS_RUNNING,
            "pipeline": pipeline.name,
            "pipeline_type": pipeline.type,
            "folder_path": folder_path,
            "metric_name": metric.name,
            "metric_file": metric.file_name,
            "output_prefix": metric.output_prefix,
            "started_at": now_text,
            "updated_at": now_text,
            "finished_at": "",
            "error": "",
        }
        self.mark_dirty()

    def finish_task(
        self,
        pipeline: PipelineConfig,
        folder_path: str,
        metric: MetricModule,
        status: str,
        csv_path: Path | None = None,
        result: bool | None = None,
        details: list[Any] | None = None,
        error: str = "",
    ) -> None:
        key = self.task_key(pipeline, folder_path, metric)
        now_text = now_iso()
        message = summarize_details(details or [])
        self.tasks[key] = {
            **self.tasks.get(key, {}),
            "status": status,
            "pipeline": pipeline.name,
            "pipeline_type": pipeline.type,
            "folder_path": folder_path,
            "metric_name": metric.name,
            "metric_file": metric.file_name,
            "output_prefix": metric.output_prefix,
            "csv_path": str(csv_path) if csv_path is not None else self.tasks.get(key, {}).get("csv_path", ""),
            "result": result,
            "message": message,
            "error": error,
            "updated_at": now_text,
            "finished_at": now_text,
        }
        self.mark_dirty()

    def refresh_summary(self, expected_task_keys: set[str]) -> RunStateSummary:
        summary = summarize_run_state(self.data, expected_task_keys)
        self.data["expected_task_count"] = summary.expected_count
        self.data["completed_task_count"] = summary.completed_count
        self.data["success_task_count"] = summary.success_count
        self.data["skipped_task_count"] = summary.skipped_count
        self.data["failed_task_count"] = summary.failed_count
        self.data["running_task_count"] = summary.running_count
        self.data["missing_task_count"] = summary.missing_count
        self.data["updated_at"] = now_iso()
        return summary

    def save(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.data["updated_at"] = now_iso()
        temp_path = self.path.with_name(f"{self.path.name}.tmp")
        with temp_path.open("w", encoding="utf-8") as file:
            json.dump(self.data, file, ensure_ascii=False, indent=2, sort_keys=True)
        temp_path.replace(self.path)
        self.dirty = False
        self.last_saved_monotonic = time.monotonic()

    def mark_dirty(self) -> None:
        self.dirty = True
        try:
            flush_interval = max(1.0, float(os.environ.get("RUN_STATE_FLUSH_INTERVAL_SECONDS", "30")))
        except (TypeError, ValueError):
            flush_interval = 30.0
        if time.monotonic() - self.last_saved_monotonic >= flush_interval:
            self.save()

    def flush(self) -> None:
        if self.dirty:
            self.save()


def load_config(config_file_path: str | Path) -> UnifiedConfig:
    config_path = Path(config_file_path)
    if not config_path.exists():
        raise FileNotFoundError(f"Config file not found: {config_path}")

    with config_path.open("r", encoding="utf-8") as file:
        raw = json.load(file)

    base_dir = config_path.resolve().parent
    required_keys = ["source_path", "max_depth", "output_dir", "report_prefix", "pipelines"]
    for key in required_keys:
        if key not in raw:
            raise ValueError(f"Missing config key: {key}")

    query_date = str(raw.get("query_date", "")).strip()
    if query_date:
        date_pattern = r"^\d{8}-\d{8}$"
        if not re.match(date_pattern, query_date):
            raise ValueError("query_date格式必须为YYYYMMDD-YYYYMMDD")
        start_date = datetime.strptime(query_date.split("-")[0], "%Y%m%d")
        end_date = datetime.strptime(query_date.split("-")[1], "%Y%m%d")
        if start_date > end_date:
            raise ValueError("query_date起始日期不能晚于终止日期")

    max_depth = int(raw["max_depth"])
    if max_depth < 0:
        raise ValueError("max_depth must be >= 0")

    output_dir = resolve_path(raw["output_dir"], base_dir)
    pipeline_configs = parse_pipeline_configs(raw["pipelines"], raw, base_dir, output_dir)
    enabled_pipelines = [pipeline for pipeline in pipeline_configs if pipeline.enabled]
    if not enabled_pipelines:
        raise ValueError("No enabled pipeline found in config.json")

    return UnifiedConfig(
        source_path=str(raw["source_path"]),
        query_date=query_date,
        query_software=str(raw.get("query_software", "")).strip(),
        max_depth=max_depth,
        suffix=str(raw.get("suffix", "_done")),
        output_dir=output_dir,
        report_prefix=str(raw["report_prefix"]),
        pipelines=enabled_pipelines,
    )


def parse_pipeline_configs(
    pipelines_raw: Any,
    root_config: dict[str, Any],
    base_dir: Path,
    default_output_dir: Path,
) -> list[PipelineConfig]:
    if isinstance(pipelines_raw, dict):
        pipeline_items = []
        for name, item in pipelines_raw.items():
            if not isinstance(item, dict):
                raise ValueError(f"Pipeline config must be an object: {name}")
            item = dict(item)
            item.setdefault("name", name)
            pipeline_items.append(item)
    elif isinstance(pipelines_raw, list):
        pipeline_items = pipelines_raw
    else:
        raise ValueError("pipelines must be a list or object")

    parsed: list[PipelineConfig] = []
    for index, item in enumerate(pipeline_items, start=1):
        if not isinstance(item, dict):
            raise ValueError(f"Pipeline #{index} must be an object")

        name = str(item.get("name", "")).strip()
        pipeline_type = str(item.get("type", name)).strip().lower()
        if not name:
            raise ValueError(f"Pipeline #{index} missing name")
        if pipeline_type not in {PIPELINE_CAN, PIPELINE_TOPIC}:
            raise ValueError(f"Unsupported pipeline type for {name}: {pipeline_type}")
        if "func_folder_path" not in item:
            raise ValueError(f"Pipeline {name} missing func_folder_path")

        pipeline_output_dir = resolve_path(item.get("output_dir", str(default_output_dir)), base_dir)
        default_category = METRIC_CATEGORY_MPI if pipeline_type == PIPELINE_CAN else METRIC_CATEGORY_CPI
        cross_done_names = item.get("cross_done_metric_names", sorted(DEFAULT_CROSS_DONE_METRICS))
        if not isinstance(cross_done_names, list):
            raise ValueError(f"Pipeline {name}: cross_done_metric_names must be a list")

        parsed.append(
            PipelineConfig(
                name=name,
                type=pipeline_type,
                enabled=bool(item.get("enabled", True)),
                func_folder_path=resolve_path(item["func_folder_path"], base_dir),
                output_dir=pipeline_output_dir,
                file_pattern=str(item.get("file_pattern", root_config.get("file_pattern", "*.h5"))),
                preprocess_signals=list(item.get("preprocess_signals", [])),
                default_metric_category=str(item.get("default_metric_category", default_category)).upper(),
                enable_cross_done_continuation=bool(item.get("enable_cross_done_continuation", False)),
                cross_done_max_gap_seconds=float(item.get("cross_done_max_gap_seconds", 10)),
                cross_done_context_seconds=float(item.get("cross_done_context_seconds", 10)),
                cross_done_metric_names={str(name) for name in cross_done_names},
                skip_empty_topic_frames=bool(item.get("skip_empty_topic_frames", True)),
            )
        )

    return parsed


def resolve_path(path_value: str | Path, base_dir: Path) -> Path:
    path = Path(path_value)
    if path.is_absolute():
        return path
    return (base_dir / path).resolve()


def load_bad_case_extract_config(config_file_path: str | Path) -> BadCaseExtractConfig:
    config_path = Path(config_file_path)
    if not config_path.exists():
        print(f"未找到 Bad Case 提取配置：{config_path}，将使用默认规则")
        raw: dict[str, Any] = {}
    else:
        with config_path.open("r", encoding="utf-8") as file:
            raw = json.load(file)

    global_rule_raw = raw.get("global_rule", {})
    if not isinstance(global_rule_raw, dict):
        raise ValueError("bad_case_extract.json: global_rule must be an object")

    pipeline_rules_raw = raw.get("pipeline_rules", {})
    if not isinstance(pipeline_rules_raw, dict):
        raise ValueError("bad_case_extract.json: pipeline_rules must be an object")

    metric_rules_raw = raw.get("metric_rules", {})
    if not isinstance(metric_rules_raw, dict):
        raise ValueError("bad_case_extract.json: metric_rules must be an object")

    global_rule = parse_bad_case_extract_rule(global_rule_raw, BadCaseExtractRule())
    pipeline_rules = {
        str(name): parse_bad_case_extract_rule(rule_raw, global_rule)
        for name, rule_raw in pipeline_rules_raw.items()
        if isinstance(rule_raw, dict)
    }
    metric_rules = {
        str(name): parse_bad_case_extract_rule(rule_raw, global_rule)
        for name, rule_raw in metric_rules_raw.items()
        if isinstance(rule_raw, dict)
    }

    return BadCaseExtractConfig(
        enabled=bool(raw.get("enabled", True)),
        output_file=str(raw.get("output_file", "Bad_Case_Report")).strip() or "Bad_Case_Report",
        global_rule=global_rule,
        pipeline_rules=pipeline_rules,
        metric_rules=metric_rules,
    )


def parse_bad_case_extract_rule(raw: dict[str, Any], base_rule: BadCaseExtractRule) -> BadCaseExtractRule:
    rule = BadCaseExtractRule(
        enabled=base_rule.enabled,
        min_severity=base_rule.min_severity,
        top_n=base_rule.top_n,
        max_cases=base_rule.max_cases,
        sort_by=base_rule.sort_by,
    )

    if "enabled" in raw:
        rule.enabled = bool(raw["enabled"])
    if "min_severity" in raw:
        rule.min_severity = parse_optional_float(raw["min_severity"], "min_severity")
    if "top_n" in raw:
        rule.top_n = parse_optional_positive_int(raw["top_n"], "top_n")
    if "max_cases" in raw:
        rule.max_cases = parse_optional_positive_int(raw["max_cases"], "max_cases")
    if "sort_by" in raw:
        rule.sort_by = str(raw["sort_by"]).strip() or "severity_desc"
    if rule.sort_by not in {"severity_desc", "severity_asc", "path_asc"}:
        raise ValueError("bad_case_extract.json: sort_by must be one of severity_desc, severity_asc, path_asc")
    return rule


def parse_optional_float(value: Any, field_name: str) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except Exception as exc:
        raise ValueError(f"bad_case_extract.json: {field_name} must be numeric or null") from exc


def parse_optional_positive_int(value: Any, field_name: str) -> int | None:
    if value is None:
        return None
    try:
        parsed = int(value)
    except Exception as exc:
        raise ValueError(f"bad_case_extract.json: {field_name} must be an integer or null") from exc
    if parsed < 0:
        raise ValueError(f"bad_case_extract.json: {field_name} must be >= 0")
    return parsed


def load_metric_modules(pipeline: PipelineConfig) -> list[MetricModule]:
    func_dir = pipeline.func_folder_path
    if not func_dir.exists():
        raise FileNotFoundError(f"Pipeline {pipeline.name}: metric folder not found: {func_dir}")

    metric_modules: list[MetricModule] = []
    failed_files: list[tuple[Path, str]] = []
    for index, path in enumerate(sorted(func_dir.glob("*.py")), start=1):
        if path.name.startswith("__") or path.name.endswith("_scene.py"):
            continue

        module_name = f"unified_{pipeline.name}_{index}_{path.stem}"
        success, module_obj, message = load_single_metric_module(path, module_name)
        if not success or module_obj is None:
            failed_files.append((path, message))
            print(f"❌ [{pipeline.name}] 加载失败：{path.name} - {message}")
            continue

        runner = getattr(module_obj, "__run_haoran_data_run__", None)
        if runner is None or not callable(runner):
            failed_files.append((path, "未找到可调用的__run_haoran_data_run__函数"))
            print(f"❌ [{pipeline.name}] 加载失败：{path.name} - 未找到__run_haoran_data_run__")
            continue

        category = get_metric_category(module_obj, path.name, pipeline)
        metric_modules.append(
            MetricModule(
                pipeline_name=pipeline.name,
                pipeline_type=pipeline.type,
                name=path.stem,
                file_name=path.name,
                path=path,
                module=module_obj,
                runner=runner,
                category=category,
            )
        )
        print(f"✅ [{pipeline.name}] 加载成功：{path.name}")

    if failed_files:
        print(f"⚠️ [{pipeline.name}] 有 {len(failed_files)} 个指标加载失败，已跳过")
    if not metric_modules:
        raise RuntimeError(f"Pipeline {pipeline.name}: no metric module loaded")
    return metric_modules


def load_single_metric_module(path: Path, module_name: str) -> tuple[bool, object | None, str]:
    spec = importlib.util.spec_from_file_location(module_name, path)
    if spec is None or spec.loader is None:
        return False, None, f"无法创建模块加载器：{path}"

    module = importlib.util.module_from_spec(spec)
    try:
        spec.loader.exec_module(module)
    except Exception as exc:
        return False, None, f"加载函数失败：{exc}"
    return True, module, "函数加载成功"


def get_metric_category(module: object, metric_file_name: str, pipeline: PipelineConfig) -> str:
    category = str(getattr(module, "METRIC_CATEGORY", "")).strip().upper()
    if category in {METRIC_CATEGORY_MPI, METRIC_CATEGORY_CPI}:
        return category
    if pipeline.type == PIPELINE_CAN and Path(metric_file_name).stem.startswith("apa_"):
        return METRIC_CATEGORY_CPI
    return pipeline.default_metric_category


def validate_can_required_signals(metrics: list[MetricModule], preprocess_signals: list[str]) -> None:
    configured = set(preprocess_signals)
    missing_by_file: dict[str, list[str]] = {}

    for metric in metrics:
        required = list(getattr(metric.module, "REQUIRED_SIGNALS", []))
        missing = [signal for signal in required if signal not in configured]
        if missing:
            missing_by_file[metric.file_name] = missing

    if missing_by_file:
        lines = ["以下 CAN 指标依赖的信号未写入 config.json 的 preprocess_signals："]
        for func_name, missing in missing_by_file.items():
            lines.append(f"- {func_name}: {', '.join(missing)}")
        raise ValueError("\n".join(lines))


def collect_eligible_folders(
    base_path: str | Path,
    max_depth: int,
    suffix: str,
    query_date: str,
    query_software: str,
) -> list[str]:
    base = Path(base_path)
    if not base.exists():
        print(f"警告：路径不存在 {base}")
        return []

    start_date = None
    end_date = None
    if query_date:
        start_date = datetime.strptime(query_date.split("-")[0], "%Y%m%d")
        end_date = datetime.strptime(query_date.split("-")[1], "%Y%m%d")

    eligible: list[str] = []
    scanned_folder_count = 0
    base_depth = len(base.resolve().parts)
    print(f"开始搜索候选文件夹：{base}")
    for root, dirs, _files in os.walk(base):
        scanned_folder_count += 1
        if scanned_folder_count == 1 or scanned_folder_count % 100 == 0:
            print(f"[搜索进度] 已扫描 {scanned_folder_count} 个目录，已命中 {len(eligible)} 个候选文件夹")

        current = Path(root)
        current_depth = len(current.resolve().parts) - base_depth
        if current_depth >= max_depth:
            dirs.clear()
            continue

        folder_name = current.name
        if suffix and not folder_name.endswith(suffix):
            continue
        if get_immediate_subfolders(current):
            continue

        if start_date and end_date:
            if len(folder_name) < 8:
                continue
            try:
                folder_date = datetime.strptime(folder_name[:8], "%Y%m%d")
            except ValueError:
                continue
            if not (start_date <= folder_date <= end_date):
                continue

        if query_software and not folder_matches_software(current, query_software):
            continue

        eligible.append(str(current))
        print(f"✅ 命中候选文件夹：{current}")

    print(f"搜索结束，共扫描 {scanned_folder_count} 个目录，命中 {len(eligible)} 个候选文件夹")
    return sort_done_folders(sorted(set(eligible)))


def get_immediate_subfolders(folder_path: str | Path) -> list[str]:
    folder = Path(folder_path)
    try:
        return sorted(str(path) for path in folder.iterdir() if path.is_dir())
    except FileNotFoundError:
        return []


def folder_matches_software(folder_path: Path, query_software: str) -> bool:
    info_file = folder_path / "_cfis_info.txt"
    if not info_file.exists():
        return False

    try:
        content = read_text_fallback(info_file)
    except Exception:
        return False

    match = re.search(r'(QWANGS[^"]+\.tgz)', content, re.DOTALL)
    if not match:
        return False
    return query_software in match.group(1)


def read_text_fallback(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return path.read_text(encoding="gbk", errors="ignore")


def parse_done_folder_time(folder_path: str) -> datetime | None:
    folder_name = os.path.basename(folder_path)
    match = DONE_FOLDER_TIME_PATTERN.match(folder_name)
    if not match:
        return None
    try:
        return datetime.strptime("".join(match.groups()), "%Y%m%d%H%M%S")
    except ValueError:
        return None


def sort_done_folders(folders: list[str]) -> list[str]:
    def sort_key(folder_path: str) -> tuple[str, datetime, str, str]:
        parsed_time = parse_done_folder_time(folder_path)
        time_key = parsed_time if parsed_time is not None else datetime.max
        return (os.path.dirname(folder_path), time_key, os.path.basename(folder_path), folder_path)

    return sorted(folders, key=sort_key)


def is_next_done_continuous(
    current_folder: str,
    next_folder: str,
    max_gap_seconds: float,
) -> tuple[bool, float | None]:
    if os.path.dirname(current_folder) != os.path.dirname(next_folder):
        return False, None

    current_time = parse_done_folder_time(current_folder)
    next_time = parse_done_folder_time(next_folder)
    if current_time is None or next_time is None:
        return False, None

    gap_seconds = (next_time - current_time).total_seconds()
    return 0 < gap_seconds <= max_gap_seconds, gap_seconds


def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def normalize_state_path(path_value: str | Path) -> str:
    try:
        return os.path.normcase(str(Path(path_value).resolve()))
    except Exception:
        return os.path.normcase(os.path.abspath(str(path_value)))


def make_task_key(pipeline_name: str, folder_path: str, metric_name: str) -> str:
    return f"{pipeline_name}|{normalize_state_path(folder_path)}|{metric_name}"


def summarize_details(details: list[Any], max_length: int = 500) -> str:
    text = "; ".join(str(item) for item in details)
    if len(text) <= max_length:
        return text
    return f"{text[:max_length - 3]}..."


def hash_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def hash_text_items(items: list[str]) -> str:
    payload = "\n".join(items)
    return hash_text(payload)


def calculate_config_hash(config_path: Path) -> str:
    with config_path.open("r", encoding="utf-8") as file:
        raw = json.load(file)
    payload = json.dumps(raw, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hash_text(payload)


def calculate_metric_code_hash(metrics_by_pipeline: dict[str, list[MetricModule]]) -> str:
    base_dir = Path(__file__).resolve().parent
    paths: set[Path] = {
        Path(__file__).resolve(),
        base_dir / "signal_lib.py",
        base_dir / "longitudinal_preprocessor.py",
    }
    for metrics in metrics_by_pipeline.values():
        for metric in metrics:
            paths.add(metric.path.resolve())

    digest = hashlib.sha256()
    for path in sorted(paths, key=lambda item: str(item)):
        if not path.exists() or not path.is_file():
            continue
        digest.update(str(path).encode("utf-8"))
        digest.update(b"\0")
        digest.update(path.read_bytes())
        digest.update(b"\0")
    return digest.hexdigest()


def calculate_folder_list_hash(eligible_folders: list[str]) -> str:
    return hash_text_items([normalize_state_path(folder) for folder in eligible_folders])


def build_expected_task_keys(
    config: UnifiedConfig,
    metrics_by_pipeline: dict[str, list[MetricModule]],
    eligible_folders: list[str],
) -> set[str]:
    expected: set[str] = set()
    for folder_path in eligible_folders:
        for pipeline in config.pipelines:
            for metric in metrics_by_pipeline.get(pipeline.name, []):
                expected.add(make_task_key(pipeline.name, folder_path, metric.name))
    return expected


def summarize_run_state(data: dict[str, Any], expected_task_keys: set[str]) -> RunStateSummary:
    tasks = data.get("tasks", {})
    if not isinstance(tasks, dict):
        tasks = {}

    completed_count = 0
    success_count = 0
    skipped_count = 0
    failed_count = 0
    running_count = 0
    missing_count = 0
    for key in expected_task_keys:
        task = tasks.get(key)
        if not isinstance(task, dict):
            missing_count += 1
            continue

        status = str(task.get("status", ""))
        if status == TASK_STATUS_SUCCESS:
            completed_count += 1
            success_count += 1
        elif status in COMPLETED_TASK_STATUSES:
            completed_count += 1
            skipped_count += 1
        elif status == TASK_STATUS_FAILED:
            failed_count += 1
        elif status == TASK_STATUS_RUNNING:
            running_count += 1
        else:
            missing_count += 1

    return RunStateSummary(
        expected_count=len(expected_task_keys),
        completed_count=completed_count,
        success_count=success_count,
        skipped_count=skipped_count,
        failed_count=failed_count,
        running_count=running_count,
        missing_count=missing_count,
    )


def load_run_state(path: Path) -> dict[str, Any] | None:
    try:
        with path.open("r", encoding="utf-8") as file:
            data = json.load(file)
    except Exception:
        return None
    if not isinstance(data, dict):
        return None
    if int(data.get("schema_version", 0) or 0) != RUN_STATE_SCHEMA_VERSION:
        return None
    return data


def find_latest_resume_candidate(output_dir: Path, config_hash: str) -> tuple[Path, dict[str, Any]] | None:
    state_dir = output_dir / RUN_STATE_DIR_NAME
    if not state_dir.exists():
        return None

    candidates: list[tuple[float, Path, dict[str, Any]]] = []
    for path in state_dir.glob("run_*.json"):
        data = load_run_state(path)
        if data is None or data.get("config_hash") != config_hash:
            continue
        candidates.append((path.stat().st_mtime, path, data))

    if not candidates:
        return None
    _, path, data = max(candidates, key=lambda item: item[0])
    return path, data


def should_resume_state(
    state_path: Path,
    state_data: dict[str, Any],
    summary: RunStateSummary,
    metric_code_hash: str,
    folder_list_hash: str,
    resume_mode: str,
) -> bool:
    if summary.pending_count <= 0:
        print(f"检测到相同配置的运行状态已全部完成：{state_path.name}，本次将开始新运行。")
        return False

    code_changed = state_data.get("metric_code_hash") != metric_code_hash
    folder_list_changed = state_data.get("folder_list_hash") != folder_list_hash
    if resume_mode == RESUME_MODE_NEVER:
        return False

    if resume_mode == RESUME_MODE_AUTO:
        if code_changed or folder_list_changed:
            print("检测到可续跑状态，但指标代码或本次文件夹列表已变化，--resume auto 不自动续跑。")
            return False
        return True

    print("检测到上一次相同 config.json 的未完成运行状态：")
    print(f"  状态文件：{state_path}")
    print(f"  run_timestamp：{state_data.get('run_timestamp', '')}")
    print(
        "  进度："
        f"{summary.completed_count}/{summary.expected_count} 已完成 "
        f"(success={summary.success_count}, skipped={summary.skipped_count}, "
        f"failed={summary.failed_count}, running={summary.running_count}, missing={summary.missing_count})"
    )
    if code_changed:
        print("  ⚠️ 指标代码 hash 与上次不同；继续会复用已完成旧结果，只补跑失败/未跑任务。")
    if folder_list_changed:
        print("  ⚠️ 本次扫描到的文件夹列表与上次不同；继续可能保留上次 CSV 中已移除文件夹的结果。")

    default_yes = not code_changed and not folder_list_changed
    prompt = "是否继续上一次运行状态？[Y/n] " if default_yes else "是否仍要继续上一次运行状态？[y/N] "
    try:
        answer = input(prompt).strip().lower()
    except EOFError:
        answer = ""

    if not answer:
        return default_yes
    return answer in {"y", "yes", "1", "true", "是", "继续"}


def prepare_run_state(
    config: UnifiedConfig,
    config_path: Path,
    metrics_by_pipeline: dict[str, list[MetricModule]],
    eligible_folders: list[str],
    resume_mode: str,
) -> tuple[RunState, bool, set[str]]:
    config_hash = calculate_config_hash(config_path)
    metric_code_hash = calculate_metric_code_hash(metrics_by_pipeline)
    folder_list_hash = calculate_folder_list_hash(eligible_folders)
    expected_task_keys = build_expected_task_keys(config, metrics_by_pipeline, eligible_folders)

    config.output_dir.mkdir(parents=True, exist_ok=True)
    state_dir = config.output_dir / RUN_STATE_DIR_NAME
    state_dir.mkdir(parents=True, exist_ok=True)

    candidate = find_latest_resume_candidate(config.output_dir, config_hash)
    if candidate is not None:
        candidate_path, candidate_data = candidate
        summary = summarize_run_state(candidate_data, expected_task_keys)
        if should_resume_state(
            candidate_path,
            candidate_data,
            summary,
            metric_code_hash,
            folder_list_hash,
            resume_mode,
        ):
            candidate_data["status"] = TASK_STATUS_RUNNING
            candidate_data["last_resume_at"] = now_iso()
            candidate_data["resume_count"] = int(candidate_data.get("resume_count", 0) or 0) + 1
            candidate_data["last_metric_code_hash"] = metric_code_hash
            candidate_data["last_folder_list_hash"] = folder_list_hash
            run_state = RunState(candidate_path, candidate_data)
            run_state.refresh_summary(expected_task_keys)
            run_state.save()
            print(f"继续上一次运行：run_timestamp={run_state.run_timestamp}")
            return run_state, True, expected_task_keys

    run_timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    state_path = state_dir / f"run_{run_timestamp}.json"
    state_data: dict[str, Any] = {
        "schema_version": RUN_STATE_SCHEMA_VERSION,
        "status": TASK_STATUS_RUNNING,
        "run_timestamp": run_timestamp,
        "created_at": now_iso(),
        "updated_at": now_iso(),
        "finished_at": "",
        "config_path": str(config_path),
        "config_hash": config_hash,
        "metric_code_hash": metric_code_hash,
        "last_metric_code_hash": metric_code_hash,
        "folder_list_hash": folder_list_hash,
        "last_folder_list_hash": folder_list_hash,
        "source_path": config.source_path,
        "query_date": config.query_date,
        "query_software": config.query_software,
        "suffix": config.suffix,
        "tasks": {},
    }
    run_state = RunState(state_path, state_data)
    run_state.refresh_summary(expected_task_keys)
    run_state.save()
    print(f"创建新的运行状态：{run_state.path}")
    return run_state, False, expected_task_keys


def finalize_run_state(run_state: RunState, expected_task_keys: set[str]) -> RunStateSummary:
    summary = run_state.refresh_summary(expected_task_keys)
    run_state.data["status"] = "completed" if summary.pending_count == 0 else "incomplete"
    run_state.data["finished_at"] = now_iso()
    run_state.save()
    print(
        "运行状态已更新："
        f"{summary.completed_count}/{summary.expected_count} 已完成，"
        f"失败 {summary.failed_count}，未跑 {summary.missing_count}，状态文件：{run_state.path}"
    )
    return summary


def run_all(
    config: UnifiedConfig,
    dry_run: bool = False,
    list_metrics: bool = False,
    bad_case_config_path: Path | None = None,
    config_path: Path | None = None,
    resume_mode: str = RESUME_MODE_ASK,
) -> int:
    ensure_sys_path(Path(__file__).resolve().parent)
    if config_path is None:
        config_path = Path(__file__).resolve().parent / "config.json"
    if bad_case_config_path is None:
        bad_case_config_path = Path(__file__).resolve().parent / "bad_case_extract.json"
    bad_case_extract_config = load_bad_case_extract_config(bad_case_config_path)

    metrics_by_pipeline: dict[str, list[MetricModule]] = {}
    metric_category_map: dict[str, str] = {}
    print("========== 开始加载指标函数 ==========")
    for pipeline in config.pipelines:
        print(f"--- Pipeline: {pipeline.name} ({pipeline.type}) ---")
        metrics = load_metric_modules(pipeline)
        if pipeline.type == PIPELINE_CAN:
            validate_can_required_signals(metrics, pipeline.preprocess_signals)
        metrics_by_pipeline[pipeline.name] = metrics
        for metric in metrics:
            metric_category_map[metric.output_prefix] = metric.category
    print("========== 指标函数加载完成 ==========")

    if list_metrics or dry_run:
        for pipeline in config.pipelines:
            metrics = metrics_by_pipeline[pipeline.name]
            print(f"[{pipeline.name}] {len(metrics)} 个指标：")
            for metric in metrics:
                print(f"  - {metric.name} ({metric.category})")
        if dry_run:
            if bad_case_extract_config.enabled:
                print(f"Bad Case 提取配置加载正常：{bad_case_config_path}")
            else:
                print(f"Bad Case 提取已禁用：{bad_case_config_path}")
            print("Dry run 完成：配置与指标加载正常，未扫描数据目录。")
            return 0

    eligible_folders = collect_eligible_folders(
        base_path=config.source_path,
        max_depth=config.max_depth,
        suffix=config.suffix,
        query_date=config.query_date,
        query_software=config.query_software,
    )
    print(f"共找到 {len(eligible_folders)} 个符合条件的文件夹")
    if not eligible_folders:
        print("未找到候选文件夹，程序结束。")
        return 0

    config.output_dir.mkdir(parents=True, exist_ok=True)
    for pipeline in config.pipelines:
        pipeline.output_dir.mkdir(parents=True, exist_ok=True)

    run_state, is_resumed_run, expected_task_keys = prepare_run_state(
        config=config,
        config_path=config_path,
        metrics_by_pipeline=metrics_by_pipeline,
        eligible_folders=eligible_folders,
        resume_mode=resume_mode,
    )
    run_timestamp = run_state.run_timestamp
    start_time = time.time()
    total_folders = len(eligible_folders)
    print("========== 开始按文件夹执行所有指标 ==========")

    try:
        for folder_index, folder_path in enumerate(eligible_folders, start=1):
            print(f"\n========== [{folder_index}/{total_folders}] {folder_path} ==========")
            for pipeline in config.pipelines:
                metrics = metrics_by_pipeline[pipeline.name]
                next_folder_path = find_next_folder_for_pipeline(
                    pipeline=pipeline,
                    folders=eligible_folders,
                    folder_index=folder_index,
                )
                if pipeline.type == PIPELINE_CAN:
                    process_can_folder(
                        pipeline=pipeline,
                        metrics=metrics,
                        folder_path=folder_path,
                        folder_index=folder_index,
                        total_folders=total_folders,
                        run_timestamp=run_timestamp,
                        next_folder_path=next_folder_path,
                        run_state=run_state,
                        replace_existing_results=is_resumed_run,
                    )
                elif pipeline.type == PIPELINE_TOPIC:
                    process_topic_folder(
                        pipeline=pipeline,
                        metrics=metrics,
                        folder_path=folder_path,
                        folder_index=folder_index,
                        total_folders=total_folders,
                        run_timestamp=run_timestamp,
                        next_folder_path=next_folder_path,
                        run_state=run_state,
                        replace_existing_results=is_resumed_run,
                    )
            run_state.flush()
    finally:
        run_state.flush()

    report_path = generate_unified_report(
        output_dir=config.output_dir,
        report_prefix=config.report_prefix,
        run_timestamp=run_timestamp,
        metric_category_map=metric_category_map,
    )
    bad_case_report_path = generate_bad_case_extract_report(
        output_dir=config.output_dir,
        run_timestamp=run_timestamp,
        extract_config=bad_case_extract_config,
    )
    finalize_run_state(run_state, expected_task_keys)
    elapsed = time.time() - start_time
    print("========== 全部处理完成 ==========")
    print(f"共处理 {total_folders} 个文件夹，用时 {elapsed:.1f} 秒")
    if report_path:
        print(f"统一统计报告：{report_path}")
    if bad_case_report_path:
        print(f"Bad Case 提取报告：{bad_case_report_path}")
    return 0


def ensure_sys_path(base_dir: Path) -> None:
    for path in (base_dir, base_dir / "metrics" / "can" / "code", base_dir / "metrics" / "topic" / "code"):
        text = str(path)
        if text not in sys.path:
            sys.path.insert(0, text)


def find_next_folder_for_pipeline(
    pipeline: PipelineConfig,
    folders: list[str],
    folder_index: int,
) -> str | None:
    if not pipeline.enable_cross_done_continuation:
        return None
    if folder_index >= len(folders):
        return None

    current_folder = folders[folder_index - 1]
    candidate_next_folder = folders[folder_index]
    if os.path.dirname(current_folder) != os.path.dirname(candidate_next_folder):
        return None

    is_continuous, gap_seconds = is_next_done_continuous(
        current_folder,
        candidate_next_folder,
        pipeline.cross_done_max_gap_seconds,
    )

    print(
        f"[{pipeline.name}] Cross done candidate: {os.path.basename(current_folder)} -> "
        f"{os.path.basename(candidate_next_folder)}"
        + (f" ({gap_seconds:.1f}s)" if gap_seconds is not None else "")
    )
    return candidate_next_folder


def filter_pending_metrics(
    run_state: RunState | None,
    pipeline: PipelineConfig,
    folder_path: str,
    metrics: list[MetricModule],
) -> list[MetricModule]:
    pending: list[MetricModule] = []
    for metric in metrics:
        if run_state is not None and run_state.should_skip(pipeline, folder_path, metric):
            print(f"[{pipeline.name}]  ⏭️ 跳过已完成：{metric.file_name}")
            continue
        pending.append(metric)
    return pending


def process_can_folder(
    pipeline: PipelineConfig,
    metrics: list[MetricModule],
    folder_path: str,
    folder_index: int,
    total_folders: int,
    run_timestamp: str,
    next_folder_path: str | None = None,
    run_state: RunState | None = None,
    replace_existing_results: bool = False,
) -> None:
    print(f"[{pipeline.name}] [{folder_index}/{total_folders}] 开始 CAN 指标")
    pending_metrics = filter_pending_metrics(run_state, pipeline, folder_path, metrics)
    if not pending_metrics:
        print(f"[{pipeline.name}] 本文件夹 CAN 指标均已完成，跳过")
        return

    h5_files = sorted(str(path) for path in Path(folder_path).glob(pipeline.file_pattern) if path.is_file())
    if not h5_files:
        print(f"[{pipeline.name}] 跳过：未找到匹配文件 {pipeline.file_pattern}")
        for metric in pending_metrics:
            if run_state is not None:
                run_state.finish_task(
                    pipeline=pipeline,
                    folder_path=folder_path,
                    metric=metric,
                    status=TASK_STATUS_SKIPPED_NO_H5,
                    details=[f"未找到匹配文件 {pipeline.file_pattern}"],
                )
        return

    try:
        folder_context = build_can_folder_context(
            folder_path=folder_path,
            file_pattern=pipeline.file_pattern,
            preprocess_signals=pipeline.preprocess_signals,
            h5_files=h5_files,
            next_folder_path=next_folder_path,
            cross_done_max_gap_seconds=pipeline.cross_done_max_gap_seconds,
            cross_done_context_seconds=pipeline.cross_done_context_seconds,
        )
    except Exception as exc:
        print(f"[{pipeline.name}] 构建 CAN 会话失败：{exc}")
        for metric in pending_metrics:
            csv_path = save_metric_result(
                pipeline=pipeline,
                metric=metric,
                folder_path=folder_path,
                result=False,
                details=[f"构建 CAN 会话失败: {exc}"],
                statistics={},
                mileage_km=0.0,
                run_timestamp=run_timestamp,
                replace_existing=replace_existing_results,
            )
            if run_state is not None:
                run_state.finish_task(
                    pipeline=pipeline,
                    folder_path=folder_path,
                    metric=metric,
                    status=TASK_STATUS_FAILED,
                    csv_path=csv_path,
                    result=False,
                    details=[f"构建 CAN 会话失败: {exc}"],
                    error=str(exc),
                )
        return

    session = folder_context.get("session")
    folder_mileage_km = calculate_folder_mileage_km(folder_context.get("primary_h5_file"))
    segment_start_timestamp = calculate_h5_start_timestamp(folder_context.get("primary_h5_file"))
    apa_total_count = _safe_int(folder_context.get("apa_total_count", 0))

    for metric in pending_metrics:
        print(f"[{pipeline.name}]  └── 执行：{metric.file_name}")
        if run_state is not None:
            run_state.start_task(pipeline, folder_path, metric)
        required_signals = list(getattr(metric.module, "REQUIRED_SIGNALS", []))
        missing_signals = session.missing_signals(required_signals) if session is not None else required_signals
        if missing_signals:
            details = [f"缺少信号: {', '.join(missing_signals)}"]
            statistics = {
                "All Case Count": apa_total_count if metric.category == METRIC_CATEGORY_CPI else 0,
                "Bad Case": 0,
                "Severity Score": 0.0,
            }
            csv_path = save_metric_result(
                pipeline=pipeline,
                metric=metric,
                folder_path=folder_path,
                result=False,
                details=details,
                statistics=statistics,
                mileage_km=folder_mileage_km,
                run_timestamp=run_timestamp,
                segment_start_timestamp=segment_start_timestamp,
                replace_existing=replace_existing_results,
            )
            if run_state is not None:
                run_state.finish_task(
                    pipeline=pipeline,
                    folder_path=folder_path,
                    metric=metric,
                    status=TASK_STATUS_SKIPPED_MISSING_SIGNALS,
                    csv_path=csv_path,
                    result=False,
                    details=details,
                )
            print(f"[{pipeline.name}]  ⚠️ {metric.file_name} 跳过：{details[0]}")
            continue

        try:
            check_result, check_details, check_statistics = normalize_metric_result(metric.runner(folder_context))
            if metric.category == METRIC_CATEGORY_CPI:
                check_statistics = dict(check_statistics or {})
                check_statistics["All Case Count"] = apa_total_count
            csv_path = save_metric_result(
                pipeline=pipeline,
                metric=metric,
                folder_path=folder_path,
                result=check_result,
                details=check_details,
                statistics=check_statistics,
                mileage_km=folder_mileage_km,
                run_timestamp=run_timestamp,
                segment_start_timestamp=segment_start_timestamp,
                replace_existing=replace_existing_results,
            )
            if run_state is not None:
                run_state.finish_task(
                    pipeline=pipeline,
                    folder_path=folder_path,
                    metric=metric,
                    status=TASK_STATUS_SUCCESS,
                    csv_path=csv_path,
                    result=check_result,
                    details=check_details,
                )
            print(f"[{pipeline.name}]  ✅ {metric.file_name} 完成")
        except Exception as exc:
            csv_path = save_metric_result(
                pipeline=pipeline,
                metric=metric,
                folder_path=folder_path,
                result=False,
                details=[f"执行失败: {exc}"],
                statistics={},
                mileage_km=folder_mileage_km,
                run_timestamp=run_timestamp,
                segment_start_timestamp=segment_start_timestamp,
                replace_existing=replace_existing_results,
            )
            if run_state is not None:
                run_state.finish_task(
                    pipeline=pipeline,
                    folder_path=folder_path,
                    metric=metric,
                    status=TASK_STATUS_FAILED,
                    csv_path=csv_path,
                    result=False,
                    details=[f"执行失败: {exc}"],
                    error=str(exc),
                )
            print(f"[{pipeline.name}]  ❌ {metric.file_name} 执行失败：{exc}")


def build_can_folder_context(
    folder_path: str,
    file_pattern: str,
    preprocess_signals: list[str],
    h5_files: list[str] | None = None,
    next_folder_path: str | None = None,
    cross_done_max_gap_seconds: float = 10.0,
    cross_done_context_seconds: float = 10.0,
) -> dict[str, Any]:
    from signal_lib import FolderSession

    folder = Path(folder_path)
    if h5_files is None:
        h5_files = sorted(str(path) for path in folder.glob(file_pattern) if path.is_file())
    else:
        h5_files = list(h5_files)
    primary_h5_file = select_primary_h5_file(h5_files, preprocess_signals)
    additional_h5_files: list[str] = []
    context_folder_paths = [str(folder)]
    context_end_timestamp: float | None = None
    primary_h5_range = calculate_h5_time_range(primary_h5_file)

    if primary_h5_file and next_folder_path:
        next_folder = Path(next_folder_path)
        next_h5_files = sorted(str(path) for path in next_folder.glob(file_pattern) if path.is_file())
        next_primary_h5_file = select_primary_h5_file(next_h5_files, preprocess_signals)
        if next_primary_h5_file and h5_files_are_continuous(
            primary_h5_file,
            next_primary_h5_file,
            cross_done_max_gap_seconds,
        ):
            additional_h5_files.append(next_primary_h5_file)
            context_folder_paths.append(str(next_folder))
            if primary_h5_range is not None:
                context_end_timestamp = primary_h5_range[1] + float(cross_done_context_seconds)
            print(
                f"[can] 使用相邻 H5 续接：{Path(primary_h5_file).name} + "
                f"{next_folder.name}/{Path(next_primary_h5_file).name} "
                f"(context={float(cross_done_context_seconds):.1f}s)"
            )

    folder_context: dict[str, Any] = {
        "folder_path": str(folder),
        "context_folder_paths": context_folder_paths,
        "h5_files": h5_files,
        "primary_h5_file": primary_h5_file,
        "additional_h5_files": additional_h5_files,
        "context_h5_files": [primary_h5_file] + additional_h5_files if primary_h5_file else [],
        "context_end_timestamp": context_end_timestamp,
        "session": None,
    }
    if primary_h5_file:
        folder_context["session"] = FolderSession(
            primary_h5_file,
            preprocess_signals=preprocess_signals,
            additional_h5_files=additional_h5_files,
            context_end_timestamp=context_end_timestamp,
        )
    if primary_h5_file and additional_h5_files and primary_h5_range is not None:
        folder_context["apa_total_count"] = calculate_apa_total_count(
            folder_context.get("session"),
            start_timestamp=primary_h5_range[0],
            end_timestamp=primary_h5_range[1],
        )
    else:
        folder_context["apa_total_count"] = calculate_apa_total_count(folder_context.get("session"))
    return folder_context


def select_primary_h5_file(h5_files: list[str], preprocess_signals: list[str]) -> str | None:
    if not h5_files:
        return None

    preferred = None
    for h5_file in h5_files:
        if Path(h5_file).name == "_msd_can_data.h5":
            preferred = h5_file
            break

    if preferred and h5_contains_all_signals(preferred, preprocess_signals):
        return preferred

    for h5_file in h5_files:
        if h5_file == preferred:
            continue
        if h5_contains_all_signals(h5_file, preprocess_signals):
            return h5_file

    return preferred or h5_files[0]


def _h5_file_cache_key(h5_path: str | Path | None) -> tuple[str, int, int] | None:
    if not h5_path:
        return None
    try:
        path = Path(h5_path)
        stat = path.stat()
        return str(path.resolve()), int(stat.st_mtime_ns), int(stat.st_size)
    except Exception:
        return None


def _remember_limited(cache: dict[Any, Any], key: Any, value: Any, max_size: int) -> None:
    if max_size <= 0:
        return
    if key in cache:
        cache.pop(key)
    cache[key] = value
    while len(cache) > max_size:
        cache.pop(next(iter(cache)))


def _env_int(name: str, default: int) -> int:
    try:
        return max(0, int(os.environ.get(name, str(default))))
    except Exception:
        return default


def _h5_meta_cache_size() -> int:
    return _env_int("H5_META_CACHE_SIZE", DEFAULT_H5_META_CACHE_SIZE)


def _topic_frames_cache_size() -> int:
    return _env_int("TOPIC_FRAMES_CACHE_SIZE", DEFAULT_TOPIC_FRAMES_CACHE_SIZE)


def h5_contains_all_signals(h5_path: str | Path, signal_names: list[str]) -> bool:
    if not signal_names:
        return True

    file_key = _h5_file_cache_key(h5_path)
    signal_key = tuple(sorted({str(signal_name) for signal_name in signal_names}))
    cache_key = (file_key, signal_key) if file_key is not None else None
    if cache_key is not None and cache_key in _H5_CONTAINS_SIGNALS_CACHE:
        return _H5_CONTAINS_SIGNALS_CACHE[cache_key]

    result = False
    try:
        import h5py

        remaining = set(signal_key)
        with h5py.File(h5_path, "r") as h5_file:
            root_names = list(h5_file.keys())
            if not root_names:
                result = False
            else:
                channel = h5_file[root_names[0]].get("Channel1")
                if channel is None:
                    result = False
                else:
                    for message_name in channel.keys():
                        dataset = channel[message_name]
                        for field_name in dataset.dtype.names or ():
                            if field_name == "timestamp":
                                continue
                            remaining.discard(f"{message_name}.{field_name}")
                            remaining.discard(str(field_name))
                        if not remaining:
                            result = True
                            break
                    else:
                        result = not remaining
    except Exception:
        result = False

    if cache_key is not None:
        _remember_limited(_H5_CONTAINS_SIGNALS_CACHE, cache_key, result, _h5_meta_cache_size())
    return result


def h5_files_are_continuous(
    current_h5_file: str | Path,
    next_h5_file: str | Path,
    max_gap_seconds: float,
) -> bool:
    current_range = calculate_h5_time_range(current_h5_file)
    next_range = calculate_h5_time_range(next_h5_file)
    if current_range is None or next_range is None:
        return False

    current_start, current_end = current_range
    next_start, next_end = next_range
    if next_end <= current_start:
        return False

    gap_seconds = next_start - current_end
    if gap_seconds < -float(max_gap_seconds):
        return False
    return gap_seconds <= float(max_gap_seconds)


def calculate_apa_total_count(
    session: Any,
    start_timestamp: float | None = None,
    end_timestamp: float | None = None,
) -> int:
    if session is None:
        return 0

    required_signals = [APA_CPI_TYPE_SIGNAL, APA_CPI_STATUS_SIGNAL]
    if session.missing_signals(required_signals):
        return 0

    try:
        if start_timestamp is None or end_timestamp is None:
            start_timestamp, end_timestamp = session.common_time_range(required_signals)
        axis, parking_status_raw = session.get_signal_slice(APA_CPI_STATUS_SIGNAL, start_timestamp, end_timestamp)
        if len(axis) < 2:
            return 0

        apa_type_values = session.sample_signal_to_axis(
            APA_CPI_TYPE_SIGNAL,
            axis,
            method="nearest",
            start_timestamp=start_timestamp,
            end_timestamp=end_timestamp,
        )

        total_count = 0
        for index in range(1, len(axis)):
            prev_status = normalize_numeric_value(parking_status_raw[index - 1])
            curr_status = normalize_numeric_value(parking_status_raw[index])
            prev_type = normalize_numeric_value(apa_type_values[index - 1])
            if prev_type == 1 and prev_status == 5 and curr_status != 5:
                total_count += 1
        return total_count
    except Exception:
        return 0


def calculate_folder_mileage_km(primary_h5_file: str | Path | None) -> float:
    if not primary_h5_file:
        return 0.0

    file_key = _h5_file_cache_key(primary_h5_file)
    if file_key is not None and file_key in _H5_MILEAGE_CACHE:
        return _H5_MILEAGE_CACHE[file_key]

    mileage_km = 0.0
    try:
        import h5py

        with h5py.File(primary_h5_file, "r") as h5_file:
            root_names = list(h5_file.keys())
            if not root_names:
                mileage_km = 0.0
            else:
                channel = h5_file[root_names[0]].get("Channel1")
                if channel is None or ODOMETER_DATASET_NAME not in channel:
                    mileage_km = 0.0
                else:
                    dataset = channel[ODOMETER_DATASET_NAME]
                    dtype_names = dataset.dtype.names or ()
                    if ODOMETER_FIELD_NAME not in dtype_names:
                        mileage_km = 0.0
                    else:
                        odometer_values = dataset[ODOMETER_FIELD_NAME]
                        if len(odometer_values) == 0:
                            mileage_km = 0.0
                        else:
                            if ODOMETER_VALID_FIELD_NAME in dtype_names:
                                valid_mask = dataset[ODOMETER_VALID_FIELD_NAME] == 0
                                if valid_mask.any():
                                    odometer_values = odometer_values[valid_mask]

                            if len(odometer_values) == 0:
                                mileage_km = 0.0
                            else:
                                mileage_km = max(float(odometer_values[-1] - odometer_values[0]), 0.0)
    except Exception:
        mileage_km = 0.0

    if file_key is not None:
        _remember_limited(_H5_MILEAGE_CACHE, file_key, mileage_km, _h5_meta_cache_size())
    return mileage_km


def calculate_folder_start_timestamp(session: Any, primary_h5_file: str | Path | None = None) -> float | None:
    try:
        if session is not None:
            store = session.get_store()
            start_values: list[float] = []
            for signal in getattr(store, "signals", {}).values():
                timestamps = signal.get("timestamps") if isinstance(signal, dict) else None
                if timestamps:
                    start_values.append(float(timestamps[0]))
            if start_values:
                return min(start_values)
    except Exception:
        pass

    return calculate_h5_start_timestamp(primary_h5_file)


def calculate_h5_start_timestamp(primary_h5_file: str | Path | None) -> float | None:
    h5_range = calculate_h5_time_range(primary_h5_file)
    return h5_range[0] if h5_range is not None else None


def calculate_h5_time_range(primary_h5_file: str | Path | None) -> tuple[float, float] | None:
    if not primary_h5_file:
        return None

    file_key = _h5_file_cache_key(primary_h5_file)
    if file_key is not None and file_key in _H5_TIME_RANGE_CACHE:
        return _H5_TIME_RANGE_CACHE[file_key]

    h5_range: tuple[float, float] | None = None
    try:
        import h5py

        with h5py.File(primary_h5_file, "r") as h5_file:
            root_names = list(h5_file.keys())
            if not root_names:
                h5_range = None
            else:
                channel = h5_file[root_names[0]].get("Channel1")
                if channel is None:
                    h5_range = None
                else:
                    start_values: list[float] = []
                    end_values: list[float] = []
                    for message_name in channel.keys():
                        dataset = channel[message_name]
                        dtype_names = dataset.dtype.names or ()
                        if "timestamp" not in dtype_names or len(dataset) == 0:
                            continue
                        timestamps = dataset["timestamp"][:]
                        if len(timestamps) == 0:
                            continue
                        start_values.append(float(min(timestamps)))
                        end_values.append(float(max(timestamps)))

                    if start_values and end_values:
                        h5_range = (min(start_values), max(end_values))
    except Exception:
        h5_range = None

    if file_key is not None:
        _remember_limited(_H5_TIME_RANGE_CACHE, file_key, h5_range, _h5_meta_cache_size())
    return h5_range


def topic_scene_fingerprint(module: object) -> str | None:
    module_key = id(module)
    cached = _TOPIC_SCENE_FINGERPRINT_CACHE.get(module_key)
    if cached is not None and cached[0] is module:
        return cached[1]

    extractor = getattr(module, "_find_risk_segments", None)
    if not callable(extractor):
        _TOPIC_SCENE_FINGERPRINT_CACHE[module_key] = (module, None)
        return None

    hasher = hashlib.sha256()
    for function_name in TOPIC_SCENE_FINGERPRINT_FUNCTIONS:
        function = getattr(module, function_name, None)
        if not callable(function):
            continue
        try:
            function_tree = ast.parse(inspect.getsource(function))
        except (OSError, TypeError, SyntaxError):
            _TOPIC_SCENE_FINGERPRINT_CACHE[module_key] = (module, None)
            return None
        hasher.update(function_name.encode("utf-8"))
        hasher.update(ast.dump(function_tree, include_attributes=False).encode("utf-8"))

    fingerprint = hasher.hexdigest()
    _TOPIC_SCENE_FINGERPRINT_CACHE[module_key] = (module, fingerprint)
    return fingerprint


def clone_topic_scene_segments(segments: Any) -> Any:
    if not isinstance(segments, list):
        return segments
    return [segment[:] if isinstance(segment, list) else segment for segment in segments]


def run_topic_metric_with_shared_scenes(
    metric: MetricModule,
    frames: list[dict[str, Any]],
    scene_cache: dict[tuple[str, int], tuple[list[dict[str, Any]], Any]],
) -> Any:
    extractor = getattr(metric.module, "_find_risk_segments", None)
    fingerprint = topic_scene_fingerprint(metric.module)
    if not callable(extractor) or fingerprint is None:
        return metric.runner(frames)

    def cached_extractor(
        comprehensive_frames: list[dict[str, Any]],
        *args: Any,
        **kwargs: Any,
    ) -> Any:
        cache_key = (fingerprint, id(comprehensive_frames))
        cached = scene_cache.get(cache_key)
        if cached is not None and cached[0] is comprehensive_frames:
            return clone_topic_scene_segments(cached[1])

        segments = extractor(comprehensive_frames, *args, **kwargs)
        scene_cache[cache_key] = (comprehensive_frames, segments)
        return clone_topic_scene_segments(segments)

    setattr(metric.module, "_find_risk_segments", cached_extractor)
    try:
        return metric.runner(frames)
    finally:
        setattr(metric.module, "_find_risk_segments", extractor)


def process_topic_folder(
    pipeline: PipelineConfig,
    metrics: list[MetricModule],
    folder_path: str,
    folder_index: int,
    total_folders: int,
    run_timestamp: str,
    next_folder_path: str | None,
    run_state: RunState | None = None,
    replace_existing_results: bool = False,
) -> None:
    print(f"[{pipeline.name}] [{folder_index}/{total_folders}] 开始 Topic 指标")
    pending_metrics = filter_pending_metrics(run_state, pipeline, folder_path, metrics)
    if not pending_metrics:
        print(f"[{pipeline.name}] 本文件夹 Topic 指标均已完成，跳过")
        return

    try:
        curr_frames = load_topic_frames_cached(folder_path)
    except Exception as exc:
        print(f"[{pipeline.name}] Topic 预处理失败：{exc}")
        for metric in pending_metrics:
            csv_path = save_metric_result(
                pipeline=pipeline,
                metric=metric,
                folder_path=folder_path,
                result=False,
                details=[f"Topic 预处理失败: {exc}"],
                statistics={},
                mileage_km=0.0,
                run_timestamp=run_timestamp,
                replace_existing=replace_existing_results,
            )
            if run_state is not None:
                run_state.finish_task(
                    pipeline=pipeline,
                    folder_path=folder_path,
                    metric=metric,
                    status=TASK_STATUS_FAILED,
                    csv_path=csv_path,
                    result=False,
                    details=[f"Topic 预处理失败: {exc}"],
                    error=str(exc),
                )
        return

    if not curr_frames and pipeline.skip_empty_topic_frames:
        print(f"[{pipeline.name}] 跳过：Topic 预处理未产生综合帧")
        for metric in pending_metrics:
            if run_state is not None:
                run_state.finish_task(
                    pipeline=pipeline,
                    folder_path=folder_path,
                    metric=metric,
                    status=TASK_STATUS_SKIPPED_EMPTY_TOPIC_FRAMES,
                    details=["Topic 预处理未产生综合帧"],
                )
        return

    cross_done_names = normalize_cross_done_names(pipeline.cross_done_metric_names)
    next_frames: list[dict[str, Any]] | None = None
    cross_done_frames: list[dict[str, Any]] | None = None
    cross_done_frames_ready = False
    scene_cache: dict[tuple[str, int], tuple[list[dict[str, Any]], Any]] = {}
    for metric in pending_metrics:
        print(f"[{pipeline.name}]  └── 执行：{metric.file_name}")
        if run_state is not None:
            run_state.start_task(pipeline, folder_path, metric)
        frames_for_metric = curr_frames

        if pipeline.enable_cross_done_continuation and metric_matches_cross_done(metric, cross_done_names):
            if not cross_done_frames_ready:
                if next_folder_path:
                    try:
                        next_frames = load_topic_frames_cached(next_folder_path)
                        cross_done_frames = build_two_done_frames(
                            curr_frames,
                            folder_path,
                            next_frames,
                            next_folder_path,
                            pipeline.cross_done_context_seconds,
                        )
                    except Exception as exc:
                        print(f"[{pipeline.name}]  ⚠️ 下一个 done 预处理失败，按单段执行：{exc}")
                        cross_done_frames = build_two_done_frames(
                            curr_frames,
                            folder_path,
                            None,
                            None,
                            pipeline.cross_done_context_seconds,
                        )
                else:
                    cross_done_frames = build_two_done_frames(
                        curr_frames,
                        folder_path,
                        None,
                        None,
                        pipeline.cross_done_context_seconds,
                    )
                cross_done_frames_ready = True
            frames_for_metric = cross_done_frames if cross_done_frames is not None else curr_frames

        try:
            metric_result = run_topic_metric_with_shared_scenes(metric, frames_for_metric, scene_cache)
            check_result, check_details, check_statistics = normalize_metric_result(metric_result)
            csv_path = save_metric_result(
                pipeline=pipeline,
                metric=metric,
                folder_path=folder_path,
                result=check_result,
                details=check_details,
                statistics=check_statistics,
                mileage_km=0.0,
                run_timestamp=run_timestamp,
                replace_existing=replace_existing_results,
            )
            if run_state is not None:
                run_state.finish_task(
                    pipeline=pipeline,
                    folder_path=folder_path,
                    metric=metric,
                    status=TASK_STATUS_SUCCESS,
                    csv_path=csv_path,
                    result=check_result,
                    details=check_details,
                )
            print(f"[{pipeline.name}]  ✅ {metric.file_name} 完成")
        except Exception as exc:
            csv_path = save_metric_result(
                pipeline=pipeline,
                metric=metric,
                folder_path=folder_path,
                result=False,
                details=[f"执行失败: {exc}"],
                statistics={},
                mileage_km=0.0,
                run_timestamp=run_timestamp,
                replace_existing=replace_existing_results,
            )
            if run_state is not None:
                run_state.finish_task(
                    pipeline=pipeline,
                    folder_path=folder_path,
                    metric=metric,
                    status=TASK_STATUS_FAILED,
                    csv_path=csv_path,
                    result=False,
                    details=[f"执行失败: {exc}"],
                    error=str(exc),
                )
            print(f"[{pipeline.name}]  ❌ {metric.file_name} 执行失败：{exc}")


def load_topic_frames_cached(folder_path: str) -> list[dict[str, Any]]:
    cache_key = str(Path(folder_path).resolve())
    cached = _TOPIC_FRAMES_CACHE.get(cache_key)
    if cached is not None:
        _TOPIC_FRAMES_CACHE.pop(cache_key)
        _TOPIC_FRAMES_CACHE[cache_key] = cached
        return cached

    import longitudinal_preprocessor

    frames = longitudinal_preprocessor.process_comprehensive_frames(folder_path)
    _remember_limited(_TOPIC_FRAMES_CACHE, cache_key, frames, _topic_frames_cache_size())
    return frames


def normalize_cross_done_names(names: set[str]) -> set[str]:
    normalized: set[str] = set()
    for name in names:
        clean_name = str(name).strip()
        if not clean_name:
            continue
        normalized.add(clean_name)
        normalized.add(Path(clean_name).stem)
        if not clean_name.endswith(".py"):
            normalized.add(f"{clean_name}.py")
    return normalized


def metric_matches_cross_done(metric: MetricModule, cross_done_names: set[str]) -> bool:
    return "*" in cross_done_names or metric.name in cross_done_names or metric.file_name in cross_done_names


def tag_sequence_frames(
    frames: list[dict[str, Any]],
    folder_path: str,
    source_index: int,
    primary_frame_count: int,
    peer_folder_path: str | None = None,
) -> list[dict[str, Any]]:
    tagged_frames: list[dict[str, Any]] = []
    for frame_index, frame in enumerate(frames):
        tagged_frame = dict(frame)
        tagged_frame["_sequence_meta"] = {
            "source_index": source_index,
            "is_primary": source_index == 0,
            "source_folder_path": folder_path,
            "source_folder_name": os.path.basename(folder_path),
            "frame_index_in_source": frame_index,
            "primary_frame_count": primary_frame_count,
            "peer_folder_path": peer_folder_path,
            "peer_folder_name": os.path.basename(peer_folder_path) if peer_folder_path else "",
        }
        tagged_frames.append(tagged_frame)
    return tagged_frames


def build_two_done_frames(
    current_frames: list[dict[str, Any]],
    current_folder: str,
    next_frames: list[dict[str, Any]] | None,
    next_folder: str | None,
    context_seconds: float = 10.0,
) -> list[dict[str, Any]]:
    primary_count = len(current_frames)
    combined_frames = tag_sequence_frames(
        current_frames,
        current_folder,
        source_index=0,
        primary_frame_count=primary_count,
        peer_folder_path=next_folder,
    )
    if next_frames and next_folder:
        next_frames = limit_next_frames_for_context(current_frames, next_frames, context_seconds)
        combined_frames.extend(
            tag_sequence_frames(
                next_frames,
                next_folder,
                source_index=1,
                primary_frame_count=primary_count,
                peer_folder_path=current_folder,
            )
        )
    return combined_frames


def limit_next_frames_for_context(
    current_frames: list[dict[str, Any]],
    next_frames: list[dict[str, Any]],
    context_seconds: float,
) -> list[dict[str, Any]]:
    current_timestamps = [
        timestamp
        for timestamp in (extract_frame_timestamp_ms(frame) for frame in current_frames)
        if timestamp is not None
    ]
    if not current_timestamps:
        return next_frames

    current_end_ms = max(current_timestamps)
    cutoff_ms = current_end_ms + max(float(context_seconds), 0.0) * 1000.0
    limited_frames = []
    for frame in next_frames:
        timestamp_ms = extract_frame_timestamp_ms(frame)
        if timestamp_ms is None:
            continue
        if current_end_ms <= timestamp_ms <= cutoff_ms:
            limited_frames.append(frame)
    return limited_frames


def extract_frame_timestamp_ms(frame: dict[str, Any]) -> float | None:
    timestamp = frame.get("unp_timestamp_ms")
    if timestamp is None:
        timestamp = frame.get("timestamp_ms")
    try:
        return None if timestamp is None else float(timestamp)
    except Exception:
        return None


def normalize_metric_result(result: Any) -> tuple[bool, list[Any], dict[str, Any]]:
    if not validate_function_return(result):
        raise ValueError("指标返回值结构错误，必须为(bool, list[, dict])")

    if len(result) == 3:
        check_result, check_details, check_statistics = result
    else:
        check_result, check_details = result
        check_statistics = {}

    return bool(check_result), list(check_details), dict(check_statistics or {})


def validate_function_return(result: Any) -> bool:
    if not isinstance(result, tuple):
        return False
    if len(result) not in (2, 3):
        return False
    if not isinstance(result[0], bool):
        return False
    if not isinstance(result[1], list):
        return False
    if len(result) == 3 and not isinstance(result[2], dict):
        return False
    return True


def save_metric_result(
    pipeline: PipelineConfig,
    metric: MetricModule,
    folder_path: str,
    result: bool,
    details: list[Any],
    statistics: dict[str, Any],
    mileage_km: float,
    run_timestamp: str,
    segment_start_timestamp: float | None = None,
    replace_existing: bool = False,
) -> Path:
    output_path = pipeline.output_dir / f"{metric.output_prefix}_{run_timestamp}.csv"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    fieldnames = ["指标来源", "文件夹路径", "检查结果", "符合要求的内容", "总Case", "Bad Case", "严重度得分", "段里程(km)"]
    need_header = not output_path.exists() or output_path.stat().st_size == 0
    row = {
        "指标来源": pipeline.name,
        "文件夹路径": folder_path,
        "检查结果": "是" if result else "否",
        "符合要求的内容": "; ".join(
            stringify_detail(item, segment_start_timestamp=segment_start_timestamp) for item in details
        ),
        "总Case": _safe_int(statistics.get("All Case Count", 0)),
        "Bad Case": _safe_int(statistics.get("Bad Case", 0)),
        "严重度得分": _safe_float(statistics.get("Severity Score", 0.0)),
        "段里程(km)": round(_safe_float(mileage_km), 6),
    }

    if replace_existing and output_path.exists():
        upsert_metric_csv_row(output_path, fieldnames, row)
    else:
        with output_path.open("a", newline="", encoding="utf-8-sig") as file:
            writer = csv.DictWriter(file, fieldnames=fieldnames)
            if need_header:
                writer.writeheader()
            writer.writerow(row)

    return output_path


def upsert_metric_csv_row(output_path: Path, fieldnames: list[str], new_row: dict[str, Any]) -> None:
    existing_rows: list[dict[str, str]] = []
    with output_path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        for row in reader:
            existing_rows.append({str(key): str(value) for key, value in row.items()})

    target_folder = normalize_state_path(str(new_row.get("文件夹路径", "")))
    output_rows: list[dict[str, Any]] = []
    inserted = False
    for row in existing_rows:
        row_folder = normalize_state_path(row.get("文件夹路径", ""))
        if row_folder == target_folder:
            if not inserted:
                output_rows.append(new_row)
                inserted = True
            continue
        output_rows.append({field: row.get(field, "") for field in fieldnames})

    if not inserted:
        output_rows.append(new_row)

    temp_path = output_path.with_name(f"{output_path.name}.tmp")
    with temp_path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        for row in output_rows:
            writer.writerow({field: row.get(field, "") for field in fieldnames})
    temp_path.replace(output_path)


def stringify_detail(item: Any, segment_start_timestamp: float | None = None) -> str:
    if isinstance(item, (dict, list)):
        try:
            return json.dumps(
                convert_event_times_in_detail(item, segment_start_timestamp),
                ensure_ascii=False,
                sort_keys=True,
            )
        except Exception:
            return str(item)

    detail_text = str(item)
    if segment_start_timestamp is None:
        return detail_text
    return convert_event_times_in_text(detail_text, segment_start_timestamp)


def convert_event_times_in_detail(item: Any, segment_start_timestamp: float | None) -> Any:
    if segment_start_timestamp is None:
        return item

    if isinstance(item, dict):
        converted: dict[Any, Any] = {}
        for key, value in item.items():
            if str(key).lower() in EVENT_TIME_DETAIL_KEYS:
                relative_value = calculate_relative_event_time(value, segment_start_timestamp)
                converted[key] = round(relative_value, 3) if relative_value is not None else value
            else:
                converted[key] = convert_event_times_in_detail(value, segment_start_timestamp)
        return converted

    if isinstance(item, list):
        return [convert_event_times_in_detail(value, segment_start_timestamp) for value in item]

    return item


def convert_event_times_in_text(detail_text: str, segment_start_timestamp: float) -> str:
    converted_text = EVENT_TIME_TEXT_PATTERN.sub(
        lambda match: replace_event_time_match(match, segment_start_timestamp),
        detail_text,
    )
    return EVENT_BLOCK_PATTERN.sub(
        lambda match: EVENT_BLOCK_START_END_PATTERN.sub(
            lambda inner_match: replace_event_time_match(inner_match, segment_start_timestamp),
            match.group(0),
        ),
        converted_text,
    )


def replace_event_time_match(match: re.Match[str], segment_start_timestamp: float) -> str:
    relative_value = calculate_relative_event_time(match.group("value"), segment_start_timestamp)
    if relative_value is None:
        return match.group(0)
    return f"{match.group('label')}{relative_value:.3f}"


def calculate_relative_event_time(value: Any, segment_start_timestamp: float) -> float | None:
    try:
        event_timestamp = float(value)
        start_timestamp = float(segment_start_timestamp)
    except Exception:
        return None

    if event_timestamp + 1e-6 < start_timestamp:
        return None

    relative_timestamp = event_timestamp - start_timestamp
    if abs(relative_timestamp) < 0.0005:
        relative_timestamp = 0.0
    return relative_timestamp


def generate_unified_report(
    output_dir: Path,
    report_prefix: str,
    run_timestamp: str,
    metric_category_map: dict[str, str],
) -> Path | None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font

    output_dir.mkdir(parents=True, exist_ok=True)
    report_stem = report_prefix[:-5] if report_prefix.lower().endswith(".xlsx") else report_prefix
    output_file = output_dir / f"{report_stem}_{run_timestamp}.xlsx"

    statistics_result: list[dict[str, Any]] = []
    bad_case_folders_cache: dict[str, list[str]] = {}
    bad_case_severity_score_cache: dict[str, list[float]] = {}
    bad_case_content_cache: dict[str, list[str]] = {}

    current_run_suffix = f"_{run_timestamp}.csv"
    for file_path in sorted(output_dir.glob("*.csv")):
        if not file_path.name.endswith(current_run_suffix):
            continue

        prefix = extract_csv_prefix(file_path.name)
        csv_data = read_csv_file(file_path)
        if not csv_data:
            continue

        case_sums = calculate_case_sums(csv_data)
        pipeline_name, metric_name = split_output_prefix(prefix, csv_data)
        metric_category = metric_category_map.get(prefix, METRIC_CATEGORY_CPI)
        total_case_sum = case_sums["total_case_sum"]
        bad_case_sum = case_sums["bad_case_sum"]
        total_mileage_km = case_sums["total_mileage_km"]

        bad_case_folders_cache[prefix] = case_sums["bad_case_folders"]
        bad_case_severity_score_cache[prefix] = case_sums["severity_score_folders"]
        bad_case_content_cache[prefix] = case_sums["detail_content_folders"]

        statistics_result.append(
            {
                "指标来源": pipeline_name,
                "指标名称": metric_name,
                "指标类型": metric_category,
                "总里程(km)": round(total_mileage_km, 6),
                "总次数": total_case_sum,
                "触发次数": bad_case_sum,
                "output_prefix": prefix,
            }
        )

    if not statistics_result:
        print("未找到可统计的CSV结果文件，跳过统计报告生成")
        return None

    statistics_result.sort(key=lambda row: (str(row["指标来源"]), str(row["指标名称"])))

    workbook = Workbook()
    ws_main = workbook.active
    ws_main.title = "所有指标统计汇总"

    header_font = Font(bold=True, size=11)
    header_alignment = Alignment(horizontal="center", vertical="center")
    green_font = Font(color="008000")
    red_font = Font(color="FF0000")

    headers = ["指标来源", "指标名称", "指标类型", "总里程(km)", "总次数", "触发次数", "MPI", "CPI"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws_main.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = header_alignment

    for row_idx, row_data in enumerate(statistics_result, 2):
        metric_category = row_data["指标类型"]
        total_mileage_km = _safe_float(row_data["总里程(km)"])
        total_case = _safe_int(row_data["总次数"])
        trigger_count = _safe_int(row_data["触发次数"])

        for col_idx, key in enumerate(["指标来源", "指标名称", "指标类型", "总里程(km)"], 1):
            ws_main.cell(row=row_idx, column=col_idx, value=row_data[key]).alignment = header_alignment

        total_cell = ws_main.cell(row=row_idx, column=5, value=total_case)
        total_cell.font = green_font
        total_cell.alignment = header_alignment
        trigger_cell = ws_main.cell(row=row_idx, column=6, value=trigger_count)
        trigger_cell.font = red_font
        trigger_cell.alignment = header_alignment

        mpi_text = ""
        cpi_text = ""
        if metric_category == METRIC_CATEGORY_MPI:
            if trigger_count == 0:
                mileage_text = f"{round(total_mileage_km, 3):g}"
                mpi_text = f">{mileage_text}({mileage_text}/0)"
            else:
                mpi_value = round(total_mileage_km / trigger_count, 1)
                mpi_text = f"{mpi_value}({round(total_mileage_km, 3)}/{trigger_count})"
        else:
            if trigger_count == 0:
                cpi_text = f">{total_case}({total_case}/0)"
            else:
                cpi_value = round(total_case / trigger_count, 3)
                cpi_text = f"{cpi_value}({total_case}/{trigger_count})"

        mpi_cell = ws_main.cell(row=row_idx, column=7, value=mpi_text)
        mpi_cell.alignment = header_alignment
        mpi_cell.comment = Comment(
            text="MPI = 总里程(km) / 触发次数\n触发次数沿用各指标CSV中的 Bad Case 统计口径",
            author="统一统计脚本",
        )
        cpi_cell = ws_main.cell(row=row_idx, column=8, value=cpi_text)
        cpi_cell.alignment = header_alignment
        cpi_cell.comment = Comment(
            text="CPI = 总次数 / 触发次数\n总次数沿用各指标CSV中的 All Case Count 统计口径",
            author="统一统计脚本",
        )

    for col in ws_main.columns:
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
        ws_main.column_dimensions[col[0].column_letter].width = min(max_length + 4, 70)

    used_sheet_names = set(workbook.sheetnames)
    for row_data in statistics_result:
        output_prefix = row_data["output_prefix"]
        sheet_name = unique_sheet_name(output_prefix, used_sheet_names)
        ws = workbook.create_sheet(title=sheet_name)
        used_sheet_names.add(sheet_name)

        headers = ["Bad_Case数据列表", "问题严重程度得分", "问题内容", "问题分析", "分析截图"]
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font if col_idx <= 3 else red_font
            cell.alignment = header_alignment

        folders = bad_case_folders_cache.get(output_prefix, [])
        scores = bad_case_severity_score_cache.get(output_prefix, [])
        contents = bad_case_content_cache.get(output_prefix, [])
        if folders and scores and contents:
            for row_idx, (bad_folder, score, content) in enumerate(zip(folders, scores, contents), 2):
                ws.cell(row=row_idx, column=1, value=bad_folder)
                ws.cell(row=row_idx, column=2, value=score)
                ws.cell(row=row_idx, column=3, value=content)
        else:
            ws.cell(row=2, column=1, value="无Bad_Case")
            ws.cell(row=2, column=2, value="无Bad_Case")
            ws.cell(row=2, column=3, value="无Bad_Case")

        ws.column_dimensions["A"].width = 80
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 100
        ws.column_dimensions["D"].width = 60
        ws.column_dimensions["E"].width = 60

    workbook.save(output_file)
    print(f"统计报告已生成：{output_file}")
    return output_file


def generate_bad_case_extract_report(
    output_dir: Path,
    run_timestamp: str,
    extract_config: BadCaseExtractConfig,
) -> Path | None:
    if not extract_config.enabled:
        print("Bad Case 提取已禁用，跳过 Bad Case 报告生成")
        return None

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    output_dir.mkdir(parents=True, exist_ok=True)
    report_stem = extract_config.output_file[:-5] if extract_config.output_file.lower().endswith(".xlsx") else extract_config.output_file
    output_file = output_dir / f"{report_stem}_{run_timestamp}.xlsx"

    report_rows: list[dict[str, Any]] = []
    current_run_suffix = f"_{run_timestamp}.csv"
    for file_path in sorted(output_dir.glob("*.csv")):
        if not file_path.name.endswith(current_run_suffix):
            continue

        output_prefix = extract_csv_prefix(file_path.name)
        csv_data = read_csv_file(file_path)
        if not csv_data:
            continue

        pipeline_name, metric_name = split_output_prefix(output_prefix, csv_data)
        rule = resolve_bad_case_extract_rule(extract_config, pipeline_name, metric_name, output_prefix)
        if not rule.enabled:
            continue

        candidates: list[dict[str, Any]] = []
        for row in csv_data:
            if _safe_int(row.get("Bad Case", 0)) <= 0:
                continue

            severity_score = _safe_float(row.get("严重度得分", 0.0))
            if rule.min_severity is not None and severity_score < rule.min_severity:
                continue

            folder_path = row.get("文件夹路径", "")
            bad_case_date, bad_case_time = extract_bad_case_datetime(folder_path)
            candidates.append(
                {
                    "pipeline_name": pipeline_name,
                    "metric_name": metric_name,
                    "output_prefix": output_prefix,
                    "folder_path": folder_path,
                    "severity_score": severity_score,
                    "bad_case_date": bad_case_date,
                    "bad_case_time": bad_case_time,
                    "detail_content": row.get("符合要求的内容", ""),
                    "rule_description": describe_bad_case_rule(rule),
                }
            )

        selected = select_bad_case_records(candidates, rule)
        for rank, item in enumerate(selected, start=1):
            item["rank"] = rank
            report_rows.append(item)

    report_rows.sort(key=lambda row: (str(row["pipeline_name"]), str(row["metric_name"]), int(row["rank"])))

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Bad Case提取报告"

    headers = [
        "指标来源",
        "指标名称",
        "指标Key",
        "序号",
        "严重度得分",
        "日期",
        "时间",
        "Bad Case文件路径",
        "问题内容",
        "提取规则",
    ]
    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F81BD")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)
    center_alignment = Alignment(horizontal="center", vertical="center")

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    if report_rows:
        for row_idx, row_data in enumerate(report_rows, start=2):
            values = [
                row_data["pipeline_name"],
                row_data["metric_name"],
                row_data["output_prefix"],
                row_data["rank"],
                row_data["severity_score"],
                row_data["bad_case_date"],
                row_data["bad_case_time"],
                row_data["folder_path"],
                row_data["detail_content"],
                row_data["rule_description"],
            ]
            for col_idx, value in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = center_alignment if col_idx in {1, 2, 3, 4, 5, 6, 7} else body_alignment
    else:
        ws.cell(row=2, column=1, value="无符合 bad_case_extract.json 规则的 Bad Case")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))

    widths = {
        "A": 12,
        "B": 32,
        "C": 38,
        "D": 10,
        "E": 14,
        "F": 12,
        "G": 10,
        "H": 100,
        "I": 120,
        "J": 38,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    workbook.save(output_file)
    print(f"Bad Case 提取报告已生成：{output_file}")
    return output_file


def resolve_bad_case_extract_rule(
    extract_config: BadCaseExtractConfig,
    pipeline_name: str,
    metric_name: str,
    output_prefix: str,
) -> BadCaseExtractRule:
    rule = extract_config.global_rule
    if pipeline_name in extract_config.pipeline_rules:
        rule = merge_bad_case_extract_rules(rule, extract_config.pipeline_rules[pipeline_name])

    for key in (output_prefix, metric_name):
        if key in extract_config.metric_rules:
            rule = merge_bad_case_extract_rules(rule, extract_config.metric_rules[key])
    return rule


def merge_bad_case_extract_rules(base_rule: BadCaseExtractRule, override_rule: BadCaseExtractRule) -> BadCaseExtractRule:
    return BadCaseExtractRule(
        enabled=override_rule.enabled,
        min_severity=override_rule.min_severity,
        top_n=override_rule.top_n,
        max_cases=override_rule.max_cases,
        sort_by=override_rule.sort_by or base_rule.sort_by,
    )


def select_bad_case_records(candidates: list[dict[str, Any]], rule: BadCaseExtractRule) -> list[dict[str, Any]]:
    if rule.sort_by == "severity_asc":
        candidates = sorted(candidates, key=lambda row: (_safe_float(row["severity_score"]), str(row["folder_path"])))
    elif rule.sort_by == "path_asc":
        candidates = sorted(candidates, key=lambda row: str(row["folder_path"]))
    else:
        candidates = sorted(candidates, key=lambda row: (-_safe_float(row["severity_score"]), str(row["folder_path"])))

    if rule.top_n is not None:
        candidates = candidates[: rule.top_n]
    if rule.max_cases is not None:
        candidates = candidates[: rule.max_cases]
    return candidates


def describe_bad_case_rule(rule: BadCaseExtractRule) -> str:
    parts = [f"sort_by={rule.sort_by}"]
    if rule.min_severity is not None:
        parts.append(f"min_severity>={rule.min_severity:g}")
    if rule.top_n is not None:
        parts.append(f"top_n={rule.top_n}")
    if rule.max_cases is not None:
        parts.append(f"max_cases={rule.max_cases}")
    return "; ".join(parts)


def extract_bad_case_datetime(path_text: str) -> tuple[str, str]:
    path_parts = [part for part in re.split(r"[\\/]+", str(path_text)) if part]
    for part in reversed(path_parts):
        parsed = parse_datetime_from_name(part)
        if parsed is not None:
            return parsed
    return "", ""


def parse_datetime_from_name(name: str) -> tuple[str, str] | None:
    name_text = re.sub(r"\.[A-Za-z0-9]+$", "", str(name))
    datetime_patterns = [
        r"(?<!\d)(20\d{2})[-_]?(\d{2})[-_]?(\d{2})[-_ T]?(\d{2})[-_]?(\d{2})(?:[-_]?(\d{2}))?(?!\d)",
        r"(?<!\d)(20\d{2})[-_](\d{2})[-_](\d{2})[-_ T](\d{2})[-_](\d{2})(?:[-_](\d{2}))?(?!\d)",
    ]
    for pattern in datetime_patterns:
        match = re.search(pattern, name_text)
        if match:
            formatted = format_bad_case_datetime(
                year=match.group(1),
                month=match.group(2),
                day=match.group(3),
                hour=match.group(4),
                minute=match.group(5),
            )
            if formatted is not None:
                return formatted

    date_patterns = [
        r"(?<!\d)(20\d{2})[-_]?(\d{2})[-_]?(\d{2})(?!\d)",
        r"(?<!\d)(20\d{2})[-_](\d{2})[-_](\d{2})(?!\d)",
    ]
    for pattern in date_patterns:
        match = re.search(pattern, name_text)
        if match:
            formatted = format_bad_case_datetime(
                year=match.group(1),
                month=match.group(2),
                day=match.group(3),
            )
            if formatted is not None:
                return formatted
    return None


def format_bad_case_datetime(
    year: str,
    month: str,
    day: str,
    hour: str | None = None,
    minute: str | None = None,
) -> tuple[str, str] | None:
    try:
        if hour is not None and minute is not None:
            datetime.strptime(f"{year}{month}{day}{hour}{minute}", "%Y%m%d%H%M")
            return f"{month}/{day}", f"{hour}:{minute}"
        datetime.strptime(f"{year}{month}{day}", "%Y%m%d")
        return f"{month}/{day}", ""
    except ValueError:
        return None


def unique_sheet_name(name: str, used_sheet_names: set[str]) -> str:
    sanitized = re.sub(r"[\[\]\:\*\?\/\\]", "_", name)
    base = sanitized[:31] if len(sanitized) > 31 else sanitized
    if base not in used_sheet_names:
        return base

    index = 1
    while True:
        suffix = f"_{index}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        if candidate not in used_sheet_names:
            return candidate
        index += 1


def extract_csv_prefix(file_name: str) -> str:
    stem = Path(file_name).stem
    match = re.match(r"^(.*)_\d{4}-\d{2}-\d{2}_\d{2}-\d{2}-\d{2}$", stem)
    if match:
        return match.group(1)
    return stem


def split_output_prefix(prefix: str, rows: list[dict[str, str]]) -> tuple[str, str]:
    if "__" in prefix:
        pipeline_name, metric_name = prefix.split("__", 1)
        return pipeline_name, metric_name
    pipeline_name = rows[0].get("指标来源", "") if rows else ""
    return pipeline_name, prefix


def read_csv_file(file_path: str | Path) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    with Path(file_path).open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        for row in reader:
            rows.append({str(key): str(value) for key, value in row.items()})
    return rows


def calculate_case_sums(csv_data: list[dict[str, str]]) -> dict[str, Any]:
    total_case_sum = 0
    bad_case_sum = 0
    total_mileage_km = 0.0
    bad_case_items: list[tuple[float, str, str]] = []

    for row in csv_data:
        total_case_sum += _safe_int(row.get("总Case", "0"))
        bad_case_value = _safe_int(row.get("Bad Case", "0"))
        bad_case_sum += bad_case_value
        total_mileage_km += _safe_float(row.get("段里程(km)", "0"))
        if bad_case_value > 0:
            folder_path = row.get("文件夹路径", "")
            severity_score = _safe_float(row.get("严重度得分", "0"))
            detail_content = row.get("符合要求的内容", "")
            bad_case_items.append((severity_score, folder_path, detail_content))

    bad_case_items.sort(key=lambda item: item[0], reverse=True)
    return {
        "total_case_sum": total_case_sum,
        "bad_case_sum": bad_case_sum,
        "total_mileage_km": total_mileage_km,
        "bad_case_folders": [item[1] for item in bad_case_items],
        "severity_score_folders": [item[0] for item in bad_case_items],
        "detail_content_folders": [item[2] for item in bad_case_items],
    }


def normalize_numeric_value(value: object) -> int | None:
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, bytes):
        value = value.decode(errors="ignore")
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip()
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _safe_int(value: object) -> int:
    try:
        return int(float(value))
    except Exception:
        return 0


def _safe_float(value: object) -> float:
    try:
        return float(value)
    except Exception:
        return 0.0


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run CAN and Topic data-mining metrics from one unified entrypoint.")
    parser.add_argument("--config", default="config.json", help="Path to unified config.json")
    parser.add_argument("--bad-case-config", default="bad_case_extract.json", help="Path to bad case extraction rules")
    parser.add_argument(
        "--resume",
        choices=[RESUME_MODE_ASK, RESUME_MODE_AUTO, RESUME_MODE_NEVER],
        default=RESUME_MODE_ASK,
        help="Run-state resume mode: ask prompts, auto resumes only when config/code/folder list match, never starts fresh",
    )
    parser.add_argument("--dry-run", action="store_true", help="Validate config and metric loading without scanning data")
    parser.add_argument("--list-metrics", action="store_true", help="Print loaded metrics before running")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    base_dir = Path(__file__).resolve().parent
    config_path = Path(args.config)
    if not config_path.is_absolute():
        config_path = base_dir / config_path
    bad_case_config_path = Path(args.bad_case_config)
    if not bad_case_config_path.is_absolute():
        bad_case_config_path = base_dir / bad_case_config_path

    try:
        config = load_config(config_path)
        return run_all(
            config,
            dry_run=bool(args.dry_run),
            list_metrics=bool(args.list_metrics),
            bad_case_config_path=bad_case_config_path,
            config_path=config_path,
            resume_mode=str(args.resume),
        )
    except Exception as exc:
        print(f"执行失败：{exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
