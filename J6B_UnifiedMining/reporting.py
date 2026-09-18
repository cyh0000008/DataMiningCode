from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, TextIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from metric_utils import MetricResult


CSV_HEADERS = [
    "指标类型",
    "文件路径",
    "检查结果",
    "符合要求的内容",
    "缺失字段",
    "缺失信号",
    "总Case",
    "Bad Case",
    "指标值",
    "严重程度得分",
]

SKIPPED_HEADERS = ["指标类型", "指标名称", "文件路径", "缺失字段", "缺失信号", "处理结果"]
SUMMARY_HEADERS = ["指标类型", "指标名称", "总Case之和", "Bad Case之和", "跳过数", "指标值"]
DETAIL_HEADERS = ["指标类型", "Bad_Case数据列表", "问题严重程度得分", "问题内容", "问题分析", "分析截图"]


class StreamingReportWriter:
    def __init__(self, output_dir: Path, report_name: str, metrics: Iterable[Any]) -> None:
        base_output_dir = Path(output_dir)
        base_output_dir.mkdir(parents=True, exist_ok=True)
        report_filename = Path(ensure_xlsx(report_name)).name
        self.output_dir = make_run_output_dir(base_output_dir, Path(report_filename).stem)
        self.report_path = self.output_dir / report_filename
        metric_list = list(metrics)
        self.metric_categories = {metric.name: metric.category for metric in metric_list}
        self.metric_order = sorted(
            [metric.name for metric in metric_list],
            key=lambda metric_name: metric_report_sort_key(
                metric_name,
                self.metric_categories.get(metric_name, "CPI"),
            ),
        )
        self.summaries: Dict[str, Dict[str, Any]] = {
            metric_name: {
                "category": self.metric_categories.get(metric_name, "CPI"),
                "total_cases": 0,
                "bad_cases": 0,
                "skipped_count": 0,
                "distance_by_path": {},
                "result_count": 0,
                "detail_count": 0,
            }
            for metric_name in self.metric_order
        }

        self.csv_files: Dict[str, TextIO] = {}
        self.csv_writers: Dict[str, csv.DictWriter] = {}
        self.json_files: Dict[str, TextIO] = {}
        self.json_first: Dict[str, bool] = {}

        for metric_name in self.metric_order:
            csv_fp = (self.output_dir / ("%s.csv" % metric_name)).open("w", newline="", encoding="utf-8-sig")
            csv_writer = csv.DictWriter(csv_fp, fieldnames=CSV_HEADERS)
            csv_writer.writeheader()
            self.csv_files[metric_name] = csv_fp
            self.csv_writers[metric_name] = csv_writer

            json_fp = (self.output_dir / ("%s_events.json" % metric_name)).open("w", encoding="utf-8")
            json_fp.write("[\n")
            self.json_files[metric_name] = json_fp
            self.json_first[metric_name] = True

        self.skipped_fp = (self.output_dir / "skipped_metrics.csv").open("w", newline="", encoding="utf-8-sig")
        self.skipped_writer = csv.DictWriter(self.skipped_fp, fieldnames=SKIPPED_HEADERS)
        self.skipped_writer.writeheader()

        self.workbook = Workbook(write_only=True)
        self.summary_sheet = self.workbook.create_sheet("所有指标统计汇总")
        self.summary_sheet.append(SUMMARY_HEADERS)
        self.metric_sheets: Dict[str, Any] = {}
        for metric_name in self.metric_order:
            category = self.metric_categories.get(metric_name, "CPI")
            sheet = self.workbook.create_sheet(safe_sheet_name("%s_%s" % (category, metric_name), self.workbook.sheetnames))
            sheet.append(DETAIL_HEADERS)
            self.metric_sheets[metric_name] = sheet
        self.skipped_sheet = self.workbook.create_sheet(safe_sheet_name("跳过指标", self.workbook.sheetnames))
        self.skipped_sheet.append(["指标类型", "指标名称", "文件路径", "缺失字段", "缺失信号"])
        self.skipped_count = 0
        self.closed = False

    def write_result(self, metric_name: str, result: MetricResult) -> None:
        self._write_metric_csv(metric_name, result)
        self._write_metric_json(metric_name, result)
        self._write_metric_sheet(metric_name, result)
        self._write_skipped(metric_name, result)
        self._update_summary(metric_name, result)

    def close(self) -> Path:
        if self.closed:
            return self.report_path
        for metric_name in self.metric_order:
            summary = self.summaries[metric_name]
            if summary["detail_count"] == 0:
                self.metric_sheets[metric_name].append([
                    summary["category"],
                    "无Bad_Case",
                    0,
                    "无Bad_Case",
                    "",
                    "",
                ])
            distance_km = sum(float(item) for item in summary["distance_by_path"].values())
            self.summary_sheet.append([
                summary["category"],
                metric_name,
                summary["total_cases"],
                summary["bad_cases"],
                summary["skipped_count"],
                metric_value_display(
                    summary["category"],
                    summary["total_cases"],
                    summary["bad_cases"],
                    distance_km,
                ),
            ])
        if self.skipped_count == 0:
            self.skipped_sheet.append(["无跳过指标", "", "", "", ""])

        for fp in self.json_files.values():
            fp.write("\n]\n")
            fp.close()
        for fp in self.csv_files.values():
            fp.close()
        self.skipped_fp.close()
        self.workbook.save(self.report_path)
        self.closed = True
        return self.report_path

    def _write_metric_csv(self, metric_name: str, result: MetricResult) -> None:
        self.csv_writers[metric_name].writerow(metric_csv_row(result))
        self.csv_files[metric_name].flush()

    def _write_metric_json(self, metric_name: str, result: MetricResult) -> None:
        fp = self.json_files[metric_name]
        if self.json_first[metric_name]:
            self.json_first[metric_name] = False
        else:
            fp.write(",\n")
        json.dump(metric_json_payload(result), fp, ensure_ascii=False)
        fp.flush()

    def _write_metric_sheet(self, metric_name: str, result: MetricResult) -> None:
        sheet = self.metric_sheets[metric_name]
        for event in result.events:
            sheet.append([
                result.category,
                "%s @ %.3f-%.3fs" % (result.source_path, event.start_s, event.end_s),
                event.severity,
                event.message,
                json.dumps(event.values, ensure_ascii=False),
                "",
            ])
            self.summaries[metric_name]["detail_count"] += 1

    def _write_skipped(self, metric_name: str, result: MetricResult) -> None:
        if result_status(result) != "SKIPPED":
            return
        row = skipped_metric_row(metric_name, result)
        self.skipped_writer.writerow(row)
        self.skipped_fp.flush()
        self.skipped_sheet.append([
            result.category,
            metric_name,
            result.source_path,
            ", ".join(result.missing_fields),
            format_missing_signal_mapping(result.missing_signals),
        ])
        self.skipped_count += 1

    def _update_summary(self, metric_name: str, result: MetricResult) -> None:
        summary = self.summaries[metric_name]
        summary["category"] = result.category
        summary["total_cases"] += result.total_cases
        summary["bad_cases"] += result.bad_cases
        summary["result_count"] += 1
        if result_status(result) == "SKIPPED":
            summary["skipped_count"] += 1
        if result.distance_km is not None and result.source_path not in summary["distance_by_path"]:
            summary["distance_by_path"][result.source_path] = float(result.distance_km)

    def __enter__(self) -> "StreamingReportWriter":
        return self

    def __exit__(self, exc_type: Any, exc: Any, tb: Any) -> None:
        self.close()


def write_outputs(results_by_metric: Mapping[str, List[MetricResult]], output_dir: Path, report_name: str) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    for metric_name, results in results_by_metric.items():
        write_metric_csv(metric_name, results, output_dir / ("%s.csv" % metric_name))
        write_metric_json(metric_name, results, output_dir / ("%s_events.json" % metric_name))
    write_skipped_metrics_csv(results_by_metric, output_dir / "skipped_metrics.csv")
    report_path = output_dir / ensure_xlsx(report_name)
    write_summary_xlsx(results_by_metric, report_path)
    return report_path


def write_metric_csv(metric_name: str, results: Iterable[MetricResult], csv_path: Path) -> None:
    with csv_path.open("w", newline="", encoding="utf-8-sig") as fp:
        writer = csv.DictWriter(fp, fieldnames=CSV_HEADERS)
        writer.writeheader()
        for result in results:
            writer.writerow(metric_csv_row(result))


def write_skipped_metrics_csv(
    results_by_metric: Mapping[str, List[MetricResult]],
    csv_path: Path,
) -> None:
    with csv_path.open("w", newline="", encoding="utf-8-sig") as fp:
        writer = csv.DictWriter(fp, fieldnames=SKIPPED_HEADERS)
        writer.writeheader()
        for metric_name in sorted_results_metric_names(results_by_metric):
            results = results_by_metric[metric_name]
            for result in results:
                if result_status(result) != "SKIPPED":
                    continue
                writer.writerow(skipped_metric_row(metric_name, result))


def write_metric_json(metric_name: str, results: Iterable[MetricResult], json_path: Path) -> None:
    payload = []
    for result in results:
        payload.append(metric_json_payload(result))
    with json_path.open("w", encoding="utf-8") as fp:
        json.dump(payload, fp, ensure_ascii=False, indent=2)


def write_summary_xlsx(results_by_metric: Mapping[str, List[MetricResult]], output_path: Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "所有指标统计汇总"
    write_summary_sheet(summary, results_by_metric)

    for metric_name in sorted_results_metric_names(results_by_metric):
        results = results_by_metric[metric_name]
        add_metric_sheet(workbook, result_category(results), metric_name, results)
    add_skipped_metrics_sheet(workbook, results_by_metric)

    for sheet in workbook.worksheets:
        for column_cells in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column_cells) + 4, 90)
            sheet.column_dimensions[column_cells[0].column_letter].width = width

    workbook.save(output_path)


def write_summary_sheet(
    sheet: Any,
    results_by_metric: Mapping[str, List[MetricResult]],
) -> None:
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    headers = ["指标类型", "指标名称", "总Case之和", "Bad Case之和", "跳过数", "指标值"]
    for col, title in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.alignment = center

    row_index = 2
    for metric_name in sorted_results_metric_names(results_by_metric):
        results = results_by_metric[metric_name]
        category = result_category(results)
        total_cases = sum(item.total_cases for item in results)
        bad_cases = sum(item.bad_cases for item in results)
        skipped_count = sum(1 for item in results if result_status(item) == "SKIPPED")
        distance_km = sum_distance_km(results)
        sheet.cell(row=row_index, column=1, value=category)
        sheet.cell(row=row_index, column=2, value=metric_name)
        sheet.cell(row=row_index, column=3, value=total_cases)
        sheet.cell(row=row_index, column=4, value=bad_cases)
        sheet.cell(row=row_index, column=5, value=skipped_count)
        sheet.cell(
            row=row_index,
            column=6,
            value=metric_value_display(category, total_cases, bad_cases, distance_km),
        )
        for col in range(1, 7):
            sheet.cell(row=row_index, column=col).alignment = center
        row_index += 1


def add_metric_sheet(workbook: Workbook, category: str, metric_name: str, results: List[MetricResult]) -> None:
    sheet_name = safe_sheet_name("%s_%s" % (category, metric_name), workbook.sheetnames)
    sheet = workbook.create_sheet(sheet_name)
    headers = ["指标类型", "Bad_Case数据列表", "问题严重程度得分", "问题内容", "问题分析", "分析截图"]
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    for col, title in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.alignment = center

    row = 2
    for result in results:
        for event in result.events:
            sheet.cell(row=row, column=1, value=result.category)
            sheet.cell(row=row, column=2, value="%s @ %.3f-%.3fs" % (result.source_path, event.start_s, event.end_s))
            sheet.cell(row=row, column=3, value=event.severity)
            sheet.cell(row=row, column=4, value=event.message)
            sheet.cell(row=row, column=5, value=json.dumps(event.values, ensure_ascii=False))
            row += 1

    if row == 2:
        sheet.cell(row=2, column=1, value=category)
        sheet.cell(row=2, column=2, value="无Bad_Case")
        sheet.cell(row=2, column=3, value=0)
        sheet.cell(row=2, column=4, value="无Bad_Case")


def add_skipped_metrics_sheet(workbook: Workbook, results_by_metric: Mapping[str, List[MetricResult]]) -> None:
    sheet = workbook.create_sheet(safe_sheet_name("跳过指标", workbook.sheetnames))
    headers = ["指标类型", "指标名称", "文件路径", "缺失字段", "缺失信号"]
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    for col, title in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.alignment = center

    row = 2
    for metric_name in sorted_results_metric_names(results_by_metric):
        results = results_by_metric[metric_name]
        for result in results:
            if result_status(result) != "SKIPPED":
                continue
            sheet.cell(row=row, column=1, value=result.category)
            sheet.cell(row=row, column=2, value=metric_name)
            sheet.cell(row=row, column=3, value=result.source_path)
            sheet.cell(row=row, column=4, value=", ".join(result.missing_fields))
            sheet.cell(row=row, column=5, value=format_missing_signal_mapping(result.missing_signals))
            row += 1

    if row == 2:
        sheet.cell(row=2, column=1, value="无跳过指标")


def metric_csv_row(result: MetricResult) -> Dict[str, Any]:
    return {
        "指标类型": result.category,
        "文件路径": result.source_path,
        "检查结果": result_status(result),
        "符合要求的内容": "; ".join(result.details),
        "缺失字段": ", ".join(result.missing_fields),
        "缺失信号": format_missing_signal_mapping(result.missing_signals),
        "总Case": result.total_cases,
        "Bad Case": result.bad_cases,
        "指标值": metric_value_display(
            result.category,
            result.total_cases,
            result.bad_cases,
            result.distance_km,
        ),
        "严重程度得分": result.severity_score,
    }


def metric_json_payload(result: MetricResult) -> Dict[str, Any]:
    return {
        "metric_name": result.metric_name,
        "category": result.category,
        "result_mode": result.result_mode,
        "source_path": result.source_path,
        "distance_km": result.distance_km,
        "distance_status": result.distance_status,
        "total_cases": result.total_cases,
        "bad_cases": result.bad_cases,
        "metric_value": metric_value_value(
            result.category,
            result.total_cases,
            result.bad_cases,
            result.distance_km,
        ),
        "metric_value_display": metric_value_display(
            result.category,
            result.total_cases,
            result.bad_cases,
            result.distance_km,
        ),
        "severity_score": result.severity_score,
        "runtime_s": result.runtime_s,
        "error": result.error,
        "status": result_status(result),
        "missing_fields": result.missing_fields,
        "missing_signals": result.missing_signals,
        "events": [
            {
                "start_s": event.start_s,
                "end_s": event.end_s,
                "duration_s": event.duration_s,
                "severity": event.severity,
                "message": event.message,
                "values": event.values,
            }
            for event in result.events
        ],
    }


def skipped_metric_row(metric_name: str, result: MetricResult) -> Dict[str, Any]:
    return {
        "指标类型": result.category,
        "指标名称": metric_name,
        "文件路径": result.source_path,
        "缺失字段": ", ".join(result.missing_fields),
        "缺失信号": format_missing_signal_mapping(result.missing_signals),
        "处理结果": "该文件下该指标未运行",
    }


def result_category(results: Iterable[MetricResult]) -> str:
    for result in results:
        return result.category
    return "CPI"


def sorted_results_metric_names(results_by_metric: Mapping[str, List[MetricResult]]) -> List[str]:
    return sorted(
        results_by_metric,
        key=lambda metric_name: metric_report_sort_key(
            metric_name,
            result_category(results_by_metric[metric_name]),
        ),
    )


def metric_report_sort_key(metric_name: str, category: str) -> tuple[int, str, str]:
    category_rank = {"MPI": 0, "CPI": 1}.get(str(category).upper(), 2)
    return (category_rank, metric_scene_group(metric_name), metric_name)


def metric_scene_group(metric_name: str) -> str:
    groups = (
        "abs",
        "acc",
        "curvature",
        "follow_start",
        "follow_stop",
        "slope_no_frnt",
        "stop2go",
        "tcs",
        "vse",
    )
    for group in groups:
        if metric_name == group or metric_name.startswith(group + "_"):
            return group
    return metric_name


def sum_distance_km(results: Iterable[MetricResult]) -> float:
    total = 0.0
    seen_paths = set()
    for result in results:
        if result.source_path in seen_paths:
            continue
        seen_paths.add(result.source_path)
        if result.distance_km is None:
            continue
        total += float(result.distance_km)
    return total


def metric_value_value(
    category: str,
    total_cases: int,
    bad_cases: int,
    distance_km: Any,
) -> Any:
    if category == "MPI":
        if bad_cases <= 0 or distance_km is None:
            return None
        return float(distance_km) / float(bad_cases)
    if bad_cases == 0:
        return 100.0
    return float(total_cases) / float(bad_cases)


def metric_value_display(
    category: str,
    total_cases: int,
    bad_cases: int,
    distance_km: Any,
) -> str:
    value = metric_value_value(category, total_cases, bad_cases, distance_km)
    if category == "MPI":
        if distance_km is None:
            return "N/A"
        if bad_cases <= 0:
            return ">%.3f(%.3f/0)" % (float(distance_km), float(distance_km))
        return "%.3f(%.3f/%d)" % (float(value), float(distance_km), bad_cases)
    if value is None:
        return "N/A"
    if bad_cases == 0:
        return "100.0(%d/%d)" % (total_cases, bad_cases)
    return "%.1f(%d/%d)" % (float(value), total_cases, bad_cases)


def safe_sheet_name(name: str, existing: Iterable[str]) -> str:
    invalid = set(r"[]:*?/\\")
    cleaned = "".join("_" if ch in invalid else ch for ch in name)[:31] or "metric"
    if cleaned not in existing:
        return cleaned
    base = cleaned[:28]
    suffix = 1
    while "%s_%d" % (base, suffix) in existing:
        suffix += 1
    return "%s_%d" % (base, suffix)


def ensure_xlsx(name: str) -> str:
    return name if name.lower().endswith(".xlsx") else "%s.xlsx" % name


def make_run_output_dir(base_dir: Path, stem: str) -> Path:
    cleaned = stem.strip() or "run"
    candidate = base_dir / cleaned
    if not candidate.exists():
        candidate.mkdir(parents=True, exist_ok=False)
        return candidate
    suffix = 1
    while True:
        candidate = base_dir / ("%s_%02d" % (cleaned, suffix))
        if not candidate.exists():
            candidate.mkdir(parents=True, exist_ok=False)
            return candidate
        suffix += 1


def result_status(result: MetricResult) -> str:
    if result.status:
        return result.status
    if result.error:
        return "ERROR"
    if result.bad_cases > 0:
        return "BAD"
    return "OK"


def format_missing_signal_mapping(missing_signals: Mapping[str, str]) -> str:
    if not missing_signals:
        return ""
    return "; ".join("%s=%s" % (field_name, signal_path) for field_name, signal_path in missing_signals.items())
