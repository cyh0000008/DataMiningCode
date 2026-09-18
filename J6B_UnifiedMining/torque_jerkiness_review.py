from __future__ import annotations

import argparse
import json
import math
import shutil
import tempfile
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

import numpy as np
from openpyxl import Workbook
from openpyxl.drawing.image import Image as WorkbookImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageDraw, ImageFont

from DataMining import load_config, resolve_path
from metric_utils import derivative, time_window_moving_average
from signal_preprocessor import build_signal_frame


METRIC_NAME = "acc_torque_jerkiness_count"
DEFAULT_OUTPUT = "acc_torque_jerkiness_review.xlsx"
ACCELERATION_SMOOTHING_WINDOW_S = 0.12
ACCELERATION_TREND_WINDOW_S = 0.80
TORQUE_SMOOTHING_WINDOW_S = 0.08


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = parse_args(argv)
    config_path = args.config.resolve()
    config = load_config(config_path)
    report_dir = resolve_report_dir(args.report_dir, config_path.parent, config)
    event_path = report_dir / ("%s_events.json" % METRIC_NAME)
    if not event_path.exists():
        raise FileNotFoundError("Metric event file not found: %s" % event_path)

    output_path = args.output
    if output_path is None:
        output_path = config_path.parent / DEFAULT_OUTPUT
    elif not output_path.is_absolute():
        output_path = (Path.cwd() / output_path).resolve()

    rows = read_event_rows(event_path)
    write_review_workbook(rows, config, output_path)
    print("Review workbook written: %s" % output_path)
    print("Torque-validated jerkiness events: %d" % len(rows))
    return 0


def parse_args(argv: Optional[Sequence[str]]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate torque jerkiness review workbook.")
    parser.add_argument("--config", type=Path, default=Path("config.json"), help="config.json path")
    parser.add_argument("--report-dir", type=Path, default=None, help="DataMining output run folder")
    parser.add_argument("--output", type=Path, default=None, help="xlsx output path")
    return parser.parse_args(argv)


def resolve_report_dir(
    report_dir: Optional[Path],
    base_dir: Path,
    config: Mapping[str, Any],
) -> Path:
    if report_dir is not None:
        return report_dir if report_dir.is_absolute() else (Path.cwd() / report_dir).resolve()
    output_base = resolve_path(base_dir, config.get("output_dir", "Result_Test"))
    candidates = [
        path
        for path in output_base.iterdir()
        if path.is_dir() and (path / ("%s_events.json" % METRIC_NAME)).exists()
    ]
    if not candidates:
        raise FileNotFoundError(
            "No output folder with %s_events.json found under %s" % (METRIC_NAME, output_base)
        )
    return max(candidates, key=lambda item: item.stat().st_mtime)


def read_event_rows(event_path: Path) -> List[Dict[str, Any]]:
    payloads = json.loads(event_path.read_text(encoding="utf-8"))
    rows: List[Dict[str, Any]] = []
    for payload in payloads:
        source_path = str(payload.get("source_path", ""))
        for event in payload.get("events", []):
            rows.append(
                {
                    "source_path": source_path,
                    "file_name": Path(source_path).name,
                    "start_s": float(event.get("start_s", 0.0)),
                    "end_s": float(event.get("end_s", 0.0)),
                    "duration_s": float(event.get("duration_s", 0.0)),
                    "severity": float(event.get("severity", 0.0)),
                    "message": str(event.get("message", "")),
                    "values": dict(event.get("values", {})),
                }
            )
    return sorted(rows, key=lambda item: (item["file_name"], item["start_s"], item["end_s"]))


def write_review_workbook(
    rows: List[Dict[str, Any]],
    config: Mapping[str, Any],
    output_path: Path,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "顿挫确认列表"
    headers = [
        "序号",
        "文件名",
        "文件路径",
        "开始时间(s)",
        "结束时间(s)",
        "持续时间(s)",
        "严重度",
        "正向jerk峰值",
        "负向jerk峰值",
        "高频加速度RMS",
        "驱动扭矩变化率峰值",
        "制动扭矩变化率峰值",
        "扭矩-jerk最大延时(s)",
        "顿挫区域信号图示",
        "人工确认是否真正顿挫",
        "备注",
    ]
    sheet.append(headers)
    style_header(sheet)
    configure_columns(sheet)
    add_confirmation_validation(sheet, max(len(rows) + 50, 200))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:P1"

    if not rows:
        sheet.append([1, "无筛选出的扭矩校验顿挫", "", "", "", "", "", "", "", "", "", "", "", "", "", ""])
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        return

    temp_dir = Path(tempfile.mkdtemp(prefix="torque_jerkiness_plots_"))
    try:
        row_index = 2
        grouped: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
        for row in rows:
            grouped[row["source_path"]].append(row)

        sequence = 1
        for source_path, event_rows in grouped.items():
            frame = build_signal_frame(source_path, config)
            for row in event_rows:
                values = row["values"]
                image_path = temp_dir / ("event_%04d.png" % sequence)
                draw_event_plot(frame, row, image_path)
                sheet.append(
                    [
                        sequence,
                        row["file_name"],
                        row["source_path"],
                        row["start_s"],
                        row["end_s"],
                        row["duration_s"],
                        row["severity"],
                        values.get("positive_jerk_peak", ""),
                        values.get("negative_jerk_peak", ""),
                        values.get("highpass_acc_rms", ""),
                        values.get("drive_rate_peak_near_positive_jerk", ""),
                        values.get("brake_rate_peak_near_negative_jerk", ""),
                        max_float(
                            values.get("drive_rate_peak_lag_s"),
                            values.get("brake_rate_peak_lag_s"),
                        ),
                        "",
                        "",
                        "",
                    ]
                )
                sheet.row_dimensions[row_index].height = 245
                image = WorkbookImage(str(image_path))
                image.width = 760
                image.height = 305
                sheet.add_image(image, "N%d" % row_index)
                style_data_row(sheet, row_index)
                row_index += 1
                sequence += 1
    finally:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        shutil.rmtree(temp_dir, ignore_errors=True)


def style_header(sheet: Any) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 28


def configure_columns(sheet: Any) -> None:
    widths = {
        "A": 8,
        "B": 32,
        "C": 48,
        "D": 12,
        "E": 12,
        "F": 12,
        "G": 12,
        "H": 13,
        "I": 18,
        "J": 18,
        "K": 20,
        "L": 20,
        "M": 18,
        "N": 105,
        "O": 22,
        "P": 34,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def add_confirmation_validation(sheet: Any, max_row: int) -> None:
    validation = DataValidation(type="list", formula1='"是,否,待确认"', allow_blank=True)
    validation.error = "请选择：是、否、待确认"
    validation.errorTitle = "无效输入"
    validation.prompt = "确认该行是否为真正顿挫"
    validation.promptTitle = "人工确认"
    sheet.add_data_validation(validation)
    validation.add("O2:O%d" % max_row)


def style_data_row(sheet: Any, row_index: int) -> None:
    for column in range(1, 17):
        cell = sheet.cell(row=row_index, column=column)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column in ("D", "E", "F", "G", "H", "I", "J", "K", "L", "M"):
        sheet["%s%d" % (column, row_index)].number_format = "0.000"


def draw_event_plot(frame: Any, row: Mapping[str, Any], output_path: Path) -> None:
    t = frame.time_s
    event_start = float(row["start_s"])
    event_end = float(row["end_s"])
    start = max(0, int(np.searchsorted(t, event_start - 1.2, side="left")))
    end = min(frame.length - 1, int(np.searchsorted(t, event_end + 1.2, side="right") - 1))
    if end <= start:
        end = min(frame.length - 1, start + 1)
    rel_t = t[start : end + 1] - event_start
    values = row["values"]
    acceleration = frame.col("acceleration").astype(float)
    acceleration_smooth = time_window_moving_average(t, acceleration, ACCELERATION_SMOOTHING_WINDOW_S)
    acceleration_trend = time_window_moving_average(t, acceleration, ACCELERATION_TREND_WINDOW_S)
    acceleration_highpass = acceleration_smooth - acceleration_trend
    jerk = derivative(t, acceleration_smooth)
    drive_torque = time_window_moving_average(
        t,
        frame.col("actual_drive_torque").astype(float),
        TORQUE_SMOOTHING_WINDOW_S,
    )
    brake_torque = time_window_moving_average(
        t,
        frame.col("actual_brake_torque").astype(float),
        TORQUE_SMOOTHING_WINDOW_S,
    )
    drive_rate = derivative(t, drive_torque)
    brake_rate = derivative(t, brake_torque)
    window = slice(start, end + 1)
    series_groups = [
        (
            "加速度 m/s^2",
            [(acceleration_smooth[window], (36, 104, 179), "")],
        ),
        (
            "jerk m/s^3",
            [(jerk[window], (111, 66, 193), "")],
        ),
        (
            "高频加速度",
            [(acceleration_highpass[window], (226, 139, 29), "")],
        ),
        (
            "扭矩",
            [
                (drive_torque[window], (38, 139, 80), "驱动"),
                (brake_torque[window], (196, 70, 62), "制动"),
            ],
        ),
        (
            "扭矩变化率",
            [
                (drive_rate[window], (38, 139, 80), "驱动"),
                (brake_rate[window], (196, 70, 62), "制动"),
            ],
        ),
    ]

    image = Image.new("RGB", (980, 370), "white")
    draw = ImageDraw.Draw(image)
    fonts = FontSet()
    title = "%s  %.3f-%.3fs" % (row["file_name"], event_start, event_end)
    draw.text((18, 10), title, fill=(28, 38, 52), font=fonts.header)
    summary = (
        "jerk +%s / %s | 高频RMS %s | 驱动峰值 %s | 制动峰值 %s | 最大延时 %ss"
        % (
            fmt(values.get("positive_jerk_peak")),
            fmt(values.get("negative_jerk_peak")),
            fmt(values.get("highpass_acc_rms")),
            fmt(values.get("drive_rate_peak_near_positive_jerk")),
            fmt(values.get("brake_rate_peak_near_negative_jerk")),
            fmt(
                max_float(
                    values.get("drive_rate_peak_lag_s"),
                    values.get("brake_rate_peak_lag_s"),
                )
            ),
        )
    )
    draw.text((18, 36), summary, fill=(83, 96, 112), font=fonts.small)

    plot_left = 95
    plot_right = 905
    lane_top = 68
    lane_height = 45
    lane_gap = 8
    x_min = float(np.nanmin(rel_t)) if rel_t.size else -1.0
    x_max = float(np.nanmax(rel_t)) if rel_t.size else 1.0
    if abs(x_max - x_min) < 1e-9:
        x_max = x_min + 1.0

    def x_map(value: float) -> int:
        return plot_left + int((value - x_min) / (x_max - x_min) * (plot_right - plot_left))

    event_x1 = x_map(0.0)
    event_x2 = x_map(max(event_end - event_start, 0.0))
    for idx, (label, lane_series) in enumerate(series_groups):
        top = lane_top + idx * (lane_height + lane_gap)
        bottom = top + lane_height
        draw.rectangle((plot_left, top, plot_right, bottom), outline=(189, 198, 210), width=1)
        draw.rectangle((event_x1, top, event_x2, bottom), fill=(255, 244, 190))
        draw.rectangle((plot_left, top, plot_right, bottom), outline=(189, 198, 210), width=1)
        draw.text((12, top + 15), label, fill=(40, 50, 65), font=fonts.tiny)
        draw_multi_series(draw, rel_t, lane_series, plot_left, plot_right, top, bottom, fonts)
        draw.line((event_x1, top, event_x1, bottom), fill=(125, 88, 25), width=1)
        draw.line((event_x2, top, event_x2, bottom), fill=(125, 88, 25), width=1)

    axis_y = lane_top + len(series_groups) * (lane_height + lane_gap) - lane_gap + 14
    for tick in nice_ticks(x_min, x_max, 6):
        x = x_map(tick)
        draw.line((x, axis_y - 5, x, axis_y), fill=(140, 150, 164), width=1)
        draw.text((x, axis_y + 2), "%.1fs" % tick, fill=(83, 96, 112), font=fonts.tiny, anchor="mt")

    image.save(output_path, quality=95)


def draw_multi_series(
    draw: ImageDraw.ImageDraw,
    x_values: np.ndarray,
    series: Sequence[Tuple[np.ndarray, Tuple[int, int, int], str]],
    left: int,
    right: int,
    top: int,
    bottom: int,
    fonts: "FontSet",
) -> None:
    finite_values = [
        np.asarray(y_values, dtype=float)[np.isfinite(y_values)]
        for y_values, _color, _label in series
    ]
    finite_values = [values for values in finite_values if values.size]
    if not finite_values:
        draw.text(
            ((left + right) // 2, (top + bottom) // 2),
            "无有效数据",
            fill=(120, 130, 145),
            font=fonts.tiny,
            anchor="mm",
        )
        return

    combined = np.concatenate(finite_values)
    y_min = float(np.nanmin(combined))
    y_max = float(np.nanmax(combined))
    if y_min < 0.0 < y_max:
        zero_y = y_to_pixel(0.0, y_min, y_max, top, bottom)
        draw.line((left, zero_y, right, zero_y), fill=(214, 220, 229), width=1)

    for series_index, (y_values, color, label) in enumerate(series):
        draw_series_line(draw, x_values, y_values, left, right, top, bottom, y_min, y_max, color)
        if label:
            label_x = left + 8 + 48 * series_index
            draw.text((label_x, top + 2), label, fill=color, font=fonts.tiny)

    draw.text((right + 6, top + 2), compact_number(y_max), fill=(83, 96, 112), font=fonts.tiny)
    draw.text((right + 6, bottom - 14), compact_number(y_min), fill=(83, 96, 112), font=fonts.tiny)


def draw_series_line(
    draw: ImageDraw.ImageDraw,
    x_values: np.ndarray,
    y_values: np.ndarray,
    left: int,
    right: int,
    top: int,
    bottom: int,
    y_min: float,
    y_max: float,
    color: Tuple[int, int, int],
) -> None:
    finite = np.isfinite(y_values)
    if not np.any(finite):
        return
    if abs(y_max - y_min) < 1e-9:
        padding = max(abs(y_max) * 0.1, 1.0)
        y_min -= padding
        y_max += padding
    else:
        padding = (y_max - y_min) * 0.08
        y_min -= padding
        y_max += padding
    x_min = float(np.nanmin(x_values))
    x_max = float(np.nanmax(x_values))

    def px(value: float) -> int:
        return left + int((value - x_min) / max(x_max - x_min, 1e-9) * (right - left))

    def py(value: float) -> int:
        return y_to_pixel(value, y_min, y_max, top, bottom)

    points = [(px(float(x)), py(float(y))) for x, y, ok in zip(x_values, y_values, finite) if ok]
    if len(points) >= 2:
        draw.line(points, fill=color, width=2)
    elif points:
        x, y = points[0]
        draw.ellipse((x - 2, y - 2, x + 2, y + 2), fill=color)


def y_to_pixel(value: float, y_min: float, y_max: float, top: int, bottom: int) -> int:
    if abs(y_max - y_min) < 1e-9:
        return (top + bottom) // 2
    return bottom - int((value - y_min) / (y_max - y_min) * (bottom - top))


class FontSet:
    def __init__(self) -> None:
        regular, bold = find_chinese_fonts()
        self.header = ImageFont.truetype(str(bold), 17)
        self.small = ImageFont.truetype(str(regular), 14)
        self.tiny = ImageFont.truetype(str(regular), 12)


def find_chinese_fonts() -> Tuple[Path, Path]:
    candidates = [
        (Path(r"C:\Windows\Fonts\msyh.ttc"), Path(r"C:\Windows\Fonts\msyhbd.ttc")),
        (Path(r"C:\Windows\Fonts\simhei.ttf"), Path(r"C:\Windows\Fonts\simhei.ttf")),
        (Path(r"C:\Windows\Fonts\simsun.ttc"), Path(r"C:\Windows\Fonts\simsun.ttc")),
    ]
    for regular, bold in candidates:
        if regular.exists() and bold.exists():
            return regular, bold
    raise FileNotFoundError("未找到可渲染中文的字体。")


def nice_ticks(value_min: float, value_max: float, count: int) -> List[float]:
    if not math.isfinite(value_min) or not math.isfinite(value_max) or value_min == value_max:
        return [value_min]
    span = value_max - value_min
    raw_step = span / max(1, count - 1)
    magnitude = 10 ** math.floor(math.log10(abs(raw_step)))
    step = magnitude
    for multiplier in (1, 2, 2.5, 5, 10):
        step = multiplier * magnitude
        if span / step <= count:
            break
    current = math.floor(value_min / step) * step
    ticks: List[float] = []
    while current <= value_max + step * 0.5:
        if current >= value_min - step * 0.5:
            ticks.append(float(current))
        current += step
    return ticks


def fmt(value: Any) -> str:
    try:
        return "%.3f" % float(value)
    except (TypeError, ValueError):
        return ""


def max_float(*values: Any) -> Any:
    parsed = []
    for value in values:
        try:
            parsed.append(float(value))
        except (TypeError, ValueError):
            continue
    if not parsed:
        return ""
    return max(parsed)


def compact_number(value: float) -> str:
    if abs(value) >= 1000:
        return "%.0f" % value
    if abs(value) >= 10:
        return "%.1f" % value
    return "%.2f" % value


if __name__ == "__main__":
    raise SystemExit(main())
