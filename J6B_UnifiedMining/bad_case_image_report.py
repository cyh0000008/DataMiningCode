from __future__ import annotations

import argparse
import json
import math
import queue
import shutil
import sys
import tempfile
import threading
import traceback
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

import numpy as np
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as WorkbookImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFont

from DataMining import h5_files_are_continuous, load_config
from metric_utils import derivative, display_speed_array, time_window_moving_average
from signal_preprocessor import SignalFrame, build_signal_frame, stitch_signal_frames


BAD_CASE_SHEET_NAME = "Bad Case提取"
DEFAULT_OUTPUT_NAME = "bad_case_report_with_image.xlsx"
IMAGE_HEADER = "关键信号图"
STATUS_HEADER = "绘图状态"
PLOT_WIDTH = 1180
PLOT_HEIGHT = 650
EXCEL_IMAGE_WIDTH = 930
EXCEL_IMAGE_HEIGHT = 512


@dataclass
class CaseRow:
    row_index: int
    metric_name: str
    source_path: str
    start_s: float
    end_s: float
    severity: Any
    message: str


@dataclass
class PlotLine:
    label: str
    values: np.ndarray
    color: Tuple[int, int, int]
    step: bool = False


@dataclass
class PlotPanel:
    title: str
    unit: str
    lines: List[PlotLine]
    fixed_range: Optional[Tuple[float, float]] = None


class FontSet:
    def __init__(self) -> None:
        self.title = load_font(25, True)
        self.header = load_font(18, True)
        self.normal = load_font(15, False)
        self.small = load_font(12, False)


class FrameCache:
    def __init__(self, config: Mapping[str, Any], manifest: Mapping[str, Any]) -> None:
        self.config = config
        self.manifest = manifest
        self.cache: Dict[str, SignalFrame] = {}

    def get(self, source_path: str, needed_end_s: float) -> SignalFrame:
        normalized = str(Path(source_path).resolve())
        cache_key = "%s|%.1f" % (normalized, math.ceil(max(needed_end_s, 0.0) / 300.0))
        if cache_key not in self.cache:
            self.cache[cache_key] = self._build_frame(normalized, needed_end_s)
        return self.cache[cache_key]

    def _build_frame(self, source_path: str, needed_end_s: float) -> SignalFrame:
        frame = build_signal_frame(source_path, self.config)
        if frame.length < 2 or needed_end_s <= float(frame.time_s[-1]) + 0.5:
            return frame

        h5_files = [str(Path(item).resolve()) for item in self.manifest.get("h5_files", [])]
        try:
            current_index = h5_files.index(str(Path(source_path).resolve()))
        except ValueError:
            return frame
        if current_index + 1 >= len(h5_files):
            return frame

        next_path = h5_files[current_index + 1]
        is_continuous, next_offset_s, _boundary_gap_s = h5_files_are_continuous(
            source_path,
            next_path,
            frame,
            float(self.config.get("defaults", {}).get("h5_stitch_max_gap_s", 5.0)),
        )
        if not is_continuous:
            return frame
        next_frame = build_signal_frame(next_path, self.config)
        return stitch_signal_frames(frame, next_frame, next_offset_s)


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Add key signal plots to bad_case_report.xlsx.")
    parser.add_argument("--input", type=Path, help="bad_case_report.xlsx path")
    parser.add_argument("--output", type=Path, default=None, help="output xlsx path")
    args = parser.parse_args(argv)
    if args.input:
        output = generate_report_with_images(args.input, args.output, print)
        print("Output written: %s" % output)
        return 0
    launch_gui()
    return 0


def launch_gui() -> None:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    root = tk.Tk()
    root.title("Bad Case 报告信号图生成工具")
    root.geometry("760x430")
    root.minsize(680, 360)

    selected_path = tk.StringVar()
    status_var = tk.StringVar(value="请选择 bad_case_report.xlsx")
    log_queue: "queue.Queue[Tuple[str, str]]" = queue.Queue()

    def browse_file() -> None:
        path = filedialog.askopenfilename(
            title="选择 bad_case_report.xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx"), ("所有文件", "*.*")],
        )
        if path:
            selected_path.set(path)
            status_var.set("已选择：%s" % path)

    def append_log(text: str) -> None:
        log_text.configure(state="normal")
        log_text.insert("end", text + "\n")
        log_text.see("end")
        log_text.configure(state="disabled")

    def set_running(running: bool) -> None:
        run_button.configure(state=("disabled" if running else "normal"))
        browse_button.configure(state=("disabled" if running else "normal"))
        progress.configure(mode=("indeterminate" if running else "determinate"))
        if running:
            progress.start(12)
        else:
            progress.stop()
            progress["value"] = 0

    def run_worker() -> None:
        input_text = selected_path.get().strip()
        if not input_text:
            messagebox.showwarning("请选择文件", "请先选择 bad_case_report.xlsx")
            return
        input_path = Path(input_text)
        if not input_path.exists():
            messagebox.showerror("文件不存在", str(input_path))
            return
        set_running(True)
        status_var.set("正在生成信号图，请稍等...")
        append_log("开始处理：%s" % input_path)

        def worker() -> None:
            try:
                output = generate_report_with_images(input_path, None, lambda msg: log_queue.put(("log", msg)))
                log_queue.put(("done", str(output)))
            except Exception:
                log_queue.put(("error", traceback.format_exc()))

        threading.Thread(target=worker, daemon=True).start()

    def poll_queue() -> None:
        try:
            while True:
                kind, payload = log_queue.get_nowait()
                if kind == "log":
                    append_log(payload)
                elif kind == "done":
                    set_running(False)
                    status_var.set("完成：%s" % payload)
                    append_log("完成输出：%s" % payload)
                    messagebox.showinfo("生成完成", "已生成：\n%s" % payload)
                elif kind == "error":
                    set_running(False)
                    status_var.set("生成失败")
                    append_log(payload)
                    messagebox.showerror("生成失败", payload[-1800:])
        except queue.Empty:
            pass
        root.after(150, poll_queue)

    frame = ttk.Frame(root, padding=16)
    frame.pack(fill="both", expand=True)
    ttk.Label(frame, text="选择 bad_case_report.xlsx：").grid(row=0, column=0, sticky="w")
    entry = ttk.Entry(frame, textvariable=selected_path)
    entry.grid(row=1, column=0, sticky="ew", padx=(0, 10), pady=(6, 10))
    browse_button = ttk.Button(frame, text="浏览...", command=browse_file)
    browse_button.grid(row=1, column=1, sticky="ew", pady=(6, 10))
    run_button = ttk.Button(frame, text="生成 with_image 报告", command=run_worker)
    run_button.grid(row=2, column=1, sticky="ew", pady=(0, 10))
    progress = ttk.Progressbar(frame)
    progress.grid(row=2, column=0, sticky="ew", padx=(0, 10), pady=(0, 10))
    ttk.Label(frame, textvariable=status_var).grid(row=3, column=0, columnspan=2, sticky="w", pady=(0, 8))
    log_text = tk.Text(frame, height=12, state="disabled", wrap="word")
    log_text.grid(row=4, column=0, columnspan=2, sticky="nsew")
    scrollbar = ttk.Scrollbar(frame, command=log_text.yview)
    scrollbar.grid(row=4, column=2, sticky="ns")
    log_text.configure(yscrollcommand=scrollbar.set)
    frame.columnconfigure(0, weight=1)
    frame.rowconfigure(4, weight=1)
    poll_queue()
    root.mainloop()


def generate_report_with_images(
    bad_case_report_path: Path,
    output_path: Optional[Path] = None,
    log: Optional[Callable[[str], None]] = None,
) -> Path:
    logger = log or (lambda _msg: None)
    input_path = Path(bad_case_report_path).resolve()
    if not input_path.exists():
        raise FileNotFoundError("未找到 bad_case_report.xlsx：%s" % input_path)
    if output_path is None:
        output_path = input_path.with_name(DEFAULT_OUTPUT_NAME)
    else:
        output_path = Path(output_path).resolve()

    manifest = load_manifest(input_path.parent)
    config_path = resolve_config_path(input_path.parent, manifest)
    logger("使用配置：%s" % config_path)
    config = load_config(config_path)
    frame_cache = FrameCache(config, manifest)

    workbook = load_workbook(input_path)
    sheet = workbook[BAD_CASE_SHEET_NAME] if BAD_CASE_SHEET_NAME in workbook.sheetnames else workbook.active
    header_map = header_columns(sheet)
    cases = read_cases(sheet, header_map)
    logger("读取到 %d 条 case" % len(cases))

    image_col = ensure_output_columns(sheet)
    temp_dir = Path(tempfile.mkdtemp(prefix="bad_case_signal_plots_"))
    image_paths: List[Path] = []
    try:
        for index, case in enumerate(cases, 1):
            logger("[%d/%d] %s %.3f-%.3fs" % (index, len(cases), case.metric_name, case.start_s, case.end_s))
            try:
                frame = frame_cache.get(case.source_path, case.end_s)
                image_path = temp_dir / ("case_%04d.png" % index)
                draw_case_plot(frame, case, image_path)
                image_paths.append(image_path)
                add_image_to_sheet(sheet, case.row_index, image_col, image_path)
            except Exception as exc:
                sheet.cell(row=case.row_index, column=image_col, value="绘图失败：%s" % exc)
                sheet.cell(row=case.row_index, column=image_col).alignment = Alignment(wrap_text=True, vertical="center")
                logger("  绘图失败：%s" % exc)

        style_output_workbook(sheet, image_col)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        logger("保存成功：%s" % output_path)
        return output_path
    finally:
        workbook.close()
        shutil.rmtree(temp_dir, ignore_errors=True)


def load_manifest(report_dir: Path) -> Dict[str, Any]:
    manifest_path = report_dir / "run_manifest.json"
    if not manifest_path.exists():
        return {}
    try:
        data = json.loads(manifest_path.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def resolve_config_path(report_dir: Path, manifest: Mapping[str, Any]) -> Path:
    candidates: List[Path] = []
    manifest_config = manifest.get("config_path")
    if manifest_config:
        candidates.append(Path(str(manifest_config)))
    candidates.append(report_dir / "config.json")
    candidates.append(Path.cwd() / "config.json")
    bundle_dir = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    candidates.append(bundle_dir / "config.json")
    candidates.append(Path(__file__).resolve().parent / "config.json")
    for candidate in candidates:
        try:
            if candidate.exists():
                return candidate.resolve()
        except OSError:
            continue
    raise FileNotFoundError("未找到 config.json。请确认 bad_case_report.xlsx 同目录有 run_manifest.json，或将工具放在工程目录运行。")


def header_columns(sheet: Any) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for cell in sheet[1]:
        value = normalize_header(cell.value)
        if value:
            mapping[value] = int(cell.column)
    return mapping


def normalize_header(value: Any) -> str:
    return "" if value is None else str(value).strip().replace(" ", "")


def read_cases(sheet: Any, header_map: Mapping[str, int]) -> List[CaseRow]:
    source_col = require_header(header_map, "文件路径")
    metric_col = require_header(header_map, "指标名称")
    start_col = require_header(header_map, "开始时间(s)")
    end_col = require_header(header_map, "结束时间(s)")
    severity_col = header_map.get("严重程度得分")
    message_col = header_map.get("问题内容")
    cases: List[CaseRow] = []
    for row_index in range(2, sheet.max_row + 1):
        source_path = str(sheet.cell(row=row_index, column=source_col).value or "").strip()
        if not source_path or source_path == "无匹配Bad Case":
            continue
        start_s = as_float(sheet.cell(row=row_index, column=start_col).value)
        end_s = as_float(sheet.cell(row=row_index, column=end_col).value)
        if start_s is None or end_s is None:
            continue
        if end_s < start_s:
            start_s, end_s = end_s, start_s
        cases.append(
            CaseRow(
                row_index=row_index,
                metric_name=str(sheet.cell(row=row_index, column=metric_col).value or "").strip(),
                source_path=source_path,
                start_s=float(start_s),
                end_s=float(end_s),
                severity=sheet.cell(row=row_index, column=severity_col).value if severity_col else "",
                message=str(sheet.cell(row=row_index, column=message_col).value or "") if message_col else "",
            )
        )
    return cases


def require_header(header_map: Mapping[str, int], name: str) -> int:
    key = normalize_header(name)
    if key not in header_map:
        raise KeyError("bad_case_report.xlsx 缺少列：%s" % name)
    return header_map[key]


def as_float(value: Any) -> Optional[float]:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def ensure_output_columns(sheet: Any) -> int:
    existing: Optional[int] = None
    for cell in sheet[1]:
        if normalize_header(cell.value) == IMAGE_HEADER:
            existing = int(cell.column)
            break
    image_col = existing or (sheet.max_column + 1)
    header = sheet.cell(row=1, column=image_col, value=IMAGE_HEADER)
    header.font = Font(bold=True, color="FFFFFF")
    header.fill = PatternFill("solid", fgColor="1F4E78")
    header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.column_dimensions[header.column_letter].width = 132
    return image_col


def add_image_to_sheet(sheet: Any, row_index: int, image_col: int, image_path: Path) -> None:
    image = WorkbookImage(str(image_path))
    image.width = EXCEL_IMAGE_WIDTH
    image.height = EXCEL_IMAGE_HEIGHT
    sheet.row_dimensions[row_index].height = 392
    sheet.add_image(image, "%s%d" % (sheet.cell(row=1, column=image_col).column_letter, row_index))
    cell = sheet.cell(row=row_index, column=image_col)
    cell.value = ""
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_output_workbook(sheet: Any, image_col: int) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in range(2, sheet.max_row + 1):
        for col in range(1, image_col + 1):
            sheet.cell(row=row, column=col).alignment = Alignment(vertical="center", wrap_text=True)


def draw_case_plot(frame: SignalFrame, case: CaseRow, output_path: Path) -> None:
    if frame.length < 2:
        raise ValueError("H5 信号长度不足")
    t = frame.time_s.astype(float)
    start_s, end_s = case.start_s, case.end_s
    if start_s > float(t[-1]) or end_s < float(t[0]):
        raise ValueError("case 时间 %.3f-%.3fs 不在 H5 时间范围 %.3f-%.3fs 内" % (start_s, end_s, float(t[0]), float(t[-1])))

    window_start, window_end = plot_window_bounds(t, start_s, end_s)
    mask = (t >= window_start) & (t <= window_end)
    if int(np.sum(mask)) < 2:
        raise ValueError("case 附近可绘制采样点不足")

    panels = build_panels(frame, case.metric_name, mask)
    if not panels:
        raise ValueError("没有可绘制的关键信号")

    fonts = FontSet()
    image = Image.new("RGB", (PLOT_WIDTH, PLOT_HEIGHT), (246, 248, 252))
    draw = ImageDraw.Draw(image)
    draw_plot_header(draw, fonts, case, window_start, window_end)

    panel_count = min(len(panels), 5)
    panel_top = 95
    panel_gap = 12
    panel_height = int((PLOT_HEIGHT - panel_top - 22 - panel_gap * (panel_count - 1)) / panel_count)
    rel_t = t[mask]
    for panel_index, panel in enumerate(panels[:panel_count]):
        top = panel_top + panel_index * (panel_height + panel_gap)
        box = (22, top, PLOT_WIDTH - 22, top + panel_height)
        draw_panel(draw, box, panel, rel_t, start_s, end_s, fonts)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(output_path)


def plot_window_bounds(t: np.ndarray, start_s: float, end_s: float) -> Tuple[float, float]:
    duration = max(end_s - start_s, 0.2)
    pad = min(max(duration * 0.35, 3.0), 8.0)
    window_start = max(float(t[0]), start_s - pad)
    window_end = min(float(t[-1]), end_s + pad)
    if window_end - window_start > 70.0:
        center = (start_s + end_s) / 2.0
        window_start = max(float(t[0]), center - 35.0)
        window_end = min(float(t[-1]), center + 35.0)
    if window_end <= window_start:
        window_end = min(float(t[-1]), window_start + 1.0)
    return window_start, window_end


def draw_plot_header(draw: ImageDraw.ImageDraw, fonts: FontSet, case: CaseRow, window_start: float, window_end: float) -> None:
    title = "%s | %.3f-%.3fs" % (Path(case.source_path).name, case.start_s, case.end_s)
    draw.text((22, 14), title, fill=(15, 23, 42), font=fonts.title)
    subtitle = "指标：%s    严重度：%s    展示窗口：%.1f-%.1fs" % (
        case.metric_name,
        case.severity,
        window_start,
        window_end,
    )
    draw.text((22, 46), subtitle, fill=(55, 70, 92), font=fonts.normal)
    if case.message:
        draw.text((22, 69), truncate_text(case.message, 130), fill=(78, 88, 105), font=fonts.small)


def build_panels(frame: SignalFrame, metric_name: str, mask: np.ndarray) -> List[PlotPanel]:
    metric = metric_name.lower()
    panels: List[PlotPanel] = []

    acc_lines = signal_lines(
        frame,
        mask,
        [
            ("纵向加速度", "acceleration", 1.0, (36, 104, 179), False),
        ],
    )
    if frame.has("acceleration"):
        jerk = derivative(frame.time_s.astype(float), time_window_moving_average(frame.time_s.astype(float), frame.col("acceleration").astype(float), 0.25))
        acc_lines.append(PlotLine("jerk", jerk[mask], (111, 66, 193)))
    if acc_lines:
        panels.append(PlotPanel("加速度 / jerk", "m/s2 / m/s3", acc_lines))

    speed_lines: List[PlotLine] = []
    if frame.has("ego_speed"):
        speed_lines.append(PlotLine("自车车速", frame.col("ego_speed").astype(float)[mask], (5, 150, 105), step=True))
    if frame.has("set_speed"):
        speed_lines.append(PlotLine("SetSpeed", frame.col("set_speed").astype(float)[mask], (107, 114, 128), step=True))
    if frame.has("curv_speed_limit"):
        speed_lines.append(PlotLine("弯道限速", frame.col("curv_speed_limit").astype(float)[mask] * speed_limit_scale(frame), (220, 38, 38), step=True))
    if speed_lines:
        panels.append(PlotPanel("车速 / 设速", "kph", speed_lines))

    if wants_front_panel(metric_name):
        front_lines = signal_lines(
            frame,
            mask,
            [
                ("前车距离", "front_range", 1.0, (15, 118, 110), False),
                ("TTC", "front_ttc", 1.0, (245, 158, 11), False),
                ("前车速度", "front_velocity_abs", 3.6, (124, 58, 237), True),
            ],
        )
        if front_lines:
            panels.append(PlotPanel("前车信号", "m / s / kph", front_lines))

    if wants_torque_panel(metric_name):
        torque_lines = signal_lines(
            frame,
            mask,
            [
                ("驱动请求扭矩", "control_drive_request_torque", 1.0, (22, 101, 52), False),
                ("实车驱动扭矩", "actual_drive_torque", 1.0, (34, 139, 34), False),
                ("制动请求扭矩", "control_brake_request_torque", 1.0, (37, 99, 235), False),
                ("实车制动扭矩", "actual_brake_torque", 1.0, (249, 115, 22), False),
            ],
        )
        if torque_lines:
            panels.append(PlotPanel("驱动 / 制动扭矩", "Nm", torque_lines))

    road_lines = signal_lines(
        frame,
        mask,
        [
            ("坡度", "slope_deg", 1.0, (217, 119, 6), False),
            ("曲率", "curvature", 1000.0, (190, 24, 93), False),
        ],
    )
    if wants_road_panel(metric_name) and road_lines:
        panels.append(PlotPanel("道路信号", "deg / curvature*1000", road_lines))

    state_lines = signal_lines(
        frame,
        mask,
        [
            ("ACC active", "acc_active_flag", 1.0, (22, 101, 52), True),
            ("override", "pedal_override_flag", 1.0, (220, 38, 38), True),
            ("brake type", "brake_request_type", 1.0, (37, 99, 235), True),
            ("ABS", "abs_active", 1.0, (190, 24, 93), True),
            ("TCS", "tcs_active", 1.0, (245, 158, 11), True),
            ("VSE", "vse_active", 1.0, (124, 58, 237), True),
            ("仲裁状态", "acc_request_state", 1.0, (20, 184, 166), True),
        ],
    )
    if state_lines:
        panels.append(PlotPanel("状态信号", "enum", state_lines, fixed_range=(-0.3, 6.3)))

    return prioritize_panels(metric_name, panels)


def wants_front_panel(metric_name: str) -> bool:
    text = metric_name.lower()
    tokens = ("follow", "front", "frnt", "ttc", "cutin", "brake_tap", "heavy_braking", "low_speed")
    return any(token in text for token in tokens) or text in {"点刹", "重制动", "不合理低速"}


def wants_torque_panel(metric_name: str) -> bool:
    text = metric_name.lower()
    tokens = ("torque", "jerkiness", "shake", "brake", "rollback", "override", "start_no_override", "点刹", "重制动", "顿挫")
    return any(token in text for token in tokens)


def wants_road_panel(metric_name: str) -> bool:
    text = metric_name.lower()
    return any(token in text for token in ("slope", "curv", "弯道", "坡道"))


def prioritize_panels(metric_name: str, panels: List[PlotPanel]) -> List[PlotPanel]:
    metric = metric_name.lower()
    if "curv" in metric:
        order = ["车速", "道路", "加速度", "前车", "状态", "驱动"]
    elif "slope" in metric:
        order = ["加速度", "车速", "道路", "状态", "前车", "驱动"]
    elif any(token in metric for token in ("brake", "shake", "jerk", "torque", "顿挫", "重制动", "点刹")):
        order = ["加速度", "驱动", "车速", "前车", "状态", "道路"]
    else:
        order = ["加速度", "车速", "前车", "驱动", "状态", "道路"]

    def rank(panel: PlotPanel) -> int:
        for index, token in enumerate(order):
            if panel.title.startswith(token) or token in panel.title:
                return index
        return len(order)

    return sorted(panels, key=rank)


def signal_lines(
    frame: SignalFrame,
    mask: np.ndarray,
    specs: Iterable[Tuple[str, str, float, Tuple[int, int, int], bool]],
) -> List[PlotLine]:
    lines: List[PlotLine] = []
    for label, field_name, scale, color, step in specs:
        if not frame.has(field_name):
            continue
        values = frame.col(field_name).astype(float) * float(scale)
        sliced = values[mask]
        finite = sliced[np.isfinite(sliced)]
        if finite.size == 0:
            continue
        if field_name in {"front_ttc", "curv_speed_limit"}:
            sliced = np.where(np.isfinite(sliced), np.minimum(sliced, 120.0), np.nan)
        lines.append(PlotLine(label, sliced, color, step))
    return lines


def speed_limit_scale(frame: SignalFrame) -> float:
    values = frame.col("curv_speed_limit", np.nan).astype(float)
    finite = values[np.isfinite(values)]
    if finite.size == 0:
        return 1.0
    return 3.6 if float(np.nanpercentile(finite, 95)) < 80.0 else 1.0


def draw_panel(
    draw: ImageDraw.ImageDraw,
    box: Tuple[int, int, int, int],
    panel: PlotPanel,
    t: np.ndarray,
    event_start: float,
    event_end: float,
    fonts: FontSet,
) -> None:
    x0, y0, x1, y1 = box
    draw.rectangle(box, fill=(255, 255, 255), outline=(209, 213, 219))
    draw.text((x0 + 10, y0 + 7), panel.title, fill=(17, 24, 39), font=fonts.header)
    plot_x0, plot_y0, plot_x1, plot_y1 = x0 + 92, y0 + 36, x1 - 18, y1 - 28
    ymin, ymax = panel.fixed_range or y_range(panel.lines)
    if math.isclose(ymin, ymax):
        ymax = ymin + 1.0

    shade_x0 = x_scale(event_start, t, plot_x0, plot_x1)
    shade_x1 = x_scale(event_end, t, plot_x0, plot_x1)
    draw.rectangle((shade_x0, plot_y0, max(shade_x1, shade_x0 + 1), plot_y1), fill=(254, 226, 226))
    draw.line((shade_x0, plot_y0, shade_x0, plot_y1), fill=(239, 68, 68), width=2)

    for index in range(4):
        y = plot_y1 - (plot_y1 - plot_y0) * index / 3
        value = ymin + (ymax - ymin) * index / 3
        draw.line((plot_x0, y, plot_x1, y), fill=(238, 242, 248))
        draw.text((x0 + 8, y - 7), format_tick(value), fill=(75, 85, 99), font=fonts.small)
    for index in range(5):
        x = plot_x0 + (plot_x1 - plot_x0) * index / 4
        time_value = float(t[0]) + (float(t[-1]) - float(t[0])) * index / 4
        draw.line((x, plot_y1, x, plot_y1 + 4), fill=(160, 170, 185))
        draw.text((x - 24, plot_y1 + 7), "%.1fs" % time_value, fill=(75, 85, 99), font=fonts.small)
    draw.line((plot_x0, plot_y1, plot_x1, plot_y1), fill=(160, 170, 185))
    draw.line((plot_x0, plot_y0, plot_x0, plot_y1), fill=(160, 170, 185))
    draw.text((x0 + 11, y0 + 30), panel.unit, fill=(75, 85, 99), font=fonts.small)

    for line in panel.lines:
        draw_series(draw, t, line, ymin, ymax, (plot_x0, plot_y0, plot_x1, plot_y1))

    legend_x, legend_y = plot_x0, y0 + 9
    for line in panel.lines:
        draw.line((legend_x, legend_y + 8, legend_x + 23, legend_y + 8), fill=line.color, width=4)
        draw.text((legend_x + 29, legend_y), line.label, fill=(40, 48, 62), font=fonts.small)
        legend_x += 31 + int(draw.textlength(line.label, font=fonts.small)) + 22
        if legend_x > x1 - 210:
            legend_x = plot_x0
            legend_y += 16


def draw_series(
    draw: ImageDraw.ImageDraw,
    t: np.ndarray,
    line: PlotLine,
    ymin: float,
    ymax: float,
    plot_box: Tuple[int, int, int, int],
) -> None:
    plot_x0, plot_y0, plot_x1, plot_y1 = plot_box
    points: List[Tuple[float, float]] = []
    for time_value, sample in zip(t, line.values):
        if not (np.isfinite(time_value) and np.isfinite(sample)):
            if len(points) >= 2:
                draw.line(step_points(points) if line.step else points, fill=line.color, width=2)
            points = []
            continue
        x = x_scale(float(time_value), t, plot_x0, plot_x1)
        clipped = min(max(float(sample), ymin), ymax)
        y = plot_y1 - (clipped - ymin) / (ymax - ymin) * (plot_y1 - plot_y0)
        points.append((x, y))
    if len(points) >= 2:
        draw.line(step_points(points) if line.step else points, fill=line.color, width=2)


def step_points(points: List[Tuple[float, float]]) -> List[Tuple[float, float]]:
    stepped: List[Tuple[float, float]] = []
    for index in range(len(points) - 1):
        stepped.append(points[index])
        stepped.append((points[index + 1][0], points[index][1]))
    stepped.append(points[-1])
    return stepped


def x_scale(value: float, t: np.ndarray, x0: int, x1: int) -> float:
    left, right = float(t[0]), float(t[-1])
    if right <= left:
        return float(x0)
    return x0 + (value - left) / (right - left) * (x1 - x0)


def y_range(lines: Sequence[PlotLine]) -> Tuple[float, float]:
    values: List[float] = []
    for line in lines:
        finite = np.asarray(line.values, dtype=float)
        finite = finite[np.isfinite(finite)]
        if finite.size:
            values.extend(finite.tolist())
    if not values:
        return -1.0, 1.0
    low = float(np.nanpercentile(values, 2))
    high = float(np.nanpercentile(values, 98))
    if math.isclose(low, high):
        pad = max(abs(low) * 0.1, 1.0)
        return low - pad, high + pad
    span = high - low
    return low - span * 0.12, high + span * 0.12


def format_tick(value: float) -> str:
    if abs(value) >= 100:
        return "%.0f" % value
    if abs(value) >= 10:
        return "%.1f" % value
    return "%.2f" % value


def truncate_text(text: str, max_len: int) -> str:
    return text if len(text) <= max_len else text[: max_len - 3] + "..."


def load_font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    candidates = [
        r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
    ]
    for path_text in candidates:
        path = Path(path_text)
        if path.exists():
            try:
                return ImageFont.truetype(str(path), size)
            except Exception:
                pass
    return ImageFont.load_default()


if __name__ == "__main__":
    raise SystemExit(main())
