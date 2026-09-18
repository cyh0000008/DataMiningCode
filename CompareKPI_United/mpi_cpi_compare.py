#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generate an MPI/CPI comparison workbook from base and target reports.

Supported input layouts:
  1. Latest unified reports:
       base/J6B/Unified_Longitudinal_Report_....xlsx
       target/J6B/Unified_Longitudinal_Report_....xlsx
       base/MMT/Unified_Longitudinal_Report_....xlsx
       target/MMT/Unified_Longitudinal_Report_....xlsx

     The summary sheet uses columns like:
       指标来源, 指标名称, 指标类型, 总里程(km), 总次数, 触发次数, MPI, CPI

  Old Longitudinal_CPI_Report workbooks are converted to the latest summary
  structure under .converted_reports/ before comparison.

  2. Legacy reports copied from CompareKPI:
       base/J6B/Topic_车型_开始日期_结束日期_软件版本.xlsx
       base/J6B/Can_车型_开始日期_结束日期_软件版本.xlsx
       target/J6B/Topic_车型_开始日期_结束日期_软件版本.xlsx
       target/J6B/Can_车型_开始日期_结束日期_软件版本.xlsx
"""

from __future__ import annotations

import argparse
import json
import math
import re
import shutil
import sys
import unicodedata
from copy import copy
from dataclasses import dataclass, replace
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


SUMMARY_SHEET = "Sheet1"
MAPPING_SHEET = "信号对照表"
LEGACY_REPORT_PREFIX = "Longitudinal_CPI_Report"
LATEST_REPORT_PREFIX = "Unified_Longitudinal_Report"
REPORT_PREFIXES = (LATEST_REPORT_PREFIX, LEGACY_REPORT_PREFIX)
CONVERTED_REPORT_DIR = ".converted_reports"
DEFAULT_REPORT_DIR = "reports"
DEFAULT_TEMPLATE_DIR = "templates"

BASE_MPI_COL = "E"
BASE_CPI_COL = "F"
TARGET_MPI_COL = "G"
TARGET_CPI_COL = "H"
MISSING_DISPLAY = "\\"
REPORT_FONT_NAME = "微软雅黑"
DEFAULT_FONT_COLOR = "000000"
DEFAULT_ROW_HEIGHT = 18.0
MAX_ROW_HEIGHT = 160.0
ROW_HEIGHT_PER_LINE = 16.5

GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=REPORT_FONT_NAME, bold=True, color="FFFFFF")


# 中文指标名 -> (legacy 结果文件类型, 英文指标名, 填入模板的指标类型列)
# unified 格式会忽略 legacy 类型，只按英文指标名 + MPI/CPI 查找。
SIGNAL_MAP: dict[str, tuple[str, str, str]] = {
    "起步最大加速度": ("topic", "stop2go_max_acc", "CPI"),
    "起步最大jerk": ("topic", "stop2go_max_jerk", "CPI"),
    "初段建立时间": ("topic", "stop2go_start_time", "CPI"),
    "起步加速建立时间": ("topic", "stop2go_max_acc_time", "CPI"),
    "弯道安全速度": ("topic", "curvature_min_speed", "CPI"),
    "弯道限速最大减速度": ("topic", "curvature_min_acc", "CPI"),
    "弯道限速减速舒适性": ("topic", "curvature_min_jerk", "CPI"),
    "坡道稳速抖动": ("topic", "slope_no_frnt_vel_fluct", "CPI"),
    "坡道控速稳定性": ("topic", "slope_no_frnt_acc", "CPI"),
    "坡道控速舒适性": ("topic", "slope_no_frnt_jerk", "CPI"),
    "跟车刹停最大减速度": ("topic", "follow_stop_max_acc", "CPI"),
    "跟车刹停距离": ("topic", "follow_stop_min_frntobj_dis", "CPI"),
    "跟车刹停舒适性": ("topic", "follow_stop_min_jerk", "CPI"),
    "跟车刹停波动": ("topic", "follow_stop_fluct", "CPI"),
    "跟车刹停反应": ("topic", "follow_stop_ttc", "CPI"),
    "跟车刹停风险": ("topic", "follow_stop_min_ttc_frntobj_dec", "CPI"),
    "跟车起步响应时延": ("topic", "follow_start_latency", "CPI"),
    "跟车起步舒适性": ("topic", "follow_start_jerk", "CPI"),
    "红灯刹停距停止线距离": ("topic", "redlight_stop_dis", "CPI"),
    "红灯刹停舒适性": ("topic", "redlight_min_jerk", "CPI"),
    "绿灯起步响应时延": ("topic", "traffic_light_stop2go_time", "CPI"),
    "绿灯起步最大加速度": ("topic", "traffic_light_stop2go_max_accel", "CPI"),
    "绿灯起步加速时间": ("topic", "traffic_light_stop2go_accel_time", "CPI"),
    "绿灯起步舒适性": ("topic", "traffic_light_stop2go_max_jerk", "CPI"),
    "溜车": ("can", "acc_rollback_count", "MPI"),
    "无前车稳速未达到SetSpeed": ("can", "acc_set_speed_reach", "MPI"),
    "不合理低速": ("can", "acc_unreasonable_low_speed_count", "MPI"),
    "制动跳变导致的实车加速度抖动": ("can", "acc_brake_rapid_toggle_decel_jump_count", "MPI"),
    "点刹": ("can", "acc_brake_tap_count", "MPI"),
    "重制动": ("can", "acc_heavy_braking_count", "MPI"),
    "顿挫": ("can", "acc_jerkiness", "MPI"),
    "控制抖动": ("can", "control_shake_count", "MPI"),
    "Override后顿挫": ("can", "acc_jerkiness_after_override_release", "MPI"),
    "Override后重制动": ("can", "acc_heavy_braking_after_override_release_count", "MPI"),
    "行车断握手": ("can", "acc_lost_arbitration", "MPI"),
    "进Override后顿挫": ("can", "acc_override_brake_request_jerkiness", "MPI"),
    "起步顿挫": ("can", "acc_start_no_override_jerkiness", "MPI"),
}


REPORT_PROFILE_J6B = "J6B"
REPORT_PROFILE_MMT = "MMT"
REPORT_PROFILES = (REPORT_PROFILE_J6B, REPORT_PROFILE_MMT)

# J6B 保留当前 J6B 配置需要的指标；MMT 使用完整模板。
J6B_METRIC_NAMES = frozenset(
    {
        "起步最大加速度",
        "起步最大jerk",
        "初段建立时间",
        "起步加速建立时间",
        "弯道安全速度",
        "弯道限速最大减速度",
        "弯道限速减速舒适性",
        "坡道稳速抖动",
        "坡道控速稳定性",
        "坡道控速舒适性",
        "跟车刹停最大减速度",
        "跟车刹停距离",
        "跟车刹停舒适性",
        "跟车刹停波动",
        "跟车刹停反应",
        "跟车刹停风险",
        "跟车起步响应时延",
        "跟车起步舒适性",
        "溜车",
        "无前车稳速未达到SetSpeed",
        "不合理低速",
        "制动跳变导致的实车加速度抖动",
        "点刹",
        "重制动",
        "顿挫",
        "控制抖动",
        "Override后顿挫",
        "Override后重制动",
        "行车断握手",
        "进Override后顿挫",
        "起步顿挫",
    }
)

PROFILE_SIGNAL_MAP_OVERRIDES: dict[str, dict[str, tuple[str, str, str]]] = {
    REPORT_PROFILE_J6B: {
        "顿挫": ("can", "acc_torque_jerkiness_count", "MPI"),
    },
}


METRIC_STANDARDS: dict[str, str] = {
    "起步最大加速度": "1.0~3.0m/s2",
    "起步最大jerk": "<=3.5m/s3",
    "初段建立时间": "达到3.33m/s所需时间：1.0~3.5s",
    "起步加速建立时间": "首次acc>0.2到最大acc：1.0~4.0s",
    "弯道安全速度": "sqrt(1.0*R)-2 ~ sqrt(2.9*R)+2m/s；SetSpeed<=sqrt(1.0*R)不判bad",
    "弯道限速最大减速度": ">=-2m/s2",
    "弯道限速减速舒适性": ">=-2m/s3",
    "坡道稳速抖动": "速度波动<=1m/s",
    "坡道控速稳定性": "|加速度|<=0.75m/s2",
    "坡道控速舒适性": "|jerk|<=1m/s3",
    "跟车刹停最大减速度": ">=max(min(-初始速度(m/s)/5.5,-1.75),-4)-1m/s2",
    "跟车刹停距离": "2~5m",
    "跟车刹停舒适性": ">=-2m/s3",
    "跟车刹停反应": "TTC>=2.5s",
    "跟车刹停风险": "TTC>=1.5s",
    "跟车刹停波动": "<=3次",
    "跟车起步响应时延": "<=2.5s",
    "跟车起步舒适性": "1.0~2.5m/s3",
    "红灯刹停距停止线距离": "0~4m",
    "红灯刹停舒适性": ">=-2m/s3",
    "绿灯起步响应时延": "<=2.0s",
    "绿灯起步最大加速度": "1.0~2.3m/s2",
    "绿灯起步加速时间": "<=6.0s",
    "绿灯起步舒适性": "1.0~2.5m/s3",
    "溜车": "ACC active/hold/前进挡，车速<=0.5kph且轮速非零，持续>=0.20s",
    "无前车稳速未达到SetSpeed": "ACC active且无override，场景>=8s、车速波动<=2kph，未达到SetSpeed",
    "不合理低速": "ACC active、无override、无红灯停车/前车风险；6s内v_max-v_min<=2kph，v_min/v_max>5kph，set_speed-v_max>=5kph",
    "制动跳变导致的实车加速度抖动": "ACC active且无override，AutoBrake 2->1->2中间态<=0.10s，随后1s减速度波动>2.0m/s2",
    "点刹": "ACC active且无override，无前车稳速或有前车稳定跟车中出现制动点刹；制动扭矩短脉冲0.08~1.50s，峰值突出>=180Nm，加速度下探>=0.2m/s2",
    "重制动": "加速度<-1.8m/s2，减速度跳变量>=2.0m/s2，斜率>=1.0，持续>=0.20s，并满足jerk/深踩补充规则",
    "顿挫": "ACC active且无override，平滑加速度振荡幅值>=0.75、符号翻转>=1、时长1~4s，并通过驱动/制动扭矩校验",
    "Override后顿挫": "override释放后5s内，按顿挫规则",
    "Override后重制动": "override释放后5s内，加速度<-2.0m/s2且jerk<-3.0m/s3，持续>=0.20s",
    "行车断握手": "ACC active且仲裁状态=2，持续>=5s",
    "进Override后顿挫": "ACC active且override中制动请求<-0.2m/s2，按顿挫规则",
    "起步顿挫": "起步1s窗口内，按顿挫规则",
}

# J6B 专用标准只覆盖当前模板已有指标；不额外引入 J6B_UnifiedMining
# filters 中存在但模板未放入的指标。
J6B_METRIC_STANDARDS: dict[str, str] = {
    "弯道限速最大减速度": ">=-2.5m/s2",
    "弯道限速减速舒适性": ">=-2.5m/s3",
    "坡道稳速抖动": "速度波动<=1m/s",
    "坡道控速稳定性": "|加速度|<=1.0m/s2",
    "坡道控速舒适性": "|jerk|<=1.5m/s3",
    "跟车刹停舒适性": ">=-2.5m/s3",
    "跟车刹停波动": "<=3次",
    "跟车刹停风险": "TTC>=1.5s",
    "溜车": "ACC active，前进挡，车速<=0.5kph，轮速>=0.01，持续>=0.20s",
    "无前车稳速未达到SetSpeed": "ACC active且无override/前车/cutin，稳定>=8s、车速波动<=2kph，未达到SetSpeed",
    "不合理低速": "ACC active且无override/前车/cutin/曲率限速，稳定>=6s、车速波动<=2kph，车速>=5kph，低于SetSpeed>=5kph",
    "点刹": "ACC active且无override；无前车稳速或稳定跟车中制动点刹。点刹：制动扭矩短脉冲0.08~1.50s、峰值突出>=180Nm、加速度下探>=0.2m/s2；无前车稳速：无前车/cutin、车速>5kph、稳定>=8s、车速波动<=2kph、SetSpeed波动<=0.1kph；稳定跟车：前车有效、距离5~100m、相对速度<=1.5m/s、稳定>=8s、车速波动<=5kph、前车距离波动<=10m",
    "重制动": "ACC active且无override，加速度<-1.8m/s2，减速度跳变>=2.0m/s2，斜率>=1.0，持续>=0.20s",
    "顿挫": "ACC active且无override，车速>0，顿挫1~4s，幅值>=0.75，符号翻转>=1，扭矩/jerk校验通过",
    "控制抖动": "全程ACC active且车速>5kph，8s内驱动/制动扭矩抖动；驱动类acc_amp>=0.8m/s2，制动类acc_amp>=1.45m/s2",
    "进Override后顿挫": "ACC active且override中，制动请求<-0.2m/s2，油门>0，按顿挫规则",
}

PROFILE_METRIC_STANDARD_OVERRIDES: dict[str, dict[str, str]] = {
    REPORT_PROFILE_J6B: J6B_METRIC_STANDARDS,
}


def resolve_signal_mapping(cn_name: str, report_profile: str) -> tuple[str, str, str] | None:
    normalized_name = normalize(cn_name)
    profile_overrides = PROFILE_SIGNAL_MAP_OVERRIDES.get(report_profile, {})
    return profile_overrides.get(normalized_name) or SIGNAL_MAP.get(normalized_name)


def resolve_metric_standard(cn_name: str, report_profile: str) -> str | None:
    normalized_name = normalize(cn_name)
    profile_standard = PROFILE_METRIC_STANDARD_OVERRIDES.get(report_profile, {}).get(normalized_name)
    if profile_standard is not None:
        return profile_standard
    return METRIC_STANDARDS.get(normalized_name)


@dataclass(frozen=True)
class ResultFile:
    side: str
    kind: str
    vehicle: str
    start_date: str
    end_date: str
    version: str
    path: Path

    @property
    def date_range(self) -> str:
        if self.start_date and self.end_date and self.start_date != self.end_date:
            return f"{self.start_date}-{self.end_date}"
        return self.start_date or self.end_date


@dataclass
class MetricValue:
    raw: Any = None
    number: float | None = None
    bad_case: float | None = None
    source_column: str = ""
    source_file: str = ""
    total: Any = None
    bad: Any = None


def normalize(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value).strip())


def parse_date(token: str) -> str:
    digits = re.sub(r"\D", "", token)
    if len(digits) == 8:
        return digits
    return token


def format_date_token(match: re.Match[str]) -> str:
    return f"{match.group(1)}{match.group(2)}{match.group(3)}"


def extract_dates_from_text(text: str) -> list[str]:
    dates: list[str] = []
    patterns = (
        re.compile(r"(20\d{2})[-_]?(\d{2})[-_]?(\d{2})"),
        re.compile(r"(20\d{2})(\d{2})(\d{2})"),
    )
    for pattern in patterns:
        for match in pattern.finditer(text):
            value = format_date_token(match)
            if value not in dates:
                dates.append(value)
    return dates


def normalize_date_range_label(value: Any) -> str:
    dates = extract_dates_from_text(str(value or ""))
    if len(dates) >= 2:
        return f"{dates[0]}-{dates[1]}"
    if len(dates) == 1:
        return dates[0]
    return normalize(value)


def parse_legacy_filename(path: Path, side: str) -> ResultFile | None:
    parts = [part for part in path.stem.split("_") if part]
    if len(parts) < 5:
        return None
    kind = parts[0].lower()
    if kind not in {"topic", "can"}:
        return None
    return ResultFile(
        side=side,
        kind=kind,
        vehicle=parts[1],
        start_date=parse_date(parts[2]),
        end_date=parse_date(parts[3]),
        version="_".join(parts[4:]),
        path=path,
    )


def parse_named_report_filename(path: Path, side: str) -> ResultFile | None:
    parts = [part for part in path.stem.split("_") if part]
    if len(parts) < 4 or parts[0].lower() in {"topic", "can"}:
        return None

    date_index = None
    for idx in range(1, len(parts) - 2):
        start_date = parse_date(parts[idx])
        end_date = parse_date(parts[idx + 1])
        if re.fullmatch(r"20\d{6}", start_date) and re.fullmatch(r"20\d{6}", end_date):
            date_index = idx
            break

    if date_index is None:
        return None

    return ResultFile(
        side=side,
        kind="unified",
        vehicle="_".join(parts[:date_index]),
        start_date=parse_date(parts[date_index]),
        end_date=parse_date(parts[date_index + 1]),
        version="_".join(parts[date_index + 2:]),
        path=path,
    )


def load_manifest(report_path: Path) -> dict[str, Any]:
    manifest_path = report_path.with_name("run_manifest.json")
    if not manifest_path.exists():
        return {}
    try:
        with manifest_path.open("r", encoding="utf-8") as fh:
            data = json.load(fh)
        return data if isinstance(data, dict) else {}
    except (OSError, json.JSONDecodeError):
        return {}


def infer_vehicle(report_path: Path, manifest: dict[str, Any], side: str) -> str:
    for text in manifest.get("input_paths", []):
        match = re.search(r"([A-Za-z]\d+[A-Za-z]*)[_-]?Data", str(text))
        if match:
            return match.group(1)

    skip_names = {
        "base",
        "target",
        "dataminingresult",
        "comparekpi_united",
        "comparekpi",
    }
    for parent in [report_path.parent, *report_path.parents]:
        name = parent.name.strip()
        if not name or name.lower() in skip_names or any(name.startswith(prefix) for prefix in REPORT_PREFIXES):
            continue
        match = re.match(r"([A-Za-z]+\d+[A-Za-z0-9]*)", name)
        if match:
            return match.group(1)
    return side


def infer_dates(report_path: Path, manifest: dict[str, Any]) -> tuple[str, str]:
    dates: list[str] = []
    for item in manifest.get("input_paths", []):
        for date in extract_dates_from_text(str(item)):
            if date not in dates:
                dates.append(date)
    if not dates:
        dates = extract_dates_from_text(report_path.stem)
    if not dates:
        return "", ""
    dates = sorted(dates)
    return dates[0], dates[-1]


def infer_version(report_path: Path, manifest: dict[str, Any]) -> str:
    created_at = manifest.get("created_at")
    if isinstance(created_at, str) and created_at:
        try:
            return datetime.fromisoformat(created_at).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            return created_at

    for prefix in REPORT_PREFIXES:
        match = re.search(
            rf"{re.escape(prefix)}_(\d{{4}}-\d{{2}}-\d{{2}})_(\d{{2}}-\d{{2}}-\d{{2}})",
            report_path.stem,
        )
        if match:
            return f"{match.group(1)} {match.group(2).replace('-', ':')}"
    return report_path.stem


def make_unified_result(side: str, report_path: Path) -> ResultFile:
    filename_info = parse_named_report_filename(report_path, side)
    if filename_info is not None:
        return filename_info

    manifest = load_manifest(report_path)
    start_date, end_date = infer_dates(report_path, manifest)
    return ResultFile(
        side=side,
        kind="unified",
        vehicle=infer_vehicle(report_path, manifest, side),
        start_date=start_date,
        end_date=end_date,
        version=infer_version(report_path, manifest),
        path=report_path,
    )


def is_candidate_report(path: Path) -> bool:
    if path.suffix.lower() != ".xlsx":
        return False
    if path.name.startswith("~$"):
        return False
    if path.name.lower() == "bad_case_report.xlsx":
        return False
    return True


def report_sort_key(path: Path) -> tuple[int, float, str]:
    if parse_named_report_filename(path, "") is not None:
        preferred = 0
    elif any(path.stem.startswith(prefix) for prefix in REPORT_PREFIXES):
        preferred = 1
    else:
        preferred = 2
    try:
        mtime = path.stat().st_mtime
    except OSError:
        mtime = 0
    return preferred, -mtime, path.name.lower()


def resolve_unified_report_path(folder_or_file: Path) -> Path:
    if folder_or_file.is_file():
        if not is_candidate_report(folder_or_file):
            raise ValueError(f"不是可读取的汇总报告：{folder_or_file}")
        return folder_or_file
    if not folder_or_file.exists():
        raise FileNotFoundError(f"未找到路径：{folder_or_file}")

    direct = [path for path in folder_or_file.glob("*.xlsx") if is_candidate_report(path)]
    nested = [
        path
        for prefix in REPORT_PREFIXES
        for path in folder_or_file.glob(f"{prefix}*/*.xlsx")
        if is_candidate_report(path)
    ]
    candidates = direct + nested
    if not candidates:
        raise FileNotFoundError(f"{folder_or_file} 中未找到 unified 汇总报告")
    return sorted(candidates, key=report_sort_key)[0]


def find_legacy_reports(folder: Path, side: str) -> dict[str, ResultFile]:
    reports: dict[str, ResultFile] = {}
    if not folder.exists() or not folder.is_dir():
        return reports
    for path in folder.glob("*.xlsx"):
        if path.name.startswith("~$"):
            continue
        info = parse_legacy_filename(path, side)
        if info:
            reports.setdefault(info.kind, info)
    return reports


def find_side_reports(
    root: Path,
    side: str,
    input_path: Path | None,
    required: bool = True,
) -> dict[str, ResultFile]:
    folder_or_file = input_path if input_path is not None else root / side
    if not folder_or_file.is_absolute():
        folder_or_file = root / folder_or_file
    folder_or_file = folder_or_file.resolve()

    legacy_reports = find_legacy_reports(folder_or_file, side)
    if {"topic", "can"}.issubset(legacy_reports):
        return legacy_reports

    try:
        report_path = resolve_unified_report_path(folder_or_file)
    except FileNotFoundError:
        if not required:
            return {}
        raise
    return {"unified": make_unified_result(side, report_path)}


def profile_dir_name(report_profile: str) -> str:
    return report_profile.replace("+", "P").replace(" ", "_")


def default_side_input_path(root: Path, side: str, report_profile: str) -> Path:
    return root / side / profile_dir_name(report_profile)


def default_template_path(root: Path, report_profile: str) -> Path:
    return root / DEFAULT_TEMPLATE_DIR / f"template_{profile_dir_name(report_profile)}.xlsx"


def parse_metric_number(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    match = re.search(r"-?\d+(?:\.\d+)?", text.replace(",", ""))
    if not match:
        return None
    return float(match.group(0))


def parse_bad_case_from_metric(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    text = (
        text.replace("（", "(")
        .replace("）", ")")
        .replace("／", "/")
        .replace(",", "")
    )
    match = re.search(r"\(([^()]*)\)", text)
    if not match:
        return None
    parts = match.group(1).split("/")
    if len(parts) < 2:
        return None
    bad_match = re.search(r"-?\d+(?:\.\d+)?", parts[-1].strip())
    if not bad_match:
        return None
    return float(bad_match.group(0))


def display_raw(value: MetricValue) -> Any:
    return value.raw if value.raw not in (None, "") else MISSING_DISPLAY


def pick_summary_sheet(wb):
    for name in ("所有指标统计汇总", "汇总", "Summary", "summary"):
        if name in wb.sheetnames:
            return wb[name]
    return wb.worksheets[0]


def get_header_col(headers: dict[str, int], candidates: tuple[str, ...]) -> int | None:
    for candidate in candidates:
        key = normalize(candidate).lower()
        if key in headers:
            return headers[key]
    return None


def find_header(ws) -> tuple[int, dict[str, int]]:
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        headers: dict[str, int] = {}
        for col_idx, value in enumerate(row, start=1):
            key = normalize(value).lower()
            if key:
                headers[key] = col_idx
        if get_header_col(headers, ("指标名称", "metric", "signal")) is not None:
            return row_idx, headers
    raise ValueError(f"{ws.title} 未找到“指标名称”表头")


def is_latest_summary_headers(headers: dict[str, int]) -> bool:
    required = ("指标来源", "指标名称", "指标类型", "总里程(km)", "总次数", "触发次数")
    return all(normalize(name).lower() in headers for name in required) and (
        normalize("MPI").lower() in headers or normalize("CPI").lower() in headers
    )


def is_old_value_summary_headers(headers: dict[str, int]) -> bool:
    required = ("指标类型", "指标名称", "指标值")
    return all(normalize(name).lower() in headers for name in required) and not is_latest_summary_headers(headers)


def needs_latest_report_conversion(path: Path) -> bool:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = pick_summary_sheet(wb)
        _, headers = find_header(ws)
        return is_old_value_summary_headers(headers)
    finally:
        wb.close()


def metric_source_from_type(metric_type: str) -> str:
    return "can" if metric_type.upper() == "MPI" else "topic"


def parse_total_distance_from_metric(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    text = (
        text.replace("（", "(")
        .replace("）", ")")
        .replace("／", "/")
        .replace(",", "")
    )
    match = re.search(r"\(([^()/]+)/(?:[^()]*)\)", text)
    if not match:
        return None
    try:
        return float(match.group(1).strip())
    except ValueError:
        return None


def parse_metric_parenthetical(value: Any) -> tuple[float | None, float | None]:
    if value is None:
        return None, None
    text = str(value).strip()
    if not text:
        return None, None
    text = (
        text.replace("（", "(")
        .replace("）", ")")
        .replace("／", "/")
        .replace(",", "")
    )
    match = re.search(r"\(([^()]*)\)", text)
    if not match:
        return None, None
    parts = match.group(1).split("/")
    if len(parts) < 2:
        return None, None
    try:
        total = float(parts[0].strip())
    except ValueError:
        total = None
    try:
        bad = float(parts[-1].strip())
    except ValueError:
        bad = None
    return total, bad


def find_comparison_date_groups(ws) -> list[dict[str, Any]]:
    groups: list[dict[str, Any]] = []
    for col in range(1, max(1, ws.max_column)):
        mpi_label = normalize(ws.cell(4, col).value).upper()
        cpi_label = normalize(ws.cell(4, col + 1).value).upper()
        if (mpi_label, cpi_label) != ("MPI", "CPI"):
            continue
        date_range = normalize_date_range_label(ws.cell(2, col).value)
        if not date_range:
            continue
        start_date, end_date = split_date_range(date_range)
        groups.append(
            {
                "start_col": col,
                "vehicle": ws.cell(1, col).value or "",
                "date_range": f"{start_date}-{end_date}" if start_date and end_date else date_range,
                "start_date": start_date,
                "end_date": end_date,
                "version": ws.cell(3, col).value or "",
            }
        )
    return groups


def select_comparison_date_group(ws, requested_date_range: str | None) -> dict[str, Any] | None:
    groups = find_comparison_date_groups(ws)
    if not groups:
        return None

    if requested_date_range:
        requested = normalize_date_range_label(requested_date_range)
        for group in groups:
            if normalize_date_range_label(group["date_range"]) == requested:
                return group
        raise ValueError(
            f"{ws.parent.path if hasattr(ws.parent, 'path') else ws.title} 中未找到日期组：{requested_date_range}"
        )

    return sorted(groups, key=lambda item: (item["end_date"], item["start_date"], item["start_col"]))[-1]


def convert_comparison_report_to_latest(
    input_path: Path,
    output_path: Path,
    requested_date_range: str | None,
    report_profile: str,
) -> tuple[Path, ResultFile] | None:
    src_wb = load_workbook(input_path, data_only=True)
    try:
        if SUMMARY_SHEET in src_wb.sheetnames:
            ws = src_wb[SUMMARY_SHEET]
        else:
            ws = src_wb.worksheets[0]
        group = select_comparison_date_group(ws, requested_date_range)
        if group is None:
            return None

        output_path.parent.mkdir(parents=True, exist_ok=True)
        out_wb = Workbook()
        try:
            out_ws = out_wb.active
            out_ws.title = "所有指标统计汇总"
            latest_headers = ["指标来源", "指标名称", "指标类型", "总里程(km)", "总次数", "触发次数", "MPI", "CPI"]
            out_ws.append(latest_headers)
            for cell in out_ws[1]:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT

            for row in range(6, ws.max_row + 1):
                cn_name = normalize(ws.cell(row, 2).value)
                mapping = resolve_signal_mapping(cn_name, report_profile)
                if not cn_name or mapping is None:
                    continue
                kind, english_name, _ = mapping
                for metric_type, source_col in (("MPI", group["start_col"]), ("CPI", group["start_col"] + 1)):
                    raw_value = ws.cell(row, source_col).value
                    if raw_value in (None, "", MISSING_DISPLAY):
                        continue
                    total_value, bad_value = parse_metric_parenthetical(raw_value)
                    out_ws.append(
                        [
                            kind,
                            english_name,
                            metric_type,
                            total_value if metric_type == "MPI" else None,
                            total_value,
                            bad_value,
                            raw_value if metric_type == "MPI" else None,
                            raw_value if metric_type == "CPI" else None,
                        ]
                    )

            for idx, width in enumerate([12, 38, 12, 14, 12, 12, 20, 20], start=1):
                out_ws.column_dimensions[get_column_letter(idx)].width = width
            out_ws.freeze_panes = "A2"
            out_ws.auto_filter.ref = out_ws.dimensions
            out_wb.save(output_path)
        finally:
            out_wb.close()

        info = ResultFile(
            side="",
            kind="unified",
            vehicle=str(group["vehicle"]),
            start_date=str(group["start_date"]),
            end_date=str(group["end_date"]),
            version=str(group["version"]),
            path=output_path,
        )
        return output_path, info
    finally:
        src_wb.close()


def safe_sheet_title(title: str, used: set[str]) -> str:
    clean = re.sub(r"[\[\]:*?/\\]", "_", title)[:31]
    if not clean:
        clean = "Sheet"
    candidate = clean
    suffix = 1
    while candidate in used:
        suffix_text = f"_{suffix}"
        candidate = f"{clean[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1
    used.add(candidate)
    return candidate


def copy_cell_style(source, target) -> None:
    if source.has_style:
        target._style = copy(source._style)
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def convert_old_report_to_latest(input_path: Path, output_path: Path) -> Path:
    if output_path.exists() and output_path.stat().st_mtime >= input_path.stat().st_mtime:
        return output_path

    output_path.parent.mkdir(parents=True, exist_ok=True)
    src_wb = load_workbook(input_path)
    try:
        summary_ws = pick_summary_sheet(src_wb)
        header_row, headers = find_header(summary_ws)
        if not is_old_value_summary_headers(headers):
            return input_path

        name_col = get_header_col(headers, ("指标名称", "metric", "signal"))
        type_col = get_header_col(headers, ("指标类型", "category", "type"))
        total_col = get_header_col(headers, ("总Case之和", "总次数", "total", "总里程(km)"))
        bad_col = get_header_col(headers, ("Bad Case之和", "触发次数", "bad", "badcase"))
        value_col = get_header_col(headers, ("指标值", "value", "metricvalue"))
        if name_col is None or type_col is None or value_col is None:
            return input_path

        out_wb = Workbook()
        try:
            out_ws = out_wb.active
            out_ws.title = "所有指标统计汇总"
            latest_headers = ["指标来源", "指标名称", "指标类型", "总里程(km)", "总次数", "触发次数", "MPI", "CPI"]
            out_ws.append(latest_headers)
            for cell in out_ws[1]:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT

            for source_row in summary_ws.iter_rows(min_row=header_row + 1):
                metric_name = normalize(source_row[name_col - 1].value)
                metric_type = normalize(source_row[type_col - 1].value).upper()
                if not metric_name or metric_type not in {"MPI", "CPI"}:
                    continue

                raw_value = source_row[value_col - 1].value
                total_value = source_row[total_col - 1].value if total_col else None
                bad_value = source_row[bad_col - 1].value if bad_col else None
                out_ws.append(
                    [
                        metric_source_from_type(metric_type),
                        metric_name,
                        metric_type,
                        parse_total_distance_from_metric(raw_value),
                        total_value,
                        bad_value,
                        raw_value if metric_type == "MPI" else None,
                        raw_value if metric_type == "CPI" else None,
                    ]
                )

            for idx, width in enumerate([12, 38, 12, 14, 12, 12, 20, 20], start=1):
                out_ws.column_dimensions[get_column_letter(idx)].width = width
            out_ws.freeze_panes = "A2"
            out_ws.auto_filter.ref = out_ws.dimensions

            used_titles = {"所有指标统计汇总"}
            for src_ws in src_wb.worksheets:
                if src_ws.title == summary_ws.title:
                    continue
                parts = src_ws.title.split("_", 1)
                if len(parts) == 2 and parts[0].upper() in {"MPI", "CPI"}:
                    prefix = "can" if parts[0].upper() == "MPI" else "topic"
                    new_title = safe_sheet_title(f"{prefix}__{parts[1]}", used_titles)
                else:
                    new_title = safe_sheet_title(src_ws.title, used_titles)
                dst_ws = out_wb.create_sheet(new_title)
                skip_first_col = normalize(src_ws.cell(1, 1).value) == normalize("指标类型")
                for row in src_ws.iter_rows():
                    for cell in row:
                        if skip_first_col and cell.column == 1:
                            continue
                        target_col = cell.column - 1 if skip_first_col else cell.column
                        dst_cell = dst_ws.cell(cell.row, target_col)
                        dst_cell.value = cell.value
                        copy_cell_style(cell, dst_cell)
                for col_idx, col_dim in src_ws.column_dimensions.items():
                    if skip_first_col and col_idx == "A":
                        continue
                    target_idx = get_column_letter(max(1, src_ws[col_idx + "1"].column - (1 if skip_first_col else 0)))
                    dst_ws.column_dimensions[target_idx].width = col_dim.width

            out_wb.save(output_path)
        finally:
            out_wb.close()
    finally:
        src_wb.close()

    return output_path


def prepare_latest_reports(
    root: Path,
    report_profile: str,
    side: str,
    reports: dict[str, ResultFile],
    source_date_range: str | None = None,
) -> dict[str, ResultFile]:
    prepared: dict[str, ResultFile] = {}
    converted_root = root / CONVERTED_REPORT_DIR / profile_dir_name(report_profile) / side
    for kind, info in reports.items():
        if kind != "unified":
            prepared[kind] = info
            continue
        comparison_date_token = normalize_date_range_label(source_date_range or "").replace("-", "_")
        if comparison_date_token:
            comparison_output = (
                converted_root / f"{info.path.stem}_{comparison_date_token}_latest.xlsx"
            )
        else:
            comparison_output = converted_root / f"{info.path.stem}_latest.xlsx"
        comparison_result = convert_comparison_report_to_latest(
            info.path, comparison_output, source_date_range, report_profile
        )
        if comparison_result is not None:
            converted_path, selected_info = comparison_result
            prepared[kind] = replace(selected_info, side=info.side, path=converted_path)
            continue
        if needs_latest_report_conversion(info.path):
            output_path = converted_root / info.path.name
            converted_path = convert_old_report_to_latest(info.path, output_path)
            prepared[kind] = replace(info, path=converted_path)
        else:
            prepared[kind] = info
    return prepared


def row_value(row: tuple[Any, ...], col_idx: int | None) -> Any:
    if col_idx is None:
        return None
    if col_idx - 1 >= len(row):
        return None
    return row[col_idx - 1]


def make_metric_value(
    raw: Any,
    source_col: int,
    info: ResultFile,
    total_value: Any,
    bad_value: Any,
) -> MetricValue:
    parsed_bad = parse_bad_case_from_metric(raw)
    return MetricValue(
        raw=raw,
        number=parse_metric_number(raw),
        bad_case=parsed_bad if parsed_bad is not None else parse_metric_number(bad_value),
        source_column=get_column_letter(source_col),
        source_file=info.path.name,
        total=total_value,
        bad=bad_value,
    )


def load_metric_values(info: ResultFile) -> dict[str, dict[str, MetricValue]]:
    wb = load_workbook(info.path, data_only=True, read_only=True)
    ws = pick_summary_sheet(wb)
    header_row, headers = find_header(ws)

    name_col = get_header_col(headers, ("指标名称", "metric", "signal"))
    source_kind_col = get_header_col(headers, ("指标来源", "source", "sourcekind", "kind"))
    type_col = get_header_col(headers, ("指标类型", "category", "type"))
    value_col = get_header_col(headers, ("指标值", "value", "metricvalue"))
    total_col = get_header_col(headers, ("总Case之和", "总次数", "total", "总里程(km)"))
    bad_col = get_header_col(headers, ("Bad Case之和", "触发次数", "bad", "badcase"))
    mpi_col = get_header_col(headers, ("MPI",))
    cpi_col = get_header_col(headers, ("CPI",))

    if name_col is None:
        wb.close()
        raise ValueError(f"{info.path} 未找到“指标名称”列")
    if value_col is None and mpi_col is None and cpi_col is None:
        wb.close()
        raise ValueError(f"{info.path} 未找到“指标值”或 MPI/CPI 列")

    values: dict[str, dict[str, MetricValue]] = {}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        metric = normalize(row_value(row, name_col))
        if not metric:
            continue
        source_kind = normalize(row_value(row, source_kind_col)).lower()
        metric_keys: list[tuple[str, bool]] = []
        if source_kind:
            metric_keys.append((f"{source_kind}::{metric}", True))
        metric_keys.append((metric, not source_kind))
        total_value = row_value(row, total_col)
        bad_value = row_value(row, bad_col)

        def store_metric_value(metric_type: str, metric_value: MetricValue) -> None:
            for metric_key, allow_overwrite in metric_keys:
                metric_values = values.setdefault(metric_key, {})
                if allow_overwrite or metric_type not in metric_values:
                    metric_values[metric_type] = metric_value

        if value_col is not None and type_col is not None:
            metric_type = normalize(row_value(row, type_col)).upper()
            if metric_type in {"MPI", "CPI"}:
                raw = row_value(row, value_col)
                store_metric_value(
                    metric_type,
                    make_metric_value(raw, value_col, info, total_value, bad_value),
                )
            continue

        if mpi_col is not None:
            raw = row_value(row, mpi_col)
            store_metric_value(
                "MPI",
                make_metric_value(raw, mpi_col, info, total_value, bad_value),
            )
        if cpi_col is not None:
            raw = row_value(row, cpi_col)
            store_metric_value(
                "CPI",
                make_metric_value(raw, cpi_col, info, total_value, bad_value),
            )

    wb.close()
    return values


def get_value(
    datasets: dict[tuple[str, str], dict[str, dict[str, MetricValue]]],
    side: str,
    kind: str,
    english_name: str,
    metric_type: str,
) -> MetricValue:
    metric_key = normalize(english_name)
    source_metric_key = f"{normalize(kind).lower()}::{metric_key}"
    for dataset_kind in (kind, "unified"):
        by_metric = datasets.get((side, dataset_kind), {})
        for lookup_key in (source_metric_key, metric_key):
            values = by_metric.get(lookup_key, {})
            preferred = values.get(metric_type)
            if preferred and preferred.raw not in (None, ""):
                return preferred
            for fallback in ("MPI", "CPI"):
                candidate = values.get(fallback)
                if candidate and candidate.raw not in (None, ""):
                    return candidate
    return MetricValue()


def compare_metric(base: MetricValue, target: MetricValue) -> tuple[str, float | None, float | None, str]:
    number_diff = (
        target.number - base.number
        if base.number is not None and target.number is not None
        else None
    )
    bad_diff = (
        target.bad_case - base.bad_case
        if base.bad_case is not None and target.bad_case is not None
        else None
    )

    if base.raw in (None, "") or target.raw in (None, ""):
        return "缺失", number_diff, bad_diff, "缺少base或target数据"

    if base.bad_case is not None and target.bad_case is not None:
        if base.bad_case == 0 and target.bad_case == 0:
            return "一致", number_diff, bad_diff, "base和target bad case均为0"
        if target.bad_case == 0:
            return "提升", number_diff, bad_diff, "target bad case为0"
        if base.bad_case == 0:
            return "降低", number_diff, bad_diff, "base bad case为0"

    if number_diff is not None:
        if abs(number_diff) < 1e-12:
            return "持平", number_diff, bad_diff, "MPI/CPI数值相同"
        if number_diff > 0:
            return "提升", number_diff, bad_diff, "MPI/CPI数值提升"
        return "降低", number_diff, bad_diff, "MPI/CPI数值降低"

    if base.bad_case is not None and target.bad_case is not None:
        if target.bad_case < base.bad_case:
            return "提升", number_diff, bad_diff, "target bad case更少"
        if target.bad_case > base.bad_case:
            return "降低", number_diff, bad_diff, "target bad case更多"
        return "持平", number_diff, bad_diff, "base和target bad case相同"

    if number_diff is None:
        return "无法比较", number_diff, bad_diff, "缺少可比较数值"


def color_target_cell(cell, result: str) -> None:
    if result == "提升":
        cell.fill = GREEN_FILL
    elif result == "降低":
        cell.fill = RED_FILL
    elif result == "缺失":
        cell.fill = YELLOW_FILL


def set_font_name(cell, font_name: str = REPORT_FONT_NAME, color: str | None = None) -> None:
    font = copy(cell.font)
    font.name = font_name
    font.scheme = None
    if color is not None:
        font.color = color
    cell.font = font


def set_wrapped_alignment(cell) -> None:
    alignment = copy(cell.alignment)
    alignment.wrap_text = True
    alignment.vertical = "center"
    cell.alignment = alignment


def text_display_width(value: Any) -> int:
    if value is None:
        return 0
    width = 0
    for char in str(value):
        if unicodedata.east_asian_width(char) in {"F", "W", "A"}:
            width += 2
        else:
            width += 1
    return width


def column_display_width(ws, col: int) -> float:
    width = ws.column_dimensions[get_column_letter(col)].width
    return float(width) if width is not None else 8.43


def effective_cell_width(ws, row: int, col: int) -> float:
    for merged_range in ws.merged_cells.ranges:
        if (
            merged_range.min_row <= row <= merged_range.max_row
            and merged_range.min_col <= col <= merged_range.max_col
        ):
            if row != merged_range.min_row or col != merged_range.min_col:
                return 0.0
            return sum(
                column_display_width(ws, merged_col)
                for merged_col in range(merged_range.min_col, merged_range.max_col + 1)
            )
    return column_display_width(ws, col)


def estimate_cell_line_count(ws, row: int, col: int) -> int:
    cell = ws.cell(row, col)
    value = cell.value
    if value is None:
        return 1

    width = max(effective_cell_width(ws, row, col) - 1.0, 1.0)
    lines = 0
    for part in str(value).splitlines() or [""]:
        lines += max(1, math.ceil(text_display_width(part) / width))
    return max(lines, 1)


def apply_auto_row_heights(ws) -> None:
    for row in range(1, ws.max_row + 1):
        if ws.row_dimensions[row].hidden:
            continue
        line_count = 1
        for col in range(1, ws.max_column + 1):
            line_count = max(line_count, estimate_cell_line_count(ws, row, col))
        ws.row_dimensions[row].height = min(
            max(DEFAULT_ROW_HEIGHT, line_count * ROW_HEIGHT_PER_LINE + 2.0),
            MAX_ROW_HEIGHT,
        )


def apply_report_format(ws) -> None:
    header_rows = 5 if ws.title == SUMMARY_SHEET else 1
    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            set_font_name(cell)
            set_wrapped_alignment(cell)
            if row > header_rows:
                set_font_name(cell, color=DEFAULT_FONT_COLOR)
            if row >= 6 and 2 <= col <= 4:
                cell.fill = WHITE_FILL
                set_font_name(cell, color=DEFAULT_FONT_COLOR)
    apply_auto_row_heights(ws)


def normalize_workbook_font_table(wb) -> None:
    """Ensure inherited/default styles, including merged-cell placeholders, use the report font."""
    for index, font in enumerate(list(wb._fonts)):
        normalized_font = copy(font)
        normalized_font.name = REPORT_FONT_NAME
        normalized_font.scheme = None
        wb._fonts[index] = normalized_font

    for named_style in getattr(wb, "_named_styles", []):
        normalized_font = copy(named_style.font)
        normalized_font.name = REPORT_FONT_NAME
        normalized_font.scheme = None
        named_style.font = normalized_font


def apply_workbook_format(wb) -> None:
    for ws in wb.worksheets:
        apply_report_format(ws)
    normalize_workbook_font_table(wb)


def keep_sheets(wb, keep_names: set[str]) -> None:
    for sheet_name in list(wb.sheetnames):
        if sheet_name not in keep_names:
            del wb[sheet_name]
    wb.active = wb.sheetnames.index(SUMMARY_SHEET)


def write_top_headers(ws, base_info: ResultFile, target_info: ResultFile) -> None:
    # E/F and G/H are merged in the template.
    ws["E1"] = base_info.vehicle
    ws["G1"] = target_info.vehicle
    ws["E2"] = base_info.date_range
    ws["G2"] = target_info.date_range
    ws["E3"] = base_info.version
    ws["G3"] = target_info.version


def hide_unmapped_template_rows(
    ws, report_profile: str, metric_filter: set[str] | None = None
) -> set[int]:
    hidden_rows: set[int] = set()
    for row in range(6, ws.max_row + 1):
        cn_name = normalize(ws.cell(row, 2).value)
        if not cn_name:
            continue
        mapping = resolve_signal_mapping(cn_name, report_profile)
        if mapping is None or (metric_filter is not None and cn_name not in metric_filter):
            ws.row_dimensions[row].hidden = True
            hidden_rows.add(row)
        else:
            ws.row_dimensions[row].hidden = False
    return hidden_rows


def apply_metric_standards(ws, report_profile: str) -> None:
    for merged_range in list(ws.merged_cells.ranges):
        if merged_range.min_col <= 4 <= merged_range.max_col and merged_range.max_row >= 6:
            anchor = ws.cell(merged_range.min_row, merged_range.min_col)
            ws.unmerge_cells(str(merged_range))
            for row in range(merged_range.min_row, merged_range.max_row + 1):
                for col in range(merged_range.min_col, merged_range.max_col + 1):
                    copy_cell_style(anchor, ws.cell(row, col))

    for row in range(6, ws.max_row + 1):
        standard = resolve_metric_standard(ws.cell(row, 2).value, report_profile)
        if standard is not None:
            ws.cell(row, 4).value = standard


def write_mapping_sheet(wb, rows: list[list[Any]]) -> None:
    if MAPPING_SHEET in wb.sheetnames:
        del wb[MAPPING_SHEET]
    ws = wb.create_sheet(MAPPING_SHEET)
    headers = [
        "序号",
        "模板行",
        "指标类型",
        "指标名称",
        "英文名",
        "取值列",
        "base来源文件",
        "base原始值",
        "base数值",
        "base总量",
        "base bad case",
        "target来源文件",
        "target原始值",
        "target数值",
        "target总量",
        "target bad case",
        "MPI/CPI差值(target-base)",
        "bad case差值(target-base)",
        "比较结果",
        "比较依据",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for row in rows:
        ws.append(row)
        result = row[-2]
        result_cell = ws.cell(ws.max_row, 19)
        target_value_cell = ws.cell(ws.max_row, 13)
        number_diff_cell = ws.cell(ws.max_row, 17)
        bad_diff_cell = ws.cell(ws.max_row, 18)
        if result == "提升":
            for cell in (result_cell, target_value_cell, number_diff_cell, bad_diff_cell):
                cell.fill = GREEN_FILL
        elif result == "降低":
            for cell in (result_cell, target_value_cell, number_diff_cell, bad_diff_cell):
                cell.fill = RED_FILL
        elif result == "缺失":
            result_cell.fill = YELLOW_FILL

    widths = [8, 8, 12, 28, 38, 10, 36, 18, 12, 12, 14, 36, 18, 12, 12, 14, 18, 20, 12, 28]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def choose_header_info(reports: dict[str, ResultFile]) -> ResultFile:
    return reports.get("topic") or reports.get("unified") or next(iter(reports.values()))


def blank_header_info(side: str) -> ResultFile:
    return ResultFile(
        side=side,
        kind="",
        vehicle="",
        start_date="",
        end_date="",
        version="",
        path=Path(),
    )


def normalize_report_profile(value: str) -> str:
    key = normalize(value).lower()
    if key in {"1", "j6b"}:
        return REPORT_PROFILE_J6B
    if key in {"2", "mmt"}:
        return REPORT_PROFILE_MMT
    raise ValueError(f"未知报告配置：{value}，可选：J6B 或 MMT")


def choose_report_profile(value: str | None) -> str:
    if value:
        return normalize_report_profile(value)

    if sys.stdin.isatty():
        while True:
            choice = input("请选择报告配置：1) J6B  2) MMT（默认 2）：").strip()
            if not choice:
                return REPORT_PROFILE_MMT
            try:
                return normalize_report_profile(choice)
            except ValueError as exc:
                print(exc)

    print("未检测到交互输入，默认使用 MMT。如需 J6B，请加 --report-profile J6B。")
    return REPORT_PROFILE_MMT


def metric_filter_for_profile(report_profile: str) -> set[str] | None:
    if report_profile == REPORT_PROFILE_J6B:
        return {normalize(name) for name in J6B_METRIC_NAMES}
    return None


def safe_filename_token(value: str) -> str:
    token = normalize(value).replace("+", "P")
    token = re.sub(r'[<>:"/\\|?*]+', "_", token)
    return token.strip("._ ") or "NA"


def make_default_output_path(root: Path, report_profile: str, target_info: ResultFile) -> Path:
    parts = ["MPI_CPI_汇总对比", safe_filename_token(report_profile)]
    if target_info.date_range:
        parts.append(safe_filename_token(target_info.date_range))
    return root / DEFAULT_REPORT_DIR / f"{'_'.join(parts)}.xlsx"


def split_date_range(value: str) -> tuple[str, str]:
    dates = extract_dates_from_text(value)
    if len(dates) >= 2:
        dates = sorted(dates)
        return dates[0], dates[-1]
    if len(dates) == 1:
        return dates[0], dates[0]
    return value.strip(), ""


def apply_header_overrides(
    info: ResultFile,
    vehicle: str | None,
    date_range: str | None,
    version: str | None,
) -> ResultFile:
    start_date = info.start_date
    end_date = info.end_date
    if date_range:
        start_date, end_date = split_date_range(date_range)
    return replace(
        info,
        vehicle=vehicle or info.vehicle,
        start_date=start_date,
        end_date=end_date,
        version=version or info.version,
    )


def build_report(
    root: Path,
    template: Path,
    output: Path | None,
    base_path: Path | None,
    target_path: Path | None,
    include_details: bool,
    report_profile: str,
    base_vehicle: str | None = None,
    target_vehicle: str | None = None,
    base_date: str | None = None,
    target_date: str | None = None,
    base_version: str | None = None,
    target_version: str | None = None,
    base_source_date: str | None = None,
    target_source_date: str | None = None,
) -> Path:
    base_input_path = base_path or default_side_input_path(root, "base", report_profile)
    target_input_path = target_path or default_side_input_path(root, "target", report_profile)
    base_reports = find_side_reports(root, "base", base_input_path, required=False)
    target_reports = find_side_reports(root, "target", target_input_path, required=True)
    base_reports = prepare_latest_reports(root, report_profile, "base", base_reports, base_source_date)
    target_reports = prepare_latest_reports(root, report_profile, "target", target_reports, target_source_date)
    base_available = bool(base_reports)
    base_info = apply_header_overrides(
        choose_header_info(base_reports) if base_available else blank_header_info("base"),
        base_vehicle,
        base_date,
        base_version,
    )
    target_info = apply_header_overrides(
        choose_header_info(target_reports), target_vehicle, target_date, target_version
    )
    if output is None:
        output = make_default_output_path(root, report_profile, target_info)

    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, output)

    wb = load_workbook(output)
    if SUMMARY_SHEET not in wb.sheetnames:
        raise ValueError(f"模板中未找到 {SUMMARY_SHEET}")
    ws = wb[SUMMARY_SHEET]

    write_top_headers(ws, base_info, target_info)
    hidden_rows = hide_unmapped_template_rows(
        ws, report_profile, metric_filter_for_profile(report_profile)
    )
    apply_metric_standards(ws, report_profile)

    datasets: dict[tuple[str, str], dict[str, dict[str, MetricValue]]] = {}
    for kind, info in base_reports.items():
        datasets[("base", kind)] = load_metric_values(info)
    for kind, info in target_reports.items():
        datasets[("target", kind)] = load_metric_values(info)

    mapping_rows: list[list[Any]] = []
    sequence = 1
    for row in range(1, ws.max_row + 1):
        if row in hidden_rows:
            continue
        cn_name = normalize(ws.cell(row, 2).value)
        mapping = resolve_signal_mapping(cn_name, report_profile)
        if mapping is None:
            continue

        kind, english_name, metric_type = mapping
        base_col = BASE_MPI_COL if metric_type == "MPI" else BASE_CPI_COL
        target_col = TARGET_MPI_COL if metric_type == "MPI" else TARGET_CPI_COL

        base_value = (
            get_value(datasets, "base", kind, english_name, metric_type)
            if base_available
            else MetricValue()
        )
        target_value = get_value(datasets, "target", kind, english_name, metric_type)

        for col in (BASE_MPI_COL, BASE_CPI_COL, TARGET_MPI_COL, TARGET_CPI_COL):
            ws[f"{col}{row}"] = MISSING_DISPLAY

        if base_available:
            ws[f"{base_col}{row}"] = display_raw(base_value)
        ws[f"{target_col}{row}"] = display_raw(target_value)
        if base_available:
            result, number_diff, bad_diff, compare_reason = compare_metric(base_value, target_value)
            color_target_cell(ws[f"{target_col}{row}"], result)
        else:
            result, number_diff, bad_diff = "仅target", None, None
            compare_reason = "base为空，仅展示target结果"

        mapping_rows.append(
            [
                sequence,
                row,
                kind,
                cn_name,
                english_name,
                metric_type,
                base_value.source_file,
                display_raw(base_value),
                base_value.number,
                base_value.total,
                base_value.bad_case,
                target_value.source_file,
                display_raw(target_value),
                target_value.number,
                target_value.total,
                target_value.bad_case,
                number_diff,
                bad_diff,
                result,
                compare_reason,
            ]
        )
        sequence += 1

    if include_details:
        write_mapping_sheet(wb, mapping_rows)
        keep_sheets(wb, {SUMMARY_SHEET, MAPPING_SHEET})
    else:
        keep_sheets(wb, {SUMMARY_SHEET})
    apply_workbook_format(wb)
    wb.save(output)
    return output


def resolve_path(root: Path, value: str) -> Path:
    path = Path(value)
    if not path.is_absolute():
        path = root / path
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description="生成 MPI/CPI 汇总对比表")
    parser.add_argument("--root", default=".", help="项目根目录，默认当前目录")
    parser.add_argument("--base", default=None, help="base 报告文件或目录，默认读取 root/base/<报告配置>")
    parser.add_argument("--target", default=None, help="target 报告文件或目录，默认读取 root/target/<报告配置>")
    parser.add_argument(
        "--template",
        default=None,
        help="模板文件；默认读取 templates/template_<报告配置>.xlsx",
    )
    parser.add_argument(
        "--output",
        default=None,
        help="输出文件名；默认写入 reports/MPI_CPI_汇总对比_<报告配置>_<target日期>.xlsx",
    )
    parser.add_argument(
        "--report-profile",
        "--report-version",
        dest="report_profile",
        default=None,
        help="报告配置：J6B 或 MMT；不传时交互选择",
    )
    parser.add_argument("--details-sheet", action="store_true", help="保留“信号对照表”明细 sheet")
    parser.add_argument("--base-vehicle", default=None, help="覆盖 base 表头车型")
    parser.add_argument("--target-vehicle", default=None, help="覆盖 target 表头车型")
    parser.add_argument("--base-date", default=None, help="覆盖 base 表头日期")
    parser.add_argument("--target-date", default=None, help="覆盖 target 表头日期")
    parser.add_argument("--base-version", default=None, help="覆盖 base 表头软件版本")
    parser.add_argument("--target-version", default=None, help="覆盖 target 表头软件版本")
    parser.add_argument(
        "--base-source-date",
        default=None,
        help="当 base 是 MPI_CPI 汇总对比表时，选择其中某个日期组作为 base，例如 20260722-20260723",
    )
    parser.add_argument(
        "--target-source-date",
        default=None,
        help="当 target 是 MPI_CPI 汇总对比表时，选择其中某个日期组作为 target",
    )
    args = parser.parse_args()

    try:
        report_profile = choose_report_profile(args.report_profile)
    except ValueError as exc:
        parser.error(str(exc))

    root = Path(args.root).resolve()
    template = resolve_path(root, args.template) if args.template else default_template_path(root, report_profile)
    output = resolve_path(root, args.output) if args.output else None
    base_path = resolve_path(root, args.base) if args.base else None
    target_path = resolve_path(root, args.target) if args.target else None

    if not template.exists():
        raise FileNotFoundError(f"模板不存在：{template}")

    output = build_report(
        root,
        template,
        output,
        base_path,
        target_path,
        args.details_sheet,
        report_profile,
        base_vehicle=args.base_vehicle,
        target_vehicle=args.target_vehicle,
        base_date=args.base_date,
        target_date=args.target_date,
        base_version=args.base_version,
        target_version=args.target_version,
        base_source_date=args.base_source_date,
        target_source_date=args.target_source_date,
    )
    print(f"报告配置：{report_profile}")
    print(f"已生成：{output}")


if __name__ == "__main__":
    main()
