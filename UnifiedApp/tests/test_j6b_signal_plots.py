import json
import os
from pathlib import Path
import subprocess
import sys
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication

from UnifiedApp.unified_app.job_manager import JobManager
from UnifiedApp.unified_app.job_spec import J6BJobSpec
from UnifiedApp.unified_app.ui.j6b_form import J6BForm
from test_powertrain_reporting import Settings, ROOT


class J6BSignalPlotTests(unittest.TestCase):
    def test_ui_config_and_presets(self):
        app = QApplication.instance() or QApplication([])
        manager = JobManager()
        form = J6BForm(Settings, manager)
        try:
            self.assertFalse(form.bad_case_signal_plots_check.isChecked())
            self.assertTrue(form.lcc_sync_check.isChecked())
            form.input_list.addItem(str(ROOT))
            form.vehicle_edit.setText("SAMPLE")
            form.software_version_edit.setText("TEST")
            form.date_range_edit.setText("20261028")
            for enabled in (True, False):
                form.bad_case_signal_plots_check.setChecked(enabled)
                form.lcc_sync_check.setChecked(enabled)
                spec = J6BJobSpec.from_dict(json.loads(json.dumps(form._collect_spec().to_dict())))
                form._apply_spec(spec)
                self.assertEqual(form.bad_case_signal_plots_check.isChecked(), enabled)
                self.assertEqual(form.lcc_sync_check.isChecked(), enabled)
                with patch.object(manager, "_try_start_next"):
                    job_id = manager.submit_j6b(
                        sys.executable, ROOT / "J6B_UnifiedMining",
                        ROOT / "J6B_UnifiedMining" / "config.json", spec,
                    )
                job = manager.get_job(job_id)
                try:
                    config = json.loads(job.temp_files[0].read_text(encoding="utf-8"))
                    self.assertIs(config["bad_case_signal_plots"], enabled)
                    self.assertIs(config["lcc_sync_check"], enabled)
                    self.assertIs(config["lcc_sync_deferred"], True)
                    self.assertEqual(job.sync_status, "等待报告完成" if enabled else "未启用")
                finally:
                    manager._cleanup_temp_files(job)
            self.assertFalse(J6BJobSpec.from_dict({}).bad_case_signal_plots)
            self.assertTrue(J6BJobSpec.from_dict({}).lcc_sync_check)
        finally:
            form.close()

    def test_embedded_plots_and_threshold_evidence(self):
        result = subprocess.run(
            [sys.executable, "-c", r'''
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile
import h5py
import numpy as np
from openpyxl import load_workbook
from PIL import Image
from DataMining import load_config, load_metrics, main
from bad_case_report import write_bad_case_report
from bad_case_image_report import (
    build_evidence_panel, build_panels, evidence_markers, evidence_summary,
    generate_report_with_images, header_columns, read_cases, CaseRow,
    plot_window_bounds, draw_case_plot)
from metric_utils import Segment, jerk_from_acceleration
from reporting import metric_json_payload
from signal_preprocessor import build_signal_frame
from filters.follow_stop_min_jerk import create_metric

with tempfile.TemporaryDirectory() as folder:
    work = Path(folder)
    t = np.arange(0, 12.01, 0.05)
    accel = np.where((t >= 3) & (t < 4), -4 * (t - 3),
                     np.where((t >= 4) & (t < 6), -4.0, 0.0))
    data = {
        "acceleration": accel, "ego_speed": np.where(t < 10, 20.0, 0.0),
        "acc_active": np.where(t < 11, 1, 0),
        "abs_active": np.where((t >= 10.5) & (t < 11.5), 1, 0),
        "front_range": np.full(t.size, 10.0), "front_velocity_abs": np.zeros(t.size),
        "front_target_id": np.ones(t.size), "front_lat_range": np.zeros(t.size),
        "flt1_cutin_flag": np.zeros(t.size),
    }
    h5path = work / "synthetic.h5"
    with h5py.File(h5path, "w") as h5:
        h5.create_group("metadata")
        group = h5.create_group("sample")
        group.create_dataset("timestamp", data=t)
        signals = group.create_group("signals")
        for name, values in data.items():
            signals.create_dataset(name, data=values)
    config = {
        "input_path": [str(work)], "powertrain_type": "PHEV",
        "signals": {name: "sample/signals/" + name for name in data},
        "enabled_metrics": ["follow_stop_min_jerk", "acc_exit_after_abs"],
        "bad_case_signal_plots": True, "output_dir": str(work / "run"),
        "report_name": "Longitudinal_CPI_Report",
    }
    cfg = work / "config.json"
    cfg.write_text(json.dumps(config), encoding="utf-8")
    frame = build_signal_frame(h5path, config)
    metric = create_metric()
    metric.powertrain_type = "PHEV"
    measured = metric.run(frame)
    assert measured.bad_cases == 1, (measured.status, measured.details)
    event = measured.events[0]
    assert event.values["screening_lower"] == -2.3
    expected_jerk = jerk_from_acceleration(t, accel)
    assert abs(event.values["screening_value"] - np.nanmin(expected_jerk[:201])) < 1e-8
    assert event.values["value_time_s"] >= 3
    # Actual end-to-end entry point reopens the synthetic H5 and embeds both CPI and MPI plots.
    config["bad_case_extract_config"] = str(work / "bad_case_extract.json")
    cfg.write_text(json.dumps(config), encoding="utf-8")
    (work / "bad_case_extract.json").write_text(json.dumps({"enabled": True}), encoding="utf-8")
    assert main(["DataMining.py", str(cfg)]) == 0
    reports = list((work / "run").rglob("bad_case_report.xlsx"))
    assert len(reports) == 1, reports
    report = reports[0]
    wb = load_workbook(report)
    try:
        ws = wb.active
        headers = header_columns(ws)
        assert "关键信号图" in headers and "绘图状态" in headers
        assert len(ws._images) == 2, len(ws._images)
        cases = read_cases(ws, headers)
        cpi = next(case for case in cases if case.metric_name == "follow_stop_min_jerk")
        panel = build_evidence_panel(frame, cpi, np.ones(t.size, dtype=bool))
        np.testing.assert_allclose(panel.lines[0].values, expected_jerk)
        assert panel.thresholds == [("筛选下限", -2.3)]
        assert np.any(panel.highlight)
        assert not np.any(panel.highlight[t > cpi.end_s])
        assert "判定值时刻" in dict(evidence_markers(cpi))
        assert "-2.3" in evidence_summary(cpi)
        mpi = next(case for case in cases if case.metric_name == "acc_exit_after_abs")
        assert dict(evidence_markers(mpi))["功能激活"] == 10.5
        assert dict(evidence_markers(mpi))["ACC退出"] == 11
        assert build_evidence_panel(frame, mpi, np.ones(t.size, dtype=bool)) is None
        for case in cases:
            assert ws.cell(case.row_index, headers["绘图状态"]).value.startswith("已生成")
            assert (report.parent / ws.cell(case.row_index, headers["绘图状态"]).hyperlink.target).is_file()
    finally:
        wb.close()
    with ZipFile(report) as archive:
        assert len([name for name in archive.namelist() if name.startswith("xl/media/")]) == 2
    generate_report_with_images(report, report, config=config)
    repeated = load_workbook(report)
    try:
        assert len(repeated.active._images) == 2
    finally:
        repeated.close()
    # Scalar measurements use an event statistic, not a fictional instantaneous signal.
    scalar = CaseRow(2, "stop2go_max_acc_time", str(h5path), 1, 8, 10, "", {
        "screening_measure": "max_acc_time", "screening_value": 5,
        "screening_lower": 1, "screening_upper": 4.5,
        "value_time_s": 7, "accel_build_start_time_s": 2})
    scalar_panel = build_evidence_panel(frame, scalar, np.ones(t.size, dtype=bool))
    assert "事件统计值" in scalar_panel.title
    assert np.all(scalar_panel.lines[0].values[(t >= 1) & (t <= 8)] == 5)
    strict = CaseRow(2, "driving_uninterfered_accel_avg_acc_0_30", str(h5path), 1, 8, 0, "", {
        "screening_measure": "average_acceleration_mps2", "screening_value": 1.5,
        "screening_lower": 1.5, "screening_upper": None, "screening_strict": True})
    assert np.any(build_evidence_panel(frame, strict, np.ones(t.size, dtype=bool)).highlight)
    assert plot_window_bounds(np.array([0., 200.]), 10, 150)[1] >= 150
    # A missing H5 is visibly reported per row and never presented as a plotted success.
    with patch("bad_case_image_report.FrameCache.get", side_effect=FileNotFoundError("missing H5")):
        logs = []
        failed = generate_report_with_images(report, work / "failed.xlsx", logs.append, config)
    failed_wb = load_workbook(failed)
    try:
        failed_ws = failed_wb.active
        failed_headers = header_columns(failed_ws)
        assert len(failed_ws._images) == 0
        assert failed_ws.cell(2, failed_headers["绘图状态"]).value.startswith("绘图失败")
        assert "失败 2" in logs[-1]
    finally:
        failed_wb.close()
    # Every active metric can render its available contextual signals, even with no scalar threshold.
    for metric in load_metrics(Path("filters"), ["all"], config):
        measure = getattr(metric, "measure", None)
        values = {}
        if measure:
            values = {
                "screening_measure": measure, "screening_value": -4,
                "screening_lower": -2, "screening_upper": None, "value_time_s": 4,
            }
        case = CaseRow(2, metric.name, str(h5path), 3, 5, 1, "synthetic event", values)
        draw_case_plot(frame, case, work / "all_metrics.png")
        with Image.open(work / "all_metrics.png") as image:
            assert image.width == 1180 and image.height >= 650
    # Disabling the switch preserves the plain report and does not invoke plotting.
    config["bad_case_signal_plots"] = False
    config["output_dir"] = str(work / "disabled")
    cfg.write_text(json.dumps(config), encoding="utf-8")
    with patch("bad_case_image_report.generate_report_with_images") as plotting:
        assert main(["DataMining.py", str(cfg)]) == 0
        plotting.assert_not_called()
    plain = load_workbook(next((work / "disabled").rglob("bad_case_report.xlsx")))
    try:
        assert "关键信号图" not in header_columns(plain.active)
        assert not plain.active._images
    finally:
        plain.close()
    config["bad_case_signal_plots"] = "false"
    cfg.write_text(json.dumps(config), encoding="utf-8")
    try:
        load_config(cfg)
        raise AssertionError("invalid plot flag accepted")
    except ValueError:
        pass
'''],
            cwd=ROOT / "J6B_UnifiedMining",
            env={**os.environ, "PYTHONIOENCODING": "utf-8"},
            capture_output=True,
            encoding="utf-8",
            errors="replace",
            timeout=120,
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)


if __name__ == "__main__":
    unittest.main()
