from pathlib import Path
import unittest

from openpyxl import Workbook, load_workbook

from CompareKPI_United.mpi_cpi_compare import (
    POWERTRAIN_METRIC_STANDARD_OVERRIDES,
    POWERTRAIN_TYPES,
    apply_metric_standards,
    normalize,
    resolve_metric_standard,
)


ROOT = Path(__file__).resolve().parents[2] / "CompareKPI_United"


class CurrentScreeningStandardsTests(unittest.TestCase):
    def test_templates_match_current_standards_for_all_powertrains(self):
        for path, profile in (
            (ROOT / "template.xlsx", "MMT"),
            (ROOT / "templates" / "template_J6B.xlsx", "J6B"),
            (ROOT / "templates" / "template_MMT.xlsx", "MMT"),
        ):
            for powertrain in POWERTRAIN_TYPES:
                with self.subTest(template=path.name, powertrain=powertrain):
                    workbook = load_workbook(path)
                    self.addCleanup(workbook.close)
                    sheet = workbook.active
                    expected = POWERTRAIN_METRIC_STANDARD_OVERRIDES[powertrain][profile]
                    checked = set()
                    for row in range(6, sheet.max_row + 1):
                        name = normalize(sheet.cell(row, 2).value)
                        if name not in expected:
                            continue
                        labels = dict(
                            line.split("：", 1) for line in sheet.cell(row, 4).value.splitlines()
                        )
                        self.assertEqual(labels[powertrain], expected[name])
                        checked.add(name)
                    self.assertEqual(checked, set(expected))
                    apply_metric_standards(sheet, profile, powertrain)
                    for row in range(6, sheet.max_row + 1):
                        name = normalize(sheet.cell(row, 2).value)
                        if name in expected:
                            self.assertEqual(sheet.cell(row, 4).value, expected[name])

    def test_outdated_labeled_and_unlabeled_templates_cannot_restore_old_rules(self):
        for profile in ("J6B", "MMT"):
            for powertrain in POWERTRAIN_TYPES:
                standards = POWERTRAIN_METRIC_STANDARD_OVERRIDES[powertrain][profile]
                for old_value in ("ICE：1.0~4.5s\nPHEV：1.0~4s\nBEV：1.0~4s", "旧筛选标准"):
                    workbook = Workbook()
                    self.addCleanup(workbook.close)
                    sheet = workbook.active
                    for row, name in enumerate(standards, 6):
                        sheet.cell(row, 2, name)
                        sheet.cell(row, 4, old_value)
                    apply_metric_standards(sheet, profile, powertrain)
                    for row, (name, expected) in enumerate(standards.items(), 6):
                        self.assertEqual(sheet.cell(row, 4).value, expected)
                        self.assertEqual(resolve_metric_standard(name, profile, powertrain), expected)

    def test_concise_thresholds_units_and_removed_lower_limits(self):
        for profile in ("J6B", "MMT"):
            for powertrain in POWERTRAIN_TYPES:
                standards = POWERTRAIN_METRIC_STANDARD_OVERRIDES[powertrain][profile]
                ice = powertrain == "ICE"
                upper = (4.5 if ice else 4) if profile == "J6B" else (4 if ice else 3.5)
                self.assertEqual(standards["起步加速建立时间"], f"<={upper:g}s")
                self.assertEqual(
                    standards["跟车起步舒适性"],
                    "最大jerk<=2m/s3" if profile == "J6B" else "最大jerk<=1.5m/s3",
                )
                if profile == "MMT":
                    self.assertEqual(standards["绿灯起步舒适性"], "最大jerk<=2.5m/s3")
                fluct = standards["坡道稳速抖动"]
                threshold = "0.6" if powertrain == "ICE" else "0.3"
                if profile == "J6B":
                    self.assertEqual(fluct, "车速半极差<=" + threshold + "m/s")
                else:
                    self.assertEqual(fluct, "仪表车速极差<=" + threshold + "km/h")
                distance = standards["跟车刹停距离"]
                expected_range = "2.5~4.5m" if profile == "J6B" and powertrain == "ICE" else "3~4m"
                self.assertEqual(distance, expected_range)
                for text in standards.values():
                    self.assertNotIn("场景", text)
                    self.assertNotIn("拟合", text)
                    self.assertNotIn("观察", text)
                    self.assertNotIn("无下限", text)
                    self.assertLessEqual(len(text), 70)


if __name__ == "__main__":
    unittest.main()
