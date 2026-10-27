import os
from pathlib import Path
import subprocess
import sys
import unittest


ROOT = Path(__file__).resolve().parents[2]


class SummaryNameTests(unittest.TestCase):
    def test_summary_names_and_comparison_compatibility(self):
        for profile in ("J6B", "MMT"):
            with self.subTest(profile=profile):
                result = subprocess.run(
                    [sys.executable, "-c", """
import csv
from pathlib import Path
from types import SimpleNamespace
import sys
import tempfile
from openpyxl import load_workbook
sys.path.insert(0, str(Path.cwd().parent))
from CompareKPI_United.mpi_cpi_compare import (
    ResultFile, convert_old_report_to_latest, load_metric_values)
from metric_display_names import metric_display_name

profile = sys.argv[1]
names = ("follow_stop_min_jerk", "acc_exit_after_abs")
with tempfile.TemporaryDirectory() as folder:
    work = Path(folder)
    if profile == "J6B":
        from metric_utils import MetricResult
        from reporting import StreamingReportWriter, write_summary_xlsx
        results = {
            name: [MetricResult(name, "MPI" if "exit" in name else "CPI",
                               "case_rate", "synthetic.h5", 60.0, "OK",
                               30, 2, 5.0, [], [])]
            for name in names
        }
        batch = work / "batch.xlsx"
        write_summary_xlsx(results, batch)
        metrics = [SimpleNamespace(name=name, category=results[name][0].category,
                                   result_mode="case_rate") for name in names]
        with StreamingReportWriter(work, "Longitudinal_CPI_Report.xlsx", metrics) as writer:
            for name in names:
                writer.write_result(name, results[name][0])
        outputs = (batch, writer.report_path)
    else:
        from DataMining import generate_unified_report
        stamp = "2026-10-27_15-05-00"
        categories = {}
        for name in names:
            prefix = ("can" if "exit" in name else "topic") + "__" + name
            categories[prefix] = "MPI" if "exit" in name else "CPI"
            with (work / f"{prefix}_{stamp}.csv").open(
                "w", newline="", encoding="utf-8-sig"
            ) as stream:
                writer = csv.DictWriter(stream, fieldnames=[
                    "总Case", "Bad Case", "段里程(km)", "文件夹路径",
                    "严重度得分", "符合要求的内容"])
                writer.writeheader()
                writer.writerow({"总Case": 30, "Bad Case": 2, "段里程(km)": 60,
                                 "文件夹路径": "synthetic", "严重度得分": 5,
                                 "符合要求的内容": "test"})
        outputs = (generate_unified_report(work, "Longitudinal_CPI_Report", stamp, categories),)
    for index, output in enumerate(outputs):
        wb = load_workbook(output)
        try:
            ws = wb["所有指标统计汇总"]
            headers = {cell.value: cell.column for cell in ws[1]}
            assert headers["中文指标名称"] == headers["指标名称"] + 1
            rows = {ws.cell(r, headers["指标名称"]).value: r for r in range(2, ws.max_row + 1)}
            assert set(rows) == set(names)
            for name, row in rows.items():
                assert ws.cell(row, headers["中文指标名称"]).value == metric_display_name(name)
                total_header = "总Case之和" if profile == "J6B" else "总次数"
                assert ws.cell(row, headers[total_header]).value == 30
        finally:
            wb.close()
        readable = convert_old_report_to_latest(output, work / f"converted{index}.xlsx") \\
            if profile == "J6B" else output
        info = ResultFile("target", "unified", "SAMPLE", "", "", "", readable)
        values = load_metric_values(info)
        assert values["follow_stop_min_jerk"]["CPI"].number == 15
        assert values["acc_exit_after_abs"]["MPI"].number == 30
""", profile],
                    cwd=ROOT / f"{profile}_UnifiedMining",
                    env={**os.environ, "PYTHONIOENCODING": "utf-8"},
                    capture_output=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=60,
                )
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)


if __name__ == "__main__":
    unittest.main()
