import json
from copy import copy
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import Workbook, load_workbook
from PySide6.QtWidgets import QApplication

from CompareKPI_United.mpi_cpi_compare import (
    apply_metric_standards,
    convert_comparison_report_to_latest,
    find_comparison_date_groups,
)
from UnifiedApp.unified_app.job_manager import JobManager
from UnifiedApp.unified_app.job_spec import CompareJobSpec, J6BJobSpec, MMTJobSpec
from UnifiedApp.unified_app.ui.compare_form import CompareForm
from UnifiedApp.unified_app.ui.j6b_form import J6BForm
from UnifiedApp.unified_app.ui.mmt_form import MMTForm


ROOT = Path(__file__).resolve().parents[2]


class Settings:
    def j6b_dir(self):
        return ROOT / "J6B_UnifiedMining"

    def mmt_dir(self):
        return ROOT / "MMT_UnifiedMining"

    def compare_dir(self):
        return ROOT / "CompareKPI_United"


class PowertrainReportingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def test_mining_selection_reaches_config_and_generated_report(self):
        for profile, form_class, spec_class in (
            ("J6B", J6BForm, J6BJobSpec),
            ("MMT", MMTForm, MMTJobSpec),
        ):
            for powertrain in ("ICE", "PHEV", "BEV"):
                with self.subTest(profile=profile, powertrain=powertrain):
                    self._check_mining_and_report(profile, form_class, spec_class, powertrain)

    def _check_mining_and_report(self, profile, form_class, spec_class, powertrain):
        manager = JobManager()
        with patch.object(manager, "_try_start_next"), tempfile.TemporaryDirectory() as folder:
            work = Path(folder)
            form = form_class(Settings, manager)
            self.addCleanup(form.close)
            self.assertEqual(
                [form.powertrain_combo.itemText(i) for i in range(form.powertrain_combo.count())],
                ["ICE", "PHEV", "BEV"],
            )
            form.powertrain_combo.setCurrentText(powertrain)
            form.vehicle_edit.setText("TEST-2-CAR001")
            form.software_version_edit.setText("V1.0")
            if profile == "J6B":
                form.input_list.addItem(str(work))
                form.date_range_edit.setText("20261026-20261026")
            else:
                form.source_path_edit.setText(str(work))
                form.query_date_edit.setText("20261026-20261026")

            spec = form._collect_spec()
            self.assertEqual(spec.powertrain_type, powertrain)
            restored = spec_class.from_dict(json.loads(json.dumps(spec.to_dict())))
            form._apply_spec(restored)
            self.assertEqual(form.powertrain_combo.currentText(), powertrain)
            submit = manager.submit_j6b if profile == "J6B" else manager.submit_mmt
            job_id = submit(
                sys.executable,
                ROOT / (profile + "_UnifiedMining"),
                ROOT / (profile + "_UnifiedMining") / "config.json",
                spec,
                compare_python=sys.executable,
                compare_dir=Settings().compare_dir(),
            )
            mining = manager.get_job(job_id)
            self.addCleanup(manager._cleanup_temp_files, mining)
            config = json.loads(mining.temp_files[0].read_text(encoding="utf-8"))
            self.assertEqual(config["powertrain_type"], powertrain)
            self.assertEqual(mining.pending_compare["spec"].powertrain_type, powertrain)
            metric_check = subprocess.run(
                [sys.executable, "-c", """
import json
from pathlib import Path
import sys
from types import SimpleNamespace
from unittest.mock import patch
import numpy as np
sys.path.insert(0, str(Path.cwd()))
from DataMining import load_config
config_path = Path(sys.argv[1])
profile = sys.argv[2]
powertrain = json.loads(config_path.read_text(encoding="utf-8"))["powertrain_type"]
from powertrain_utils import value_for_powertrain
for kind, expected in (("ICE", 10), ("PHEV", 20), ("BEV", 30)):
    assert value_for_powertrain(kind, ice_value=10, phev_value=20, bev_value=30) == expected
if profile == "J6B":
    from DataMining import load_metrics
    from metric_utils import Segment
    metrics = load_metrics(Path("filters"), ["stop2go_max_jerk"], load_config(config_path))
    metric = metrics[0]
    module = sys.modules[type(metric).__module__]
    frame = SimpleNamespace(time_s=np.array([0.0, 1.0]), source_path="synthetic",
        distance_km=0.0, distance_status="OK", primary_length=2, stitched_source_paths=[])
    with patch.object(module, "find_scene_segments", return_value=[Segment(0, 1)]), \\
         patch.object(module, "measure_segment", return_value=(4.25, 0.5, {})):
        bad = metric.evaluate(frame).bad_cases
else:
    from metrics.topic.code import stop2go_max_jerk as module
    from metrics.can.code import acc_start_no_override_jerkiness as can_module
    config = load_config(config_path)
    assert all(p.powertrain_type == powertrain for p in config.pipelines)
    values = dict(max_jerk_value=3.75, max_jerk_chassis_time_ms=0.0,
        max_jerk_vehicle_speed=1.0, segment_max_vehicle_speed=1.0)
    frames = [{"_powertrain_type": powertrain}]
    with patch.object(module, "_find_risk_segments", return_value=[frames]), \\
         patch.object(module, "_calc_segment_max_jerk", return_value=values):
        bad = module._core_process(frames)[2]["Bad Case"]
    can_result = dict(all_case_count=1, bad_case_count=0, severity_score=0,
        qualified_sample_count=1, scenes=[], events=[])
    with patch.object(can_module, "analyze_acc_start_no_override_jerkiness",
                      return_value=can_result) as analyze:
        can_module.__run_haoran_data_run__({
            "session": SimpleNamespace(primary_h5_file="synthetic.h5"),
            "powertrain_type": powertrain,
        })
    assert analyze.call_args.kwargs["start_window_seconds"] == (3.0 if powertrain == "ICE" else 1.0)
assert bad == (0 if powertrain == "ICE" else 1), (powertrain, bad)
""", str(mining.temp_files[0]), profile],
                cwd=ROOT / (profile + "_UnifiedMining"),
                capture_output=True,
                encoding="utf-8",
                errors="replace",
                timeout=60,
            )
            self.assertEqual(metric_check.returncode, 0, metric_check.stdout + metric_check.stderr)

            # Supply a synthetic mining result to exercise the real auto-compare CLI.
            source = work / "TEST-2-CAR001-20261026-20261026-V1.0.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "所有指标统计汇总"
            ws.append(["指标来源", "指标名称", "指标类型", "总里程(km)", "总次数", "触发次数", "MPI", "CPI"])
            ws.append(["topic", "stop2go_max_jerk", "CPI", 10, 2, 0, None, "100.0(2/0)"])
            ws.append(["topic", "driving_uninterfered_accel_max_acc_0_30", "CPI", 10, 5, 2, None, "2.5(5/2)"])
            wb.save(source)
            wb.close()
            mining.archived_report_path = source
            output = work / "comparison.xlsx"
            mining.pending_compare["spec"].output_path = str(output)
            mining.pending_compare["spec"].base_path = str(work / "no-base")
            manager._trigger_auto_compare(mining)
            comparison = manager.jobs()[-1]
            self.assertEqual(comparison.kind, "compare")
            argument = comparison.command.index("--powertrain-type")
            self.assertEqual(comparison.command[argument + 1], powertrain)
            completed = subprocess.run(
                comparison.command,
                cwd=comparison.cwd,
                stdin=subprocess.DEVNULL,
                capture_output=True,
                encoding="utf-8",
                errors="replace",
                env={**os.environ, "PYTHONIOENCODING": "utf-8"},
                timeout=60,
            )
            self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
            self.assertTrue(output.is_file())
            generated = load_workbook(output)
            template = load_workbook(Settings().compare_dir() / "templates" / f"template_{profile}.xlsx")
            try:
                self.assertIn(powertrain, generated.properties.subject)
                report = generated["Sheet1"]
                self.assertEqual(report["A1"].value, "车型")
                self.assertEqual(report["A2"].value, "动力形式")
                self.assertEqual(report["A3"].value, "日期")
                self.assertEqual(report["A4"].value, "软件版本")
                self.assertEqual(report["E2"].value, powertrain)
                self.assertEqual(report["G2"].value, powertrain)
                self.assertEqual(report["G3"].value, "20261026")
                self.assertEqual(report["G4"].value, "V1.0")
                self.assertIn(powertrain, report["D6"].comment.text)
                self.assertEqual(report["I6"].value, "Base准出标准")
                self.assertEqual(report["J6"].value, "Base准出结论")
                self.assertEqual(report["K6"].value, "Target准出标准")
                self.assertEqual(report["L6"].value, "Target准出结论")
                self.assertEqual(report["I8"].value, "CPI ≥30")
                self.assertEqual(report["J8"].value, "数据缺失或无效，无法判定")
                self.assertEqual(report["K8"].value, "CPI ≥30")
                self.assertEqual(report["L8"].value, "可准出")
                for merged in ("A2:D2", "E2:F2", "G2:H2"):
                    self.assertIn(merged, report.merged_cells)
                for merged in template["Sheet1"].merged_cells.ranges:
                    if merged.min_col <= 4 <= merged.max_col and merged.max_row >= 6:
                        continue
                    shifted = copy(merged)
                    if shifted.min_row >= 2:
                        shifted.shift(row_shift=1)
                    self.assertIn(str(shifted), report.merged_cells)
                checked = 0
                for row in range(6, template["Sheet1"].max_row + 1):
                    original = template["Sheet1"].cell(row, 4).value
                    if not isinstance(original, str):
                        continue
                    selected = [
                        line for line in original.splitlines()
                        if line.startswith(powertrain + "：")
                    ]
                    if selected:
                        self.assertEqual(
                            report.cell(row + 1, 4).value,
                            selected[0].split("：", 1)[1].lstrip(),
                        )
                        checked += 1
                        if "1.5~2.5" in selected[0]:
                            self.assertNotIn("报告事件均值", selected[0])
                            if "0-30" in str(report.cell(row + 1, 2).value):
                                self.assertEqual(report.cell(row + 1, 8).value, "2.5(5/2)")
                self.assertGreater(checked, 20)
                expected_jerk = {
                    ("J6B", "PHEV"): "<=4m/s3",
                    ("J6B", "BEV"): "<=4m/s3",
                    ("J6B", "ICE"): "<=4.5m/s3",
                    ("MMT", "PHEV"): "<=3.5m/s3",
                    ("MMT", "BEV"): "<=3.5m/s3",
                    ("MMT", "ICE"): "<=4m/s3",
                }
                self.assertEqual(
                    report["D8"].value,
                    expected_jerk[profile, powertrain],
                )
                groups = find_comparison_date_groups(report)
                self.assertEqual(len(groups), 1)
                self.assertEqual(groups[0]["version"], "V1.0")
                converted = work / "roundtrip.xlsx"
                self.assertIsNotNone(
                    convert_comparison_report_to_latest(output, converted, None, profile)
                )
                converted_wb = load_workbook(converted, read_only=True)
                try:
                    rows = list(converted_wb.active.values)
                    self.assertTrue(any(row[1] == "stop2go_max_jerk" for row in rows[1:]))
                finally:
                    converted_wb.close()
            finally:
                generated.close()
                template.close()

    def test_manual_compare_selection_and_presets(self):
        manager = JobManager()
        form = CompareForm(Settings, manager)
        self.addCleanup(form.close)
        for powertrain in ("ICE", "PHEV", "BEV"):
            form.powertrain_combo.setCurrentText(powertrain)
            spec = CompareJobSpec.from_dict(form._collect_spec().to_dict())
            self.assertEqual(spec.powertrain_type, powertrain)
            form._apply_spec(spec)
            self.assertEqual(form.powertrain_combo.currentText(), powertrain)
            with patch.object(manager, "_try_start_next"):
                job_id = manager.submit_compare(sys.executable, Settings().compare_dir(), spec)
            command = manager.get_job(job_id).command
            self.assertEqual(command[command.index("--powertrain-type") + 1], powertrain)
        self.assertEqual(CompareJobSpec.from_dict({}).powertrain_type, "PHEV")
        with self.assertRaises(ValueError):
            manager.submit_compare(
                sys.executable, Settings().compare_dir(), CompareJobSpec(powertrain_type="invalid")
            )

    def test_missing_selected_template_standard_is_an_error(self):
        wb = Workbook()
        self.addCleanup(wb.close)
        ws = wb.active
        ws["B6"] = "起步最大jerk"
        ws["D6"] = "PHEV：<=4m/s3"
        with self.assertRaisesRegex(ValueError, "ICE"):
            apply_metric_standards(ws, "J6B", "ICE")

    def test_independent_phev_and_bev_report_standards(self):
        for powertrain, expected in (("ICE", "ice"), ("PHEV", "phev"), ("BEV", "bev")):
            wb = Workbook()
            self.addCleanup(wb.close)
            ws = wb.active
            ws["B6"] = "起步最大jerk"
            ws["D6"] = "ICE：ice\nPHEV：phev\nBEV：bev"
            apply_metric_standards(ws, "J6B", powertrain)
            self.assertEqual(ws["D6"].value, expected)

    def test_legacy_presets_and_templates(self):
        for spec_class in (J6BJobSpec, MMTJobSpec, CompareJobSpec):
            with self.assertWarns(UserWarning):
                spec = spec_class.from_dict({"powertrain_type": "PHEV/BEV"})
            self.assertEqual(spec.powertrain_type, "PHEV")
        for powertrain in ("PHEV", "BEV"):
            wb = Workbook()
            self.addCleanup(wb.close)
            ws = wb.active
            ws["B6"] = "起步最大jerk"
            ws["D6"] = "PHEV/BEV：<=4m/s3\nICE：<=4.5m/s3"
            apply_metric_standards(ws, "J6B", powertrain)
            self.assertEqual(ws["D6"].value, "<=4m/s3")

    def test_unlabeled_template_retains_legacy_fallback(self):
        wb = Workbook()
        self.addCleanup(wb.close)
        ws = wb.active
        ws["B6"] = "起步最大jerk"
        ws["D6"] = "legacy"
        apply_metric_standards(ws, "MMT", "ICE")
        self.assertEqual(ws["D6"].value, "<=3.5m/s3")


if __name__ == "__main__":
    unittest.main()
