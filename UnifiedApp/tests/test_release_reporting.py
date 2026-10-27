import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from CompareKPI_United.mpi_cpi_compare import (
    MetricValue,
    POWERTRAIN_RELEASE_THRESHOLDS,
    RELEASE_FONT_COLORS,
    assess_release,
    build_report,
    convert_comparison_report_to_latest,
    find_comparison_date_groups,
)


ROOT = Path(__file__).resolve().parents[2] / "CompareKPI_United"


class ReleaseReportingTests(unittest.TestCase):
    def test_retired_metrics_are_not_loaded_by_mining(self):
        for profile in ("J6B", "MMT"):
            with self.subTest(profile=profile):
                result = subprocess.run(
                    [sys.executable, "-c", """
from pathlib import Path
import sys
from DataMining import load_config
config = load_config(Path("config.json"))
if sys.argv[1] == "J6B":
    from DataMining import load_metrics
    names = {metric.name for metric in load_metrics(Path("filters"), ["all"], config)}
    assert "acc_brake_tap_count" in names
else:
    from DataMining import load_metric_modules
    names = {metric.name for pipeline in config.pipelines
             for metric in load_metric_modules(pipeline)}
assert "control_shake_count" not in names, names
assert "follow_stop_fluct" not in names, names
assert "follow_stop_min_jerk" in names, names
assert "slope_no_frnt_vel_fluct" in names, names
print("Loaded metrics:", len(names))
""", profile],
                    cwd=ROOT.parent / f"{profile}_UnifiedMining",
                    env={**os.environ, "PYTHONIOENCODING": "utf-8"},
                    capture_output=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=60,
                )
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_screenshot_thresholds(self):
        names = (
            "起步最大加速度", "起步最大jerk", "初段建立时间", "起步加速建立时间",
            "弯道安全速度", "弯道限速最大减速度", "弯道限速减速舒适性",
            "坡道稳速抖动", "坡道控速稳定性", "坡道控速舒适性",
            "跟车刹停最大减速度", "跟车刹停距离", "跟车刹停舒适性",
            "跟车刹停反应", "跟车刹停风险", "跟车起步响应时延", "跟车起步舒适性",
            "溜车", "无前车稳速未达到SetSpeed", "制动跳变导致的实车加速度抖动",
            "重制动", "顿挫", "Override后顿挫", "Override后重制动",
            "行车断握手", "进Override后顿挫", "起步顿挫", "不合理低速",
        )
        expected = {
            "J6B": (30, 30, 30, 10, 24, 24, 16, 30, 20, 20, 24, 8, 12,
                    24, 12, 24, 16, 500, 200, 500, 80, 80, 500, 500, 1000, 500, 320, 100),
            "MMT": (30, 30, 30, 10, 30, 30, 20, 30, 20, 20, 30, 10, 15,
                    30, 15, 30, 20, 500, 200, 500, 100, 100, 500, 500, 1000, 500, 400, 100),
        }
        for powertrain in ("ICE", "PHEV", "BEV"):
            for profile in ("J6B", "MMT"):
                with self.subTest(powertrain=powertrain, profile=profile):
                    thresholds = POWERTRAIN_RELEASE_THRESHOLDS[powertrain][profile]
                    expected_table = dict(zip(names, expected[profile]))
                    expected_table["点刹"] = 80
                    if profile == "MMT":
                        expected_table.update({
                            "红灯刹停距停止线距离": 30, "红灯刹停舒适性": 10,
                            "绿灯起步响应时延": 15, "绿灯起步最大加速度": 15,
                            "绿灯起步加速时间": 15, "绿灯起步舒适性": 15,
                        })
                    if powertrain == "ICE":
                        expected_table.update(
                            {"顿挫": 70, "起步顿挫": 280} if profile == "J6B"
                            else {"顿挫": 90, "起步顿挫": 360}
                        )
                    self.assertEqual(thresholds, expected_table)
        self.assertIsNot(
            POWERTRAIN_RELEASE_THRESHOLDS["PHEV"]["MMT"],
            POWERTRAIN_RELEASE_THRESHOLDS["BEV"]["MMT"],
        )

    def test_release_boundaries_and_invalid_data(self):
        cases = (
            ("29.9(299/10)", "CPI", "不满足准出"),
            ("30.0(300/10)", "CPI", "可准出"),
            ("30.1(301/10)", "CPI", "可准出"),
            ("30.0(2999/100)", "CPI", "不满足准出"),
            ("100.0(0/0)", "CPI", "样本不足，无法判定"),
            ("100.0(29/0)", "CPI", "样本不足，无法判定"),
            ("100.0(30/0)", "CPI", "可准出"),
            (">29(29/0)", "CPI", "样本不足，无法判定"),
            (">30(30/0)", "CPI", "可准出"),
            (">29.999(29.999/0)", "MPI", "样本不足，无法判定"),
            (">30.000(30.000/0)", "MPI", "可准出"),
            ("29.999(59.998/2)", "MPI", "不满足准出"),
            ("30.000(60.000/2)", "MPI", "可准出"),
            ("100.0（30／0）", "CPI", "可准出"),
            ("30", "CPI", "可准出"),
            (30, "CPI", "可准出"),
            (0, "CPI", "不满足准出"),
            (">20", "MPI", "样本不足，无法判定"),
            (None, "CPI", "数据缺失或无效，无法判定"),
            ("N/A", "CPI", "数据缺失或无效，无法判定"),
            ("ERROR", "CPI", "数据缺失或无效，无法判定"),
            ("N/A(300/1)", "CPI", "数据缺失或无效，无法判定"),
            ("\\", "CPI", "数据缺失或无效，无法判定"),
            ("1.182 (n=3)", "CPI", "数据缺失或无效，无法判定"),
            ("100.0(30/-1)", "CPI", "数据缺失或无效，无法判定"),
            ("100.0(30/0.5)", "CPI", "数据缺失或无效，无法判定"),
            ("1(1/2)", "CPI", "数据缺失或无效，无法判定"),
            ("100.0(nan/0)", "CPI", "数据缺失或无效，无法判定"),
        )
        for raw, kind, expected in cases:
            with self.subTest(raw=raw, kind=kind):
                self.assertEqual(assess_release(MetricValue(raw=raw), 30, kind), expected)
        self.assertEqual(
            assess_release(MetricValue(raw="100(100/1)"), None, "CPI"), "未配置准出标准"
        )

    def test_every_configured_threshold_boundary(self):
        for powertrain, profiles in POWERTRAIN_RELEASE_THRESHOLDS.items():
            for profile, thresholds in profiles.items():
                for metric, threshold in thresholds.items():
                    for kind in ("CPI", "MPI"):
                        with self.subTest(powertrain=powertrain, profile=profile,
                                          metric=metric, kind=kind):
                            for total, expected in (
                                (threshold - 1, "不满足准出"),
                                (threshold, "可准出"),
                                (threshold + 1, "可准出"),
                            ):
                                self.assertEqual(
                                    assess_release(MetricValue(raw=f"{total}({total}/1)"),
                                                   threshold, kind),
                                    expected,
                                )
                            for total, expected in (
                                (threshold - 1, "样本不足，无法判定"),
                                (threshold, "可准出"),
                            ):
                                self.assertEqual(
                                    assess_release(MetricValue(raw=f"100.0({total}/0)"),
                                                   threshold, kind),
                                    expected,
                                )

    @staticmethod
    def _write_input(path, base):
        wb = Workbook()
        ws = wb.active
        ws.title = "所有指标统计汇总"
        ws.append(["指标来源", "指标名称", "指标类型", "总里程(km)",
                   "总次数", "触发次数", "MPI", "CPI"])
        ws.append(["topic", "stop2go_max_jerk", "CPI", 800, 30, 1, None,
                   "29.0(29/1)" if base else "30.0(30/1)"])
        ws.append(["topic", "stop2go_max_acc", "CPI", 800, 29, 0, None, "100.0(29/0)"])
        ws.append(["can", "acc_jerkiness", "MPI", 75, 4, 1, "75.000(75.000/1)", None])
        ws.append(["can", "acc_torque_jerkiness_count", "MPI", 75, 4, 1,
                   "75.000(75.000/1)", None])
        ws.append(["can", "acc_rollback_count", "MPI", 500, 4, 0, ">500.000(500.000/0)", None])
        ws.append(["can", "acc_exit_after_abs", "MPI", 800, 4, 0, ">800(800/0)", None])
        ws.append(["topic", "driving_uninterfered_accel_max_acc_0_30", "CPI",
                   800, 30, 0, None, "100.0(30/0)"])
        ws.append(["can", "acc_brake_tap_count", "MPI", 80 if base else 79, 4, 1,
                   "80.000(80.000/1)" if base else "79.000(79.000/1)", None])
        ws.append(["can", "control_shake_count", "MPI", 800, 4, 1, "800(800/1)", None])
        ws.append(["topic", "follow_stop_fluct", "CPI", 800, 30, 1, None, "30(30/1)"])
        wb.save(path)
        wb.close()

    def test_generated_workbooks_for_all_profiles_and_powertrains(self):
        for profile in ("J6B", "MMT"):
            for powertrain in ("ICE", "PHEV", "BEV"):
                with self.subTest(profile=profile, powertrain=powertrain), tempfile.TemporaryDirectory() as folder:
                    work = Path(folder)
                    base = work / "CAR_20261025_20261025_V1.xlsx"
                    target = work / "CAR_20261026_20261026_V2.xlsx"
                    self._write_input(base, True)
                    self._write_input(target, False)
                    output = work / "comparison.xlsx"
                    build_report(
                        ROOT, ROOT / "templates" / f"template_{profile}.xlsx",
                        output, base, target, True, profile, powertrain_type=powertrain,
                    )
                    wb = load_workbook(output)
                    try:
                        ws = wb["Sheet1"]
                        self.assertEqual(ws.max_column, 10)
                        self.assertEqual(ws["I6"].value, "Base准出标准及结论")
                        self.assertEqual(ws["J6"].value, "Target准出标准及结论")
                        self.assertEqual(ws["E2"].value, powertrain)
                        self.assertEqual(ws["G2"].value, powertrain)
                        self.assertEqual(len(find_comparison_date_groups(ws)), 2)
                        rows = {ws.cell(r, 2).value: r for r in range(7, ws.max_row + 1)}
                        for removed in ("控制抖动", "跟车刹停波动"):
                            self.assertNotIn(removed, rows)
                        if profile == "J6B":
                            for col, expected in ((9, "MPI ≥80\n可准出"),
                                                  (10, "MPI ≥80\n不满足准出")):
                                cell = ws.cell(rows["点刹"], col)
                                self.assertEqual(cell.value, expected)
                                self.assertEqual(
                                    cell.font.color.rgb[-6:],
                                    RELEASE_FONT_COLORS[expected.split("\n")[-1]],
                                )
                        for metric, col, expected in (
                            ("起步最大jerk", 9, "CPI ≥30\n不满足准出"),
                            ("起步最大jerk", 10, "CPI ≥30\n可准出"),
                            ("起步最大加速度", 10, "CPI ≥30\n样本不足，无法判定"),
                            ("初段建立时间", 10, "CPI ≥30\n数据缺失或无效，无法判定"),
                            ("溜车", 10, "MPI ≥500\n可准出"),
                            ("ABS触发关联ACC退出", 10, "未配置准出标准"),
                            ("最大加速度（0-30 km/h）", 10, "未配置准出标准"),
                        ):
                            cell = ws.cell(rows[metric], col)
                            self.assertEqual(cell.value, expected)
                            status = expected.split("\n")[-1]
                            self.assertEqual(cell.font.color.rgb[-6:], RELEASE_FONT_COLORS[status])
                        jerk_threshold = POWERTRAIN_RELEASE_THRESHOLDS[powertrain][profile]["顿挫"]
                        status = "可准出" if 75 >= jerk_threshold else "不满足准出"
                        self.assertEqual(ws.cell(rows["顿挫"], 10).value,
                                         f"MPI ≥{jerk_threshold}\n{status}")
                        self.assertIn("$J$", str(ws.print_area))
                        self.assertTrue(all(ws.column_dimensions[c].width >= 30 for c in ("I", "J")))
                    finally:
                        wb.close()
                    self.assertIsNotNone(convert_comparison_report_to_latest(
                        output, work / "roundtrip.xlsx", "20261026-20261026", profile
                    ))


if __name__ == "__main__":
    unittest.main()
