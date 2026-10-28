import os
from pathlib import Path
import subprocess
import sys
import unittest


ROOT = Path(__file__).resolve().parents[2]


class BadCaseNameTests(unittest.TestCase):
    def test_chinese_names_and_preserved_keys(self):
        for profile in ("J6B", "MMT"):
            with self.subTest(profile=profile):
                result = subprocess.run(
                    [sys.executable, "-c", """
import csv
import json
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
from openpyxl import load_workbook
from metric_display_names import METRIC_DISPLAY_NAMES, metric_display_name

profile = sys.argv[1]
root = Path.cwd()
paths = list((root / "filters").glob("*.py")) if profile == "J6B" else [
    path for kind in ("can", "topic")
    for path in (root / "metrics" / kind / "code").glob("*.py")
]
for path in paths:
    if path.name.startswith("_"):
        continue
    assert path.stem in METRIC_DISPLAY_NAMES, path.stem
sys.path.insert(0, str(root.parent))
from CompareKPI_United.mpi_cpi_compare import SIGNAL_MAP, resolve_signal_mapping
for chinese in SIGNAL_MAP:
    mapping = resolve_signal_mapping(chinese, profile)
    if mapping and mapping[1] in METRIC_DISPLAY_NAMES:
        assert metric_display_name(mapping[1]).replace(" ", "") == chinese, (chinese, mapping)

with tempfile.TemporaryDirectory() as folder:
    work = Path(folder)
    retained = ("follow_stop_min_jerk", "acc_exit_after_abs")
    excluded = tuple(name for name in METRIC_DISPLAY_NAMES
                     if name.startswith("driving_uninterfered_accel_"))
    assert len(excluded) == 12, excluded
    names = retained + excluded
    if profile == "J6B":
        from bad_case_report import write_bad_case_report
        from bad_case_image_report import read_cases
        config = work / "bad_case_extract.json"
        config.write_text(json.dumps({"enabled": True}), encoding="utf-8")
        metrics = [SimpleNamespace(name=name, category="CPI") for name in names]
        for name in names:
            (work / f"{name}_events.json").write_text(json.dumps([{
                "metric_name": name, "source_path": "synthetic.h5",
                "events": [{"start_s": 1, "end_s": 2, "severity": 5}],
            }]), encoding="utf-8")
        output = write_bad_case_report(work, config, metrics)
    else:
        from DataMining import (BadCaseExtractConfig, BadCaseExtractRule,
                                generate_bad_case_extract_report)
        stamp = "2026-10-27_15-00-00"
        for name in names:
            path = work / f"topic__{name}_{stamp}.csv"
            with path.open("w", newline="", encoding="utf-8-sig") as stream:
                writer = csv.DictWriter(stream, fieldnames=[
                    "Bad Case", "严重度得分", "文件夹路径", "符合要求的内容"])
                writer.writeheader()
                writer.writerow({"Bad Case": 1, "严重度得分": 5,
                                 "文件夹路径": "synthetic", "符合要求的内容": "test"})
        output = generate_bad_case_extract_report(
            work, stamp, BadCaseExtractConfig(True, "bad_case_report.xlsx",
                                             BadCaseExtractRule(), {}, {}))
    wb = load_workbook(output)
    try:
        ws = wb.active
        headers = {cell.value: cell.column for cell in ws[1]}
        actual = {ws.cell(row, headers["指标名称"]).value for row in range(2, ws.max_row + 1)}
        assert actual == {metric_display_name(name) for name in retained}, actual
        keys = {ws.cell(row, headers["指标Key"]).value for row in range(2, ws.max_row + 1)}
        expected = set(retained) if profile == "J6B" else {"topic__" + name for name in retained}
        assert keys == expected, keys
        if profile == "J6B":
            cases = read_cases(ws, headers)
            assert {case.metric_name for case in cases} == set(retained)
            del headers["指标Key"]
            legacy = read_cases(ws, headers)
            assert {case.metric_name for case in legacy} == actual
    finally:
        wb.close()
    # Report filtering must leave every per-metric source record untouched.
    for name in names:
        path = (work / f"{name}_events.json" if profile == "J6B"
                else work / f"topic__{name}_{stamp}.csv")
        if profile == "J6B":
            assert json.loads(path.read_text(encoding="utf-8"))[0]["events"][0]["severity"] == 5
        else:
            with path.open(encoding="utf-8-sig", newline="") as stream:
                assert next(csv.DictReader(stream))["Bad Case"] == "1"
    # An excluded-only run still writes the standard empty report.
    if profile == "J6B":
        output = write_bad_case_report(work, config, [m for m in metrics if m.name in excluded])
    else:
        for name in retained:
            (work / f"topic__{name}_{stamp}.csv").unlink()
        output = generate_bad_case_extract_report(
            work, stamp, BadCaseExtractConfig(True, "bad_case_report.xlsx",
                                             BadCaseExtractRule(), {}, {}))
    wb = load_workbook(output)
    try:
        ws = wb.active
        key_column = next(cell.column for cell in ws[1] if cell.value == "指标Key")
        assert all(ws.cell(row, key_column).value is None for row in range(2, ws.max_row + 1))
    finally:
        wb.close()
""", profile],
                    cwd=ROOT / f"{profile}_UnifiedMining",
                    env={**os.environ, "PYTHONIOENCODING": "utf-8"},
                    capture_output=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=60,
                )
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)


if __name__ == "__main__":
    unittest.main()
